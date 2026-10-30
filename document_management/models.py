import os
import uuid

import simplejson as json
from django.contrib.postgres.fields import ArrayField
from django.db import IntegrityError, models, transaction
from django.db.models.functions import Lower

from document_management.sql_func import CYRILLIC_LOWER, CYRILLIC_UPPER
from hospitals.models import Hospitals
from podrazdeleniya.models import Podrazdeleniya
from users.models import DoctorProfile

LAYOUT_TEMPLATE_FIELD_TYPE = 41
ADDRESSEE_FIELD_TYPE = 45
MENTEE_FIELD_TYPE = 46
MENTOR_FIELD_TYPE = 47
PLAN_FIELD_TYPE = 48


def layout_template_is_simple(research):
    from directory.models import ParaclinicInputField

    if not research:
        return False
    return not ParaclinicInputField.objects.filter(group__research=research, field_type=LAYOUT_TEMPLATE_FIELD_TYPE).exists()


class GroupDocuments(models.Model):
    title = models.CharField(max_length=128, blank=True, null=True)

    class Meta:
        verbose_name = "Группа Документов"
        verbose_name_plural = "Группы документов"

    def __str__(self):
        return f"{self.title}"

    @property
    def json(self):
        return {
            "id": self.id,
            "title": self.title or "",
        }

    @staticmethod
    def get_list():
        return [row.json for row in GroupDocuments.objects.all().order_by("title", "pk")]

    @staticmethod
    def save_group(pk, title):
        title = (title or "").strip()
        if not title:
            return {"ok": False, "message": "Укажите название"}
        if pk in (None, -1, "-1"):
            obj = GroupDocuments(title=title)
        else:
            obj = GroupDocuments.objects.filter(pk=pk).first()
            if not obj:
                return {"ok": False, "message": "Группа не найдена"}
            obj.title = title
        obj.save()
        return {"ok": True, "id": obj.pk, "title": obj.title}


def _fold_doctor_name(field_name):
    return Lower(models.Func(models.F(field_name), models.Value(CYRILLIC_UPPER), models.Value(CYRILLIC_LOWER), function="TRANSLATE", output_field=models.CharField()))


def _fold_search(search):
    return search.translate(str.maketrans(CYRILLIC_UPPER, CYRILLIC_LOWER)).lower()


class AddresseeGroup(models.Model):
    ALL_ID = 0
    ALL_TITLE = "Все"
    ALL_GROUP = "Все адресаты"
    CONSTRUCTOR_GROUP = "Конструктор: ДОУ"

    title = models.CharField(max_length=128, blank=True, null=True)
    hide = models.BooleanField(default=False, blank=True, db_index=True)
    order = models.IntegerField(default=0)
    who_create = models.ForeignKey(
        DoctorProfile,
        related_name="addressee_groups",
        db_index=True,
        default=None,
        blank=True,
        null=True,
        help_text="Автор личного набора; пусто — общий набор конструктора",
        on_delete=models.SET_NULL,
    )

    class Meta:
        verbose_name = "Набор адресатов"
        verbose_name_plural = "Наборы адресатов"
        ordering = ("order", "title", "pk")

    def __str__(self):
        return self.title or ""

    @property
    def json(self):
        return {
            "id": self.id,
            "title": self.title or "",
            "hide": self.hide,
            "order": self.order,
            "own": False,
        }

    def as_json(self, doctor_id=None):
        data = self.json
        data["own"] = bool(self.who_create_id and doctor_id and self.who_create_id == doctor_id)
        return data

    @staticmethod
    def _employee_json(doctor):
        return {
            "id": doctor.pk,
            "fio": doctor.get_fio(),
            "department": doctor.podrazdeleniye.title if doctor.podrazdeleniye else "",
        }

    @classmethod
    def can_see_all(cls, user):
        if not user:
            return False
        if getattr(user, "is_superuser", False):
            return True
        return user.groups.filter(name=cls.ALL_GROUP).exists()

    @classmethod
    def can_list_all_employees(cls, user, doctor=None):
        if cls.can_see_all(user):
            return True
        if doctor and doctor.has_group(cls.CONSTRUCTOR_GROUP):
            return True
        return bool(user and user.groups.filter(name=cls.CONSTRUCTOR_GROUP).exists())

    @classmethod
    def empty_employees(cls, page=1, page_size=100):
        return {
            "result": [],
            "total": 0,
            "page": page,
            "pageSize": page_size if page_size is not None else 0,
            "hasMore": False,
        }

    @classmethod
    def get_list(cls, include_all=True, include_hidden=False, doctor_id=None, global_only=False):
        rows = cls.objects.all().order_by("order", "title", "pk")
        if global_only or not doctor_id:
            rows = rows.filter(who_create__isnull=True)
        else:
            rows = rows.filter(models.Q(who_create__isnull=True) | models.Q(who_create_id=doctor_id))
        if not include_hidden:
            rows = rows.filter(hide=False)
        result = []
        if include_all:
            result.append({"id": cls.ALL_ID, "title": cls.ALL_TITLE, "hide": False, "order": -1, "own": False})
        result.extend([row.as_json(doctor_id) for row in rows])
        return result

    @classmethod
    def get_details(cls, pk):
        if pk in (None, cls.ALL_ID, 0, "0", -1, "-1"):
            return {"ok": False, "message": "Набор не найден"}
        obj = cls.objects.filter(pk=pk).first()
        if not obj:
            return {"ok": False, "message": "Набор не найден"}
        members = [cls._employee_json(row.doctor) for row in obj.members.select_related("doctor__podrazdeleniye").order_by("doctor__family", "doctor__name", "pk") if row.doctor_id]
        payload = obj.json
        payload["ok"] = True
        payload["members"] = members
        return payload

    @classmethod
    def _set_members(cls, obj, member_ids):
        ids = []
        seen = set()
        for raw in member_ids:
            try:
                doctor_id = int(raw)
            except (TypeError, ValueError):
                continue
            if doctor_id < 1 or doctor_id in seen:
                continue
            seen.add(doctor_id)
            ids.append(doctor_id)
        valid = set(DoctorProfile.objects.filter(pk__in=ids).values_list("pk", flat=True))
        ids = [doctor_id for doctor_id in ids if doctor_id in valid]
        obj.members.all().delete()
        AddresseeGroupMember.objects.bulk_create([AddresseeGroupMember(group=obj, doctor_id=doctor_id) for doctor_id in ids])
        return ids

    @classmethod
    def save_group(cls, pk, title, hide=False, member_ids=None):
        title = (title or "").strip()
        if not title:
            return {"ok": False, "message": "Укажите название"}
        if pk in (None, -1, "-1"):
            obj = cls(title=title, hide=bool(hide))
        else:
            obj = cls.objects.filter(pk=pk, who_create__isnull=True).first()
            if not obj:
                return {"ok": False, "message": "Набор не найден"}
            obj.title = title
            obj.hide = bool(hide)
        obj.save()
        if member_ids is not None:
            cls._set_members(obj, member_ids)
        return {"ok": True, "id": obj.pk, "title": obj.title}

    @classmethod
    def save_personal(cls, who_create, title, member_ids):
        if not who_create:
            return {"ok": False, "message": "Нет пользователя"}
        title = (title or "").strip()
        if not title:
            return {"ok": False, "message": "Укажите название"}
        if not member_ids:
            return {"ok": False, "message": "Выберите сотрудников"}
        obj = cls(title=title, who_create=who_create)
        obj.save()
        cls._set_members(obj, member_ids)
        return {"ok": True, "id": obj.pk, "title": obj.title, "own": True}

    @classmethod
    def delete_personal(cls, pk, who_create):
        if not who_create:
            return {"ok": False, "message": "Нет пользователя"}
        obj = cls.objects.filter(pk=pk, who_create=who_create).first()
        if not obj:
            return {"ok": False, "message": "Набор не найден"}
        obj.delete()
        return {"ok": True}

    @classmethod
    def add_personal_member(cls, pk, who_create, doctor_id):
        if not who_create:
            return {"ok": False, "message": "Нет пользователя"}
        obj = cls.objects.filter(pk=pk, who_create=who_create).first()
        if not obj:
            return {"ok": False, "message": "Набор не найден"}
        try:
            doctor_id = int(doctor_id)
        except (TypeError, ValueError):
            doctor_id = 0
        doctor = DoctorProfile.objects.filter(pk=doctor_id, user__is_active=True, is_system_user=False).first()
        if not doctor:
            return {"ok": False, "message": "Сотрудник не найден"}
        AddresseeGroupMember.objects.get_or_create(group=obj, doctor_id=doctor.pk)
        return {"ok": True, "id": obj.pk}

    @classmethod
    def get_hospitals(cls):
        rows = Hospitals.objects.filter(hide=False).order_by("-is_default", "short_title", "title", "pk")
        return [{"id": row.pk, "title": row.short_title or row.title, "isDefault": row.is_default} for row in rows]

    @classmethod
    def _hospital_pk(cls, hospital_id):
        try:
            hospital_pk = int(hospital_id)
        except (TypeError, ValueError):
            hospital_pk = 0
        if hospital_pk > 0 and Hospitals.objects.filter(pk=hospital_pk, hide=False).exists():
            return hospital_pk
        default = Hospitals.get_default_hospital()
        return default.pk if default else None

    @classmethod
    def get_employees(cls, group_id, hospital_id, query="", page=1, page_size=100, allow_all=False):
        try:
            group_id = int(group_id)
        except (TypeError, ValueError):
            group_id = cls.ALL_ID
        hospital_pk = cls._hospital_pk(hospital_id)
        if not hospital_pk:
            return cls.empty_employees(page, page_size)
        if (not group_id or group_id <= 0) and not allow_all:
            return cls.empty_employees(page, page_size)
        doctors = DoctorProfile.objects.filter(user__is_active=True, is_system_user=False, podrazdeleniye__hospital_id=hospital_pk).select_related("user", "podrazdeleniye")
        if group_id and group_id > 0:
            doctors = doctors.filter(addressee_group_members__group_id=group_id)
        query = (query or "").strip()
        if query:
            doctors = doctors.annotate(
                family_fold=_fold_doctor_name("family"),
                name_fold=_fold_doctor_name("name"),
                patronymic_fold=_fold_doctor_name("patronymic"),
                fio_fold=_fold_doctor_name("fio"),
            )
            parts = [p for p in query.split() if p]
            q_filter = models.Q()
            for part in parts:
                needle = _fold_search(part)
                q_filter &= models.Q(family_fold__contains=needle) | models.Q(name_fold__contains=needle) | models.Q(patronymic_fold__contains=needle) | models.Q(fio_fold__contains=needle)
            doctors = doctors.filter(q_filter)
        doctors = doctors.distinct().order_by("family", "name", "patronymic", "pk")
        total = doctors.count()
        if page_size is None:
            page = 1
            rows = doctors
            has_more = False
        else:
            try:
                page = max(int(page or 1), 1)
            except (TypeError, ValueError):
                page = 1
            try:
                page_size = max(int(page_size or 100), 1)
            except (TypeError, ValueError):
                page_size = 100
            start = (page - 1) * page_size
            rows = doctors[start : start + page_size]
            has_more = start + page_size < total
        return {
            "result": [cls._employee_json(row) for row in rows],
            "total": total,
            "page": page,
            "pageSize": page_size if page_size is not None else total,
            "hasMore": has_more,
        }


class AddresseeGroupMember(models.Model):
    group = models.ForeignKey(AddresseeGroup, related_name="members", on_delete=models.CASCADE)
    doctor = models.ForeignKey(DoctorProfile, related_name="addressee_group_members", on_delete=models.CASCADE)

    class Meta:
        verbose_name = "Сотрудник набора адресатов"
        verbose_name_plural = "Сотрудники наборов адресатов"
        unique_together = ("group", "doctor")

    def __str__(self):
        return f"{self.group} {self.doctor}"


ACCESS_MODE_BLACK = "black"
ACCESS_MODE_WHITE = "white"


def normalize_access_mode(value, default=ACCESS_MODE_BLACK):
    if value in (None, ""):
        return default
    return ACCESS_MODE_WHITE if str(value) == ACCESS_MODE_WHITE else ACCESS_MODE_BLACK


def access_visible_queryset(qs, access_model, fk_name, doctor):
    if not doctor:
        return qs.exclude(access_mode=ACCESS_MODE_WHITE)
    listed = access_model.objects.filter(**{fk_name: models.OuterRef("pk"), "doctor_id": doctor.pk})
    return qs.filter((models.Q(access_mode=ACCESS_MODE_WHITE) & models.Exists(listed)) | (~models.Q(access_mode=ACCESS_MODE_WHITE) & ~models.Exists(listed)))


def access_members_json(rows):
    return [AddresseeGroup._employee_json(row.doctor) for row in rows if getattr(row, "doctor_id", None)]


def legacy_access_row(row, columns_default, extra=None):
    data = {
        "id": row.pk,
        "title": row.title or "",
        "columnsCount": row.columns_count or columns_default,
        "accessMode": ACCESS_MODE_BLACK,
        "accessMembers": [],
    }
    if extra:
        data.update(extra)
    return data


def replace_access_members(access_model, fk_name, obj, member_ids):
    ids = TypeDocuments._parse_creator_ids(member_ids)
    valid = set(DoctorProfile.objects.filter(pk__in=ids).values_list("pk", flat=True))
    ids = [doctor_id for doctor_id in ids if doctor_id in valid]
    access_model.objects.filter(**{fk_name: obj}).delete()
    access_model.objects.bulk_create([access_model(**{fk_name: obj, "doctor_id": doctor_id}) for doctor_id in ids])
    return ids


class TypeSection(models.Model):
    COLUMNS_MIN = 1
    COLUMNS_MAX = 12
    COLUMNS_DEFAULT = 10

    title = models.CharField(max_length=128, blank=True, null=True)
    access_mode = models.CharField(max_length=8, default=ACCESS_MODE_BLACK, db_index=True)
    columns_count = models.PositiveSmallIntegerField(
        default=COLUMNS_DEFAULT,
        help_text="Количество колонок значений в ResearchesPicker",
    )

    class Meta:
        verbose_name = "Тип раздела"
        verbose_name_plural = "Типы разделов"

    def __str__(self):
        return f"{self.title}"

    def get_access_members(self):
        prefetched = getattr(self, "_prefetched_objects_cache", {}).get("access_members")
        if prefetched is not None:
            rows = prefetched
        else:
            rows = self.access_members.select_related("doctor__podrazdeleniye").order_by("doctor__family", "doctor__name", "pk")
        return access_members_json(rows)

    @property
    def json(self):
        return {
            "id": self.id,
            "title": self.title or "",
            "columnsCount": self.columns_count or self.COLUMNS_DEFAULT,
            "accessMode": self.access_mode or ACCESS_MODE_BLACK,
            "accessMembers": self.get_access_members(),
        }

    @staticmethod
    def _queryset():
        return TypeSection.objects.prefetch_related(
            models.Prefetch(
                "access_members",
                queryset=TypeSectionAccess.objects.select_related("doctor__podrazdeleniye").order_by("doctor__family", "doctor__name", "pk"),
            )
        ).order_by("title", "pk")

    @staticmethod
    def get_list():
        from django.db import connection
        from django.db.utils import OperationalError, ProgrammingError

        try:
            return [row.json for row in TypeSection._queryset()]
        except (OperationalError, ProgrammingError):
            connection.rollback()
            return [legacy_access_row(row, TypeSection.COLUMNS_DEFAULT) for row in TypeSection.objects.defer("access_mode").order_by("title", "pk")]

    @staticmethod
    def get_visible_list(doctor):
        from django.db import connection
        from django.db.utils import OperationalError, ProgrammingError

        try:
            qs = access_visible_queryset(TypeSection._queryset(), TypeSectionAccess, "type_section_id", doctor)
            return [row.json for row in qs]
        except (OperationalError, ProgrammingError):
            connection.rollback()
            return TypeSection.get_list()

    @classmethod
    def normalize_columns_count(cls, value):
        if value in (None, ""):
            return None
        try:
            n = int(value)
        except (TypeError, ValueError):
            return None
        if n < cls.COLUMNS_MIN or n > cls.COLUMNS_MAX:
            return None
        return n

    @staticmethod
    def save_section(pk, title, columns_count=None, access_mode=None, access_member_ids=None):
        title = (title or "").strip()
        if not title:
            return {"ok": False, "message": "Укажите название"}
        provided = columns_count not in (None, "")
        n = TypeSection.normalize_columns_count(columns_count)
        if provided and n is None:
            return {
                "ok": False,
                "message": f"Количество колонок от {TypeSection.COLUMNS_MIN} до {TypeSection.COLUMNS_MAX}",
            }
        mode = normalize_access_mode(access_mode)
        with transaction.atomic():
            if pk in (None, -1, "-1"):
                obj = TypeSection(
                    title=title,
                    columns_count=n if n is not None else TypeSection.COLUMNS_DEFAULT,
                    access_mode=mode,
                )
            else:
                obj = TypeSection.objects.filter(pk=pk).first()
                if not obj:
                    return {"ok": False, "message": "Тип раздела не найден"}
                obj.title = title
                if provided:
                    obj.columns_count = n
                if access_mode not in (None, ""):
                    obj.access_mode = mode
            obj.save()
            if access_member_ids is not None:
                replace_access_members(TypeSectionAccess, "type_section", obj, access_member_ids)
        return {"ok": True, "id": obj.pk, "title": obj.title, "columnsCount": obj.columns_count}


class TypeSectionAccess(models.Model):
    type_section = models.ForeignKey(TypeSection, related_name="access_members", on_delete=models.CASCADE)
    doctor = models.ForeignKey(DoctorProfile, related_name="type_section_access_rows", on_delete=models.CASCADE)

    class Meta:
        verbose_name = "Доступ к типу раздела"
        verbose_name_plural = "Доступ к типам разделов"
        unique_together = ("type_section", "doctor")

    def __str__(self):
        return f"{self.type_section} {self.doctor}"


class PlaceSection(models.Model):
    COLUMNS_MIN = 1
    COLUMNS_MAX = 12
    COLUMNS_DEFAULT = 10

    title = models.CharField(max_length=128, blank=True, null=True)
    access_mode = models.CharField(max_length=8, default=ACCESS_MODE_BLACK, db_index=True)
    columns_count = models.PositiveSmallIntegerField(
        default=COLUMNS_DEFAULT,
        help_text="Количество колонок значений в ResearchesPicker",
    )
    type_section = models.ForeignKey(
        TypeSection,
        related_name="places",
        default=None,
        blank=True,
        null=True,
        db_index=True,
        help_text="Тип раздела",
        on_delete=models.SET_NULL,
    )

    class Meta:
        verbose_name = "Подраздел/место"
        verbose_name_plural = "Подразделы/места"

    def __str__(self):
        return f"{self.title}"

    def get_access_members(self):
        prefetched = getattr(self, "_prefetched_objects_cache", {}).get("access_members")
        if prefetched is not None:
            rows = prefetched
        else:
            rows = self.access_members.select_related("doctor__podrazdeleniye").order_by("doctor__family", "doctor__name", "pk")
        return access_members_json(rows)

    @property
    def json(self):
        section = self.type_section if self.type_section_id else None
        return {
            "id": self.id,
            "title": self.title or "",
            "columnsCount": self.columns_count or self.COLUMNS_DEFAULT,
            "typeSectionId": self.type_section_id,
            "typeSectionTitle": section.title if section else "",
            "accessMode": self.access_mode or ACCESS_MODE_BLACK,
            "accessMembers": self.get_access_members(),
        }

    @staticmethod
    def _queryset():
        return (
            PlaceSection.objects.select_related("type_section")
            .prefetch_related(
                models.Prefetch(
                    "access_members",
                    queryset=PlaceSectionAccess.objects.select_related("doctor__podrazdeleniye").order_by("doctor__family", "doctor__name", "pk"),
                )
            )
            .order_by("title", "pk")
        )

    @staticmethod
    def get_list():
        from django.db import connection
        from django.db.utils import OperationalError, ProgrammingError

        try:
            return [row.json for row in PlaceSection._queryset()]
        except (OperationalError, ProgrammingError):
            connection.rollback()
            rows = []
            qs = PlaceSection.objects.select_related("type_section").defer("access_mode", "type_section__access_mode").order_by("title", "pk")
            for row in qs:
                section = row.type_section if row.type_section_id else None
                rows.append(
                    legacy_access_row(
                        row,
                        PlaceSection.COLUMNS_DEFAULT,
                        {
                            "typeSectionId": row.type_section_id,
                            "typeSectionTitle": section.title if section else "",
                        },
                    )
                )
            return rows

    @classmethod
    def normalize_columns_count(cls, value):
        if value in (None, ""):
            return None
        try:
            n = int(value)
        except (TypeError, ValueError):
            return None
        if n < cls.COLUMNS_MIN or n > cls.COLUMNS_MAX:
            return None
        return n

    def as_picker_department(self, dou):
        section = self.type_section if self.type_section_id else None
        return {
            "pk": self.pk,
            "title": self.title or "",
            "type": dou,
            "extended": True,
            "e": dou,
            "columnsCount": self.columns_count or self.COLUMNS_DEFAULT,
            "typeSectionId": self.type_section_id,
            "typeSectionTitle": section.title if section else "",
        }

    @staticmethod
    def save_place(pk, title, columns_count=None, type_section_id=None, access_mode=None, access_member_ids=None):
        title = (title or "").strip()
        if not title:
            return {"ok": False, "message": "Укажите название"}
        provided = columns_count not in (None, "")
        n = PlaceSection.normalize_columns_count(columns_count)
        if provided and n is None:
            return {
                "ok": False,
                "message": f"Количество колонок от {PlaceSection.COLUMNS_MIN} до {PlaceSection.COLUMNS_MAX}",
            }
        section = None
        if type_section_id not in (None, "", -1, "-1"):
            section = TypeSection.objects.filter(pk=type_section_id).first()
            if not section:
                return {"ok": False, "message": "Тип раздела не найден"}
        mode = normalize_access_mode(access_mode)
        with transaction.atomic():
            if pk in (None, -1, "-1"):
                obj = PlaceSection(
                    title=title,
                    columns_count=n if n is not None else PlaceSection.COLUMNS_DEFAULT,
                    type_section=section,
                    access_mode=mode,
                )
            else:
                obj = PlaceSection.objects.filter(pk=pk).first()
                if not obj:
                    return {"ok": False, "message": "Подраздел/место не найден"}
                obj.title = title
                if provided:
                    obj.columns_count = n
                obj.type_section = section
                if access_mode not in (None, ""):
                    obj.access_mode = mode
            obj.save()
            if access_member_ids is not None:
                replace_access_members(PlaceSectionAccess, "place_section", obj, access_member_ids)
        return {
            "ok": True,
            "id": obj.pk,
            "title": obj.title,
            "columnsCount": obj.columns_count,
            "typeSectionId": obj.type_section_id,
        }

    @staticmethod
    def picker_departments(doctor=None):
        from podrazdeleniya.models import Podrazdeleniya

        dou = Podrazdeleniya.DOU
        rows = []
        if TypeDocuments.objects.filter(place_section__isnull=True).exists():
            rows.append(
                {
                    "pk": -1,
                    "title": "Общие",
                    "type": dou,
                    "extended": True,
                    "e": dou,
                    "columnsCount": PlaceSection.COLUMNS_DEFAULT,
                }
            )
        from django.db import connection
        from django.db.utils import OperationalError, ProgrammingError

        try:
            places = access_visible_queryset(
                PlaceSection.objects.select_related("type_section").order_by("title", "pk"),
                PlaceSectionAccess,
                "place_section_id",
                doctor,
            )
            places = list(places)
        except (OperationalError, ProgrammingError):
            connection.rollback()
            places = PlaceSection.objects.select_related("type_section").defer("access_mode", "type_section__access_mode").order_by("title", "pk")
        for place in places:
            rows.append(place.as_picker_department(dou))
        return rows

    @staticmethod
    def picker_researches(doctor=None, available_only=False):
        docs = TypeDocuments.picker_researches(doctor=doctor, available_only=available_only)
        cases = TypeCases.picker_researches()
        if doctor is None:
            return docs + cases
        from django.db import connection
        from django.db.utils import OperationalError, ProgrammingError

        try:
            visible_places = set(access_visible_queryset(PlaceSection.objects.all(), PlaceSectionAccess, "place_section_id", doctor).values_list("pk", flat=True))
        except (OperationalError, ProgrammingError):
            connection.rollback()
            return docs + cases
        return [row for row in docs + cases if not row.get("site_type") or row.get("site_type") in visible_places]


class PlaceSectionAccess(models.Model):
    place_section = models.ForeignKey(PlaceSection, related_name="access_members", on_delete=models.CASCADE)
    doctor = models.ForeignKey(DoctorProfile, related_name="place_section_access_rows", on_delete=models.CASCADE)

    class Meta:
        verbose_name = "Доступ к подразделу"
        verbose_name_plural = "Доступ к подразделам"
        unique_together = ("place_section", "doctor")

    def __str__(self):
        return f"{self.place_section} {self.doctor}"


def get_print_docx_path(instance, filename):
    return os.path.join("dou-print-docx", str(instance.pk or "new"), str(uuid.uuid4()), filename)


class TypeDocuments(models.Model):
    PICKER_PK_SHIFT = 1000000000

    title = models.CharField(max_length=128, blank=True, null=True)
    group_document = models.ForeignKey(GroupDocuments, default=None, blank=True, null=True, help_text="Группа документов", on_delete=models.SET_NULL)
    place_section = models.ForeignKey(
        PlaceSection,
        related_name="type_documents",
        default=None,
        blank=True,
        null=True,
        db_index=True,
        help_text="Подраздел/место в пикере ДОУ",
        on_delete=models.SET_NULL,
    )
    code = models.CharField(max_length=55, blank=True, null=True)
    layout_template = models.ForeignKey(
        "directory.Researches",
        default=None,
        blank=True,
        null=True,
        db_index=True,
        help_text="Шаблон документа",
        on_delete=models.SET_NULL,
    )
    monitoring = models.ForeignKey(
        "directory.Researches",
        related_name="monitoring_document_types",
        default=None,
        blank=True,
        null=True,
        db_index=True,
        help_text="Ссылка на мониторинг",
        on_delete=models.SET_NULL,
    )
    test = models.ForeignKey(
        "DouTest",
        related_name="document_types",
        default=None,
        blank=True,
        null=True,
        db_index=True,
        help_text="Ссылка на тест",
        on_delete=models.SET_NULL,
    )
    is_plan = models.BooleanField(default=False, blank=True, help_text="План")
    print_docx = models.FileField(upload_to=get_print_docx_path, default=None, null=True, blank=True, help_text="Шаблон печати docx")

    class Meta:
        verbose_name = "Вид документа"
        verbose_name_plural = "Виды документов"

    def __str__(self):
        return f"{self.title}"

    def get_layout_templates(self):
        prefetched = getattr(self, "_prefetched_objects_cache", {}).get("layout_template_links")
        if prefetched is not None:
            templates = [row.layout_template for row in prefetched if row.layout_template_id]
        else:
            templates = [row.layout_template for row in self.layout_template_links.select_related("layout_template").order_by("order", "pk") if row.layout_template_id]
        if templates:
            return templates
        if self.layout_template_id and self.layout_template:
            return [self.layout_template]
        return []

    def get_creators(self):
        prefetched = getattr(self, "_prefetched_objects_cache", {}).get("creators")
        if prefetched is not None:
            rows = prefetched
        else:
            rows = self.creators.select_related("doctor__podrazdeleniye").order_by("doctor__family", "doctor__name", "pk")
        return [AddresseeGroup._employee_json(row.doctor) for row in rows if row.doctor_id]

    def can_create(self, doctor):
        if not doctor:
            return False
        creator_ids = list(self.creators.values_list("doctor_id", flat=True))
        if not creator_ids:
            return True
        return doctor.pk in creator_ids

    @property
    def json(self):
        templates = self.get_layout_templates()
        template_ids = [row.pk for row in templates]
        return {
            "id": self.id,
            "title": self.title or "",
            "code": self.code or "",
            "groupId": self.group_document_id,
            "groupTitle": self.group_document.title if self.group_document else "",
            "placeSectionId": self.place_section_id,
            "placeSectionTitle": self.place_section.title if self.place_section else "",
            "layoutTemplateId": template_ids[0] if template_ids else self.layout_template_id,
            "layoutTemplateIds": template_ids,
            "layoutTemplates": [{"id": row.pk, "label": row.title} for row in templates],
            "monitoringId": self.monitoring_id,
            "monitoringTitle": self.monitoring.title if self.monitoring_id and self.monitoring else "",
            "testId": self.test_id,
            "testTitle": self.test.title if self.test_id and self.test else "",
            "isPlan": bool(self.is_plan),
            "creators": self.get_creators(),
            "hasPrintTemplate": bool(self.print_docx),
            "printTemplateName": os.path.basename(self.print_docx.name) if self.print_docx else "",
        }

    def save_print_docx(self, uploaded):
        filename = os.path.basename(getattr(uploaded, "name", "") or "")
        if not filename.lower().endswith(".docx"):
            return {"ok": False, "message": "Нужен файл .docx"}
        if getattr(uploaded, "size", 0) > 5 * 1024 * 1024:
            return {"ok": False, "message": "Файл больше 5 МБ"}
        if self.print_docx:
            self.print_docx.delete(save=False)
        self.print_docx = uploaded
        self.save(update_fields=["print_docx"])
        return {"ok": True, "fileName": os.path.basename(self.print_docx.name)}

    def clear_print_docx(self):
        if self.print_docx:
            self.print_docx.delete(save=True)
        return {"ok": True}

    @staticmethod
    def get_list(group_id=None, doctor=None, available_only=False):
        qs = (
            TypeDocuments.objects.select_related("group_document", "layout_template", "monitoring", "test", "place_section")
            .prefetch_related(
                models.Prefetch(
                    "layout_template_links",
                    queryset=TypeDocumentLayoutTemplate.objects.select_related("layout_template").order_by("order", "pk"),
                ),
                models.Prefetch(
                    "creators",
                    queryset=TypeDocumentCreator.objects.select_related("doctor__podrazdeleniye").order_by("doctor__family", "doctor__name", "pk"),
                ),
            )
            .order_by("title", "pk")
        )
        if group_id in (-1, "-1"):
            qs = qs.filter(group_document__isnull=True)
        elif group_id not in (None, "", 0, "0"):
            qs = qs.filter(group_document_id=group_id)
        if available_only:
            if not doctor:
                return []
            has_creators = TypeDocumentCreator.objects.filter(type_document_id=models.OuterRef("pk"))
            is_creator = TypeDocumentCreator.objects.filter(type_document_id=models.OuterRef("pk"), doctor=doctor)
            qs = qs.filter(~models.Exists(has_creators) | models.Exists(is_creator))
        return [row.json for row in qs]

    @staticmethod
    def _parse_template_ids(layout_template_ids, layout_template_id):
        raw = layout_template_ids
        if raw is None:
            raw = [layout_template_id] if layout_template_id not in (None, "", -1, "-1") else []
        if not isinstance(raw, (list, tuple)):
            raw = [raw]
        seen = set()
        result = []
        for item in raw:
            if item in (None, "", -1, "-1"):
                continue
            try:
                value = int(item)
            except (TypeError, ValueError):
                continue
            if value <= 0 or value in seen:
                continue
            seen.add(value)
            result.append(value)
        return result

    @staticmethod
    def _parse_creator_ids(creator_ids):
        if not isinstance(creator_ids, (list, tuple)):
            creator_ids = [] if creator_ids in (None, "", -1, "-1") else [creator_ids]
        seen = set()
        result = []
        for item in creator_ids:
            if isinstance(item, dict):
                item = item.get("id")
            if item in (None, "", -1, "-1"):
                continue
            try:
                value = int(item)
            except (TypeError, ValueError):
                continue
            if value <= 0 or value in seen:
                continue
            seen.add(value)
            result.append(value)
        return result

    @staticmethod
    def _set_creators(obj, creator_ids):
        ids = TypeDocuments._parse_creator_ids(creator_ids)
        valid = set(DoctorProfile.objects.filter(pk__in=ids).values_list("pk", flat=True))
        ids = [doctor_id for doctor_id in ids if doctor_id in valid]
        TypeDocumentCreator.objects.filter(type_document=obj).delete()
        TypeDocumentCreator.objects.bulk_create([TypeDocumentCreator(type_document=obj, doctor_id=doctor_id) for doctor_id in ids])
        return ids

    def as_picker_research(self):
        from podrazdeleniya.models import Podrazdeleniya

        title = self.title or ""
        code = self.code or ""
        site = self.place_section_id
        return {
            "pk": TypeDocuments.PICKER_PK_SHIFT + self.pk,
            "onlywith": -1,
            "department_pk": 2 - Podrazdeleniya.DOU,
            "title": title,
            "full_title": title,
            "doc_refferal": False,
            "treatment": False,
            "is_hospital": False,
            "is_form": False,
            "is_case": False,
            "is_application": False,
            "stom": False,
            "need_vich_code": False,
            "comment_variants": [],
            "autoadd": [],
            "addto": [],
            "code": code,
            "internal_code": code,
            "type": str(Podrazdeleniya.DOU),
            "site_type": site,
            "site_type_raw": site,
            "localizations": [],
            "service_locations": [],
            "direction_params": -1,
            "is_dou_document_type": True,
        }

    @staticmethod
    def picker_researches(doctor=None, available_only=False):
        qs = TypeDocuments.objects.select_related("place_section").order_by("title", "pk")
        if available_only:
            if not doctor:
                return []
            has_creators = TypeDocumentCreator.objects.filter(type_document_id=models.OuterRef("pk"))
            is_creator = TypeDocumentCreator.objects.filter(type_document_id=models.OuterRef("pk"), doctor=doctor)
            qs = qs.filter(~models.Exists(has_creators) | models.Exists(is_creator))
        return [row.as_picker_research() for row in qs]

    def form_research(self):
        if self.is_plan or self.test_id:
            return None
        monitoring = self.monitoring if self.monitoring_id else None
        if monitoring and monitoring.is_monitoring:
            return monitoring
        templates = self.get_layout_templates()
        if templates:
            return templates[0]
        return self.layout_template if self.layout_template_id else None

    @staticmethod
    def save_type(
        pk,
        title,
        group_id=None,
        code="",
        layout_template_id=None,
        layout_template_ids=None,
        creator_ids=None,
        place_section_id=None,
        monitoring_id=None,
        test_id=None,
        is_plan=False,
    ):
        from directory.models import Researches

        title = (title or "").strip()
        if not title:
            return {"ok": False, "message": "Укажите название"}
        group = None
        if group_id not in (None, "", -1, "-1"):
            group = GroupDocuments.objects.filter(pk=group_id).first()
            if not group:
                return {"ok": False, "message": "Группа не найдена"}
        place = None
        if place_section_id not in (None, "", -1, "-1"):
            place = PlaceSection.objects.filter(pk=place_section_id).first()
            if not place:
                return {"ok": False, "message": "Подраздел/место не найден"}
        template_ids = TypeDocuments._parse_template_ids(layout_template_ids, layout_template_id)
        templates = []
        for template_id in template_ids:
            layout_template = Researches.objects.filter(pk=template_id, is_layout_template=True).first()
            if not layout_template:
                return {"ok": False, "message": "Шаблон не найден"}
            if not layout_template_is_simple(layout_template):
                return {"ok": False, "message": f"Шаблон «{layout_template.title}» содержит вложенный шаблон"}
            templates.append(layout_template)
        first_template = templates[0] if templates else None
        monitoring = None
        if monitoring_id not in (None, "", -1, "-1"):
            monitoring = Researches.objects.filter(pk=monitoring_id, is_monitoring=True).first()
            if not monitoring:
                return {"ok": False, "message": "Мониторинг не найден"}
        test = None
        if test_id not in (None, "", -1, "-1"):
            test = DouTest.objects.filter(pk=test_id).first()
            if not test:
                return {"ok": False, "message": "Тест не найден"}
        plan_flag = bool(is_plan)
        if plan_flag and (monitoring or test):
            return {"ok": False, "message": "Укажите либо план, либо мониторинг, либо тест"}
        if monitoring and test:
            return {"ok": False, "message": "Укажите либо мониторинг, либо тест"}
        with transaction.atomic():
            if pk in (None, -1, "-1"):
                obj = TypeDocuments(
                    title=title,
                    group_document=group,
                    place_section=place,
                    code=code or "",
                    layout_template=first_template,
                    monitoring=monitoring,
                    test=test,
                    is_plan=plan_flag,
                )
            else:
                obj = TypeDocuments.objects.filter(pk=pk).first()
                if not obj:
                    return {"ok": False, "message": "Вид документа не найден"}
                obj.title = title
                obj.group_document = group
                obj.place_section = place
                obj.code = code or ""
                obj.layout_template = first_template
                obj.monitoring = monitoring
                obj.test = test
                obj.is_plan = plan_flag
            obj.save()
            TypeDocumentLayoutTemplate.objects.filter(type_document=obj).delete()
            TypeDocumentLayoutTemplate.objects.bulk_create([TypeDocumentLayoutTemplate(type_document=obj, layout_template=row, order=index) for index, row in enumerate(templates)])
            if creator_ids is not None:
                TypeDocuments._set_creators(obj, creator_ids)
            TypeDocumentsSchema.get_or_create_for_type(obj)
        return {"ok": True, "id": obj.pk, "title": obj.title}


class TypeDocumentLayoutTemplate(models.Model):
    type_document = models.ForeignKey(TypeDocuments, related_name="layout_template_links", on_delete=models.CASCADE)
    layout_template = models.ForeignKey("directory.Researches", related_name="type_document_links", on_delete=models.CASCADE)
    order = models.IntegerField(default=0, db_index=True)

    class Meta:
        verbose_name = "Шаблон вида документа"
        verbose_name_plural = "Шаблоны видов документов"
        ordering = ["order", "pk"]
        unique_together = ("type_document", "layout_template")

    def __str__(self):
        return f"{self.type_document} – {self.layout_template} ({self.order})"


class TypeDocumentCreator(models.Model):
    type_document = models.ForeignKey(TypeDocuments, related_name="creators", on_delete=models.CASCADE)
    doctor = models.ForeignKey(DoctorProfile, related_name="type_document_creators", on_delete=models.CASCADE)

    class Meta:
        verbose_name = "Создатель вида документа"
        verbose_name_plural = "Создатели видов документов"
        unique_together = ("type_document", "doctor")

    def __str__(self):
        return f"{self.type_document} {self.doctor}"


class TypeCases(models.Model):
    PICKER_PK_SHIFT = 2000000000

    title = models.CharField(max_length=128, blank=True, null=True)
    code = models.CharField(max_length=55, blank=True, null=True)
    default_type_document = models.ForeignKey(
        TypeDocuments,
        default=None,
        blank=True,
        null=True,
        db_index=True,
        help_text="Документ по умолчанию при создании дела",
        on_delete=models.SET_NULL,
    )
    place_section = models.ForeignKey(
        PlaceSection,
        related_name="type_cases",
        default=None,
        blank=True,
        null=True,
        db_index=True,
        help_text="Подраздел/место в пикере ДОУ",
        on_delete=models.SET_NULL,
    )

    class Meta:
        verbose_name = "Вид дела"
        verbose_name_plural = "Виды дел"

    def __str__(self):
        return f"{self.title}"

    @property
    def json(self):
        type_doc = self.default_type_document
        place = self.place_section if self.place_section_id else None
        return {
            "id": self.id,
            "title": self.title or "",
            "code": self.code or "",
            "defaultTypeDocumentId": self.default_type_document_id if self.default_type_document_id else -1,
            "defaultTypeDocumentTitle": type_doc.title if type_doc else "Без документа",
            "placeSectionId": self.place_section_id,
            "placeSectionTitle": place.title if place else "",
        }

    @staticmethod
    def get_list():
        qs = TypeCases.objects.select_related("default_type_document", "place_section").order_by("title", "pk")
        return [row.json for row in qs]

    def as_picker_research(self):
        from podrazdeleniya.models import Podrazdeleniya

        if not self.place_section_id:
            return None
        title = self.title or ""
        code = self.code or ""
        site = self.place_section_id
        return {
            "pk": TypeCases.PICKER_PK_SHIFT + self.pk,
            "onlywith": -1,
            "department_pk": 2 - Podrazdeleniya.DOU,
            "title": title,
            "full_title": title,
            "doc_refferal": False,
            "treatment": False,
            "is_hospital": False,
            "is_form": False,
            "is_case": False,
            "is_application": False,
            "stom": False,
            "need_vich_code": False,
            "comment_variants": [],
            "autoadd": [],
            "addto": [],
            "code": code,
            "internal_code": code,
            "type": str(Podrazdeleniya.DOU),
            "site_type": site,
            "site_type_raw": site,
            "localizations": [],
            "service_locations": [],
            "direction_params": -1,
            "is_dou_document_type": True,
            "is_dou_case_type": True,
            "defaultTypeDocumentId": self.default_type_document_id if self.default_type_document_id else -1,
        }

    @staticmethod
    def picker_researches():
        qs = TypeCases.objects.filter(place_section__isnull=False).order_by("title", "pk")
        result = []
        for row in qs:
            item = row.as_picker_research()
            if item:
                result.append(item)
        return result

    @staticmethod
    def save_case(pk, title, code="", default_type_document_id=None, place_section_id=None):
        title = (title or "").strip()
        if not title:
            return {"ok": False, "message": "Укажите название"}
        type_doc = None
        if default_type_document_id not in (None, "", -1, "-1"):
            type_doc = TypeDocuments.objects.filter(pk=default_type_document_id).first()
            if not type_doc:
                return {"ok": False, "message": "Вид документа не найден"}
        if place_section_id in (None, "", -1, "-1"):
            return {"ok": False, "message": "Укажите подраздел/место"}
        place = PlaceSection.objects.filter(pk=place_section_id).first()
        if not place:
            return {"ok": False, "message": "Подраздел/место не найден"}
        if pk in (None, -1, "-1"):
            obj = TypeCases(
                title=title,
                code=code or "",
                default_type_document=type_doc,
                place_section=place,
            )
        else:
            obj = TypeCases.objects.filter(pk=pk).first()
            if not obj:
                return {"ok": False, "message": "Вид дела не найден"}
            obj.title = title
            obj.code = code or ""
            obj.default_type_document = type_doc
            obj.place_section = place
        obj.save()
        return {"ok": True, "id": obj.pk, "title": obj.title, "placeSectionId": obj.place_section_id}


class DocumentCase(models.Model):
    type_case = models.ForeignKey(
        TypeCases,
        related_name="cases",
        blank=True,
        null=True,
        db_index=True,
        help_text="Вид дела",
        on_delete=models.SET_NULL,
    )
    topic = models.CharField(max_length=512, blank=True, default="", help_text="Тема дела")
    who_create = models.ForeignKey(
        DoctorProfile,
        related_name="document_cases",
        db_index=True,
        blank=True,
        null=True,
        help_text="Создатель дела",
        on_delete=models.SET_NULL,
    )
    created_at = models.DateTimeField(auto_now_add=True, db_index=True, help_text="Дата создания дела")
    comment = models.TextField(blank=True, default="", help_text="Комментарий")
    closed_at = models.DateTimeField(blank=True, null=True, db_index=True, help_text="Дата закрытия дела")
    who_close = models.ForeignKey(
        DoctorProfile,
        related_name="closed_document_cases",
        db_index=True,
        blank=True,
        null=True,
        help_text="Кто закрыл дело",
        on_delete=models.SET_NULL,
    )

    class Meta:
        verbose_name = "Дело"
        verbose_name_plural = "Дела"

    def __str__(self):
        return self.topic or f"Дело {self.pk}"

    def as_json(self, who=None):
        data = self._base_json()
        data["access"] = self.access_json()
        data["canEditAccess"] = self.can_edit_access(who)
        data["canEdit"] = data["canEditAccess"]
        data["canEditTopic"] = data["canEditAccess"] and not self.closed_at
        data["canClose"] = data["canEditAccess"] and not self.closed_at
        data["isFavorite"] = bool(
            who and UserFavoriteCase.objects.filter(doctor=who).filter(models.Q(document_case=self) | models.Q(document_case__isnull=True, document__document_case=self)).exists()
        )
        data["documents"] = self.documents_rows(who)
        return data

    def _base_json(self):
        creator = self.who_create.get_full_fio() if self.who_create_id and self.who_create else ""
        closed_by = self.who_close.get_full_fio() if self.who_close_id and self.who_close else ""
        return {
            "id": self.pk,
            "topic": self.topic or "",
            "createdAt": self._date_text(self.created_at),
            "creator": creator,
            "comment": self.comment or "",
            "closedAt": self._date_text(self.closed_at),
            "closedBy": closed_by,
        }

    @staticmethod
    def _date_text(moment):
        if not moment:
            return ""
        from django.utils import timezone

        local = timezone.localtime(moment) if timezone.is_aware(moment) else moment
        return local.strftime("%d.%m.%Y")

    @staticmethod
    def _is_superuser(who):
        user = getattr(who, "user", None)
        return bool(user and getattr(user, "is_superuser", False))

    def can_edit_access(self, who):
        if not who:
            return False
        if DocumentCase._is_superuser(who):
            return True
        return self.who_create_id == who.pk

    def can_access(self, who):
        if self.can_edit_access(who):
            return True
        if not who:
            return False
        return self.access_rows.filter(doctor_id=who.pk).exists()

    def access_json(self):
        prefetched = getattr(self, "_prefetched_objects_cache", {}).get("access_rows")
        rows = prefetched if prefetched is not None else self.access_rows.select_related("doctor", "doctor__podrazdeleniye")
        return [AddresseeGroup._employee_json(row.doctor) for row in rows if row.doctor_id]

    def documents_rows(self, who):
        from django.utils import timezone

        from directions.models import Issledovaniya

        docs = [
            doc
            for doc in Documents.objects.filter(models.Q(case_links__document_case=self) | models.Q(document_case=self), is_hidden=False)
            .select_related("type_document", "who_create", "document_case", "block")
            .prefetch_related("case_links__document_case", "block_links__block")
            .distinct()
            .order_by("pk")
            if Documents.can_see_document(doc, who)
        ]
        if not docs:
            return []
        iss_by_doc = {}
        for iss in Issledovaniya.objects.filter(document_id__in=[doc.pk for doc in docs]).order_by("pk"):
            iss_by_doc.setdefault(iss.document_id, iss)
        topic_by_doc = Documents.topics_for_documents(docs, iss_by_doc)
        rows = []
        for doc in docs:
            created = doc.create_at
            date_text = ""
            if created:
                local = timezone.localtime(created) if timezone.is_aware(created) else created
                date_text = local.strftime("%d.%m.%Y")
            rows.append(
                {
                    "id": doc.pk,
                    "topic": (topic_by_doc.get(doc.pk) or "").strip(),
                    "typeTitle": doc.type_document.title if doc.type_document else "",
                    "createdAt": date_text,
                    "creator": doc.who_create.get_full_fio() if doc.who_create_id and doc.who_create else "",
                }
            )
        return rows

    def set_comment(self, who, comment):
        if not self.can_edit_access(who):
            return {"ok": False, "message": "Нет прав"}
        self.comment = str(comment or "").strip()
        self.save(update_fields=["comment"])
        return {"ok": True, "comment": self.comment}

    def set_topic(self, who, topic):
        if not self.can_edit_access(who):
            return {"ok": False, "message": "Нет прав"}
        if self.closed_at:
            return {"ok": False, "message": "Дело закрыто"}
        topic_text = str(topic or "").strip()
        if not topic_text:
            return {"ok": False, "message": "Укажите тему дела"}
        self.topic = topic_text
        self.save(update_fields=["topic"])
        return {"ok": True, "topic": self.topic}

    def close_case(self, who):
        from django.utils import timezone

        if not self.can_edit_access(who):
            return {"ok": False, "message": "Нет прав"}
        if self.closed_at:
            return {"ok": False, "message": "Дело уже закрыто"}
        self.closed_at = timezone.now()
        self.who_close = who
        self.save(update_fields=["closed_at", "who_close"])
        return {"ok": True, "closedAt": self._date_text(self.closed_at), "closedBy": who.get_full_fio() if who else ""}

    @classmethod
    def available_options(cls, who):
        qs = cls.objects.filter(closed_at__isnull=True)
        if not cls._is_superuser(who):
            if not who:
                qs = cls.objects.none()
            else:
                qs = qs.filter(models.Q(who_create_id=who.pk) | models.Q(access_rows__doctor_id=who.pk)).distinct()
        return [{"id": row.pk, "label": f"{row.pk} {row.topic}".strip() or f"Дело {row.pk}"} for row in qs.order_by("-pk")]

    @staticmethod
    def assign_document(document, who, case_ids):
        if not Documents.can_see_document(document, who):
            return {"ok": False, "message": "Документ не найден"}
        posted = _link_id_list(case_ids)
        if posted is None:
            return {"ok": False, "message": "Дело недоступно"}
        Documents.ensure_case_link(document)
        existing = {link.document_case_id: link.document_case for link in document.case_links.select_related("document_case")}
        final_ids = list(existing.keys())
        seen = set(final_ids)
        for pk in posted:
            if pk in seen:
                continue
            case = DocumentCase.objects.filter(pk=pk, closed_at__isnull=True).first()
            if not case or not case.can_access(who):
                return {"ok": False, "message": "Дело недоступно"}
            seen.add(pk)
            final_ids.append(pk)
        with transaction.atomic():
            have = set(document.case_links.values_list("document_case_id", flat=True))
            DocumentCaseLink.objects.bulk_create([DocumentCaseLink(document=document, document_case_id=pk) for pk in final_ids if pk not in have])
            Documents.sync_case_field(document)
        cases = Documents.linked_cases(document)
        case_payload = cases[0].as_json(who) if len(cases) == 1 else None
        return {"ok": True, "caseIds": [case.pk for case in cases], "caseId": document.document_case_id, "case": case_payload}

    def set_access(self, who, members):
        if not self.can_edit_access(who):
            return {"ok": False, "message": "Нет прав"}
        if not isinstance(members, list):
            return {"ok": False, "message": "Некорректный список доступа"}
        ids = []
        seen = set()
        for item in members:
            raw = item.get("id") if isinstance(item, dict) else item
            try:
                doctor_id = int(raw)
            except (TypeError, ValueError):
                continue
            if doctor_id > 0 and doctor_id not in seen:
                seen.add(doctor_id)
                ids.append(doctor_id)
        valid_ids = list(DoctorProfile.objects.filter(pk__in=ids).values_list("pk", flat=True))
        valid_set = set(valid_ids)
        with transaction.atomic():
            self.access_rows.exclude(doctor_id__in=valid_set).delete()
            existing = set(self.access_rows.filter(doctor_id__in=valid_set).values_list("doctor_id", flat=True))
            DocumentCaseAccess.objects.bulk_create([DocumentCaseAccess(document_case=self, doctor_id=doctor_id) for doctor_id in valid_ids if doctor_id not in existing])
        cache = getattr(self, "_prefetched_objects_cache", None)
        if cache is not None:
            cache.pop("access_rows", None)
        return {"ok": True, "access": self.access_json()}


class UserMyCase(models.Model):
    doctor = models.ForeignKey(DoctorProfile, related_name="added_my_cases", on_delete=models.CASCADE)
    document_case = models.ForeignKey(DocumentCase, related_name="user_my_cases", on_delete=models.CASCADE)

    class Meta:
        verbose_name = "Добавленное моё дело"
        verbose_name_plural = "Добавленные мои дела"
        unique_together = ("doctor", "document_case")

    def __str__(self):
        return f"{self.doctor_id} {self.document_case_id}"

    @staticmethod
    def toggle(who, case_id):
        if not who:
            return {"ok": False, "message": "Пользователь не найден"}
        try:
            case_pk = int(case_id)
        except (TypeError, ValueError):
            return {"ok": False, "message": "Дело не найдено"}
        case = DocumentCase.objects.filter(pk=case_pk).first()
        if not case or not case.can_access(who):
            return {"ok": False, "message": "Дело не найдено"}
        if case.who_create_id == who.pk:
            return {"ok": True, "inMyCases": True, "isMine": True}
        row = UserMyCase.objects.filter(doctor=who, document_case=case).first()
        if row:
            row.delete()
            return {"ok": True, "inMyCases": False, "isMine": False}
        UserMyCase.objects.create(doctor=who, document_case=case)
        return {"ok": True, "inMyCases": True, "isMine": False}


class DocumentCaseAccess(models.Model):
    document_case = models.ForeignKey(DocumentCase, related_name="access_rows", on_delete=models.CASCADE)
    doctor = models.ForeignKey(DoctorProfile, related_name="document_case_access", on_delete=models.CASCADE)

    class Meta:
        verbose_name = "Доступ к делу"
        verbose_name_plural = "Доступ к делам"
        unique_together = ("document_case", "doctor")

    def __str__(self):
        return f"{self.document_case_id} {self.doctor_id}"


class DocumentBlock(models.Model):
    CREATE_GROUP = "Создание блоков"

    title = models.CharField(max_length=512, help_text="Название блока")
    who_create = models.ForeignKey(
        DoctorProfile,
        related_name="document_blocks",
        db_index=True,
        blank=True,
        null=True,
        help_text="Создатель блока",
        on_delete=models.SET_NULL,
    )
    created_at = models.DateTimeField(auto_now_add=True, db_index=True, help_text="Дата создания блока")
    is_published = models.BooleanField(default=False, db_index=True, help_text="Опубликован")

    class Meta:
        verbose_name = "Блок"
        verbose_name_plural = "Блоки"

    def __str__(self):
        return self.title or f"Блок {self.pk}"

    @staticmethod
    def _is_superuser(who):
        return DocumentCase._is_superuser(who)

    @classmethod
    def can_create(cls, who):
        if not who:
            return False
        if cls._is_superuser(who):
            return True
        return who.has_group(cls.CREATE_GROUP)

    def can_edit(self, who):
        if not who:
            return False
        if self._is_superuser(who):
            return True
        return self.who_create_id == who.pk

    def can_write(self, who):
        if self.can_edit(who):
            return True
        if not who:
            return False
        return self.write_rows.filter(doctor_id=who.pk).exists()

    def is_public(self):
        return bool(self.is_published) or not self.read_rows.exists()

    def can_read(self, who):
        if not who:
            return False
        if self.can_write(who) or self.is_public():
            return True
        return self.read_rows.filter(doctor_id=who.pk).exists()

    def as_json(self, who=None):
        return {
            "ok": True,
            "id": self.pk,
            "title": self.title or "",
            "createdAt": DocumentCase._date_text(self.created_at),
            "creator": self.who_create.get_full_fio() if self.who_create_id and self.who_create else "",
            "isPublished": bool(self.is_published),
            "canEdit": self.can_edit(who),
            "canWrite": self.can_write(who),
            "readAccess": self._access_json("read_rows"),
            "writeAccess": self._access_json("write_rows"),
            "documents": self.documents_rows(who),
        }

    def list_json(self):
        creator = self.who_create.get_full_fio() if self.who_create_id and self.who_create else ""
        return {
            "id": self.pk,
            "title": self.title or "",
            "creator": creator,
            "createdAt": DocumentCase._date_text(self.created_at),
        }

    def _access_json(self, relation_name):
        prefetched = getattr(self, "_prefetched_objects_cache", {}).get(relation_name)
        rows = prefetched if prefetched is not None else getattr(self, relation_name).select_related("doctor", "doctor__podrazdeleniye")
        return [AddresseeGroup._employee_json(row.doctor) for row in rows if row.doctor_id]

    def documents_rows(self, who):
        from django.utils import timezone

        from directions.models import Issledovaniya

        linked = (
            Documents.objects.filter(models.Q(block_links__block=self) | models.Q(block=self))
            .select_related("type_document", "who_create", "document_case", "block")
            .prefetch_related("case_links__document_case", "block_links__block")
            .distinct()
            .order_by("pk")
        )
        docs = [doc for doc in linked if Documents.can_see_document(doc, who)]
        if not docs:
            return []
        iss_by_doc = {}
        for iss in Issledovaniya.objects.filter(document_id__in=[doc.pk for doc in docs]).order_by("pk"):
            iss_by_doc.setdefault(iss.document_id, iss)
        topic_by_doc = Documents.topics_for_documents(docs, iss_by_doc)
        rows = []
        for doc in docs:
            created = doc.create_at
            date_text = ""
            if created:
                local = timezone.localtime(created) if timezone.is_aware(created) else created
                date_text = local.strftime("%d.%m.%Y")
            rows.append(
                {
                    "id": doc.pk,
                    "topic": (topic_by_doc.get(doc.pk) or "").strip(),
                    "typeTitle": doc.type_document.title if doc.type_document else "",
                    "createdAt": date_text,
                    "creator": doc.who_create.get_full_fio() if doc.who_create_id and doc.who_create else "",
                }
            )
        return rows

    @classmethod
    def visible_for(cls, who):
        from django.db.models import Exists, OuterRef

        if not who:
            return cls.objects.none()
        if cls._is_superuser(who):
            return cls.objects.all()
        has_readers = DocumentBlockReadAccess.objects.filter(block_id=OuterRef("pk"))
        return (
            cls.objects.annotate(_has_readers=Exists(has_readers))
            .filter(
                models.Q(is_published=True) | models.Q(_has_readers=False) | models.Q(who_create_id=who.pk) | models.Q(read_rows__doctor_id=who.pk) | models.Q(write_rows__doctor_id=who.pk)
            )
            .distinct()
        )

    @classmethod
    def list_for(cls, who):
        rows = cls.visible_for(who).select_related("who_create").order_by("-pk")
        return [row.list_json() for row in rows]

    @classmethod
    def write_options(cls, who):
        if not who:
            return []
        if cls._is_superuser(who):
            qs = cls.objects.all()
        else:
            qs = cls.objects.filter(models.Q(who_create_id=who.pk) | models.Q(write_rows__doctor_id=who.pk)).distinct()
        return [{"id": row.pk, "label": row.title or f"Блок {row.pk}"} for row in qs.order_by("-pk")]

    @classmethod
    def create_block(cls, who, title):
        if not cls.can_create(who):
            return {"ok": False, "message": "Нет прав"}
        title_text = str(title or "").strip()
        if not title_text:
            return {"ok": False, "message": "Укажите название блока"}
        row = cls.objects.create(title=title_text, who_create=who)
        return {"ok": True, "id": row.pk}

    def save_settings(self, who, title, is_published):
        if not self.can_edit(who):
            return {"ok": False, "message": "Нет прав"}
        title_text = str(title or "").strip()
        if not title_text:
            return {"ok": False, "message": "Укажите название блока"}
        self.title = title_text
        self.is_published = bool(is_published)
        self.save(update_fields=["title", "is_published"])
        return {"ok": True, "title": self.title, "isPublished": self.is_published}

    def set_read_access(self, who, members):
        return self._set_access(who, members, "read_rows", DocumentBlockReadAccess, "block")

    def set_write_access(self, who, members):
        return self._set_access(who, members, "write_rows", DocumentBlockWriteAccess, "block")

    def _set_access(self, who, members, relation_name, model_cls, fk_name):
        if not self.can_edit(who):
            return {"ok": False, "message": "Нет прав"}
        if not isinstance(members, list):
            return {"ok": False, "message": "Некорректный список доступа"}
        ids = []
        seen = set()
        for item in members:
            raw = item.get("id") if isinstance(item, dict) else item
            try:
                doctor_id = int(raw)
            except (TypeError, ValueError):
                continue
            if doctor_id > 0 and doctor_id not in seen:
                seen.add(doctor_id)
                ids.append(doctor_id)
        valid_ids = list(DoctorProfile.objects.filter(pk__in=ids).values_list("pk", flat=True))
        valid_set = set(valid_ids)
        relation = getattr(self, relation_name)
        with transaction.atomic():
            relation.exclude(doctor_id__in=valid_set).delete()
            existing = set(relation.filter(doctor_id__in=valid_set).values_list("doctor_id", flat=True))
            model_cls.objects.bulk_create([model_cls(**{fk_name: self, "doctor_id": doctor_id}) for doctor_id in valid_ids if doctor_id not in existing])
        cache = getattr(self, "_prefetched_objects_cache", None)
        if cache is not None:
            cache.pop(relation_name, None)
        key = "readAccess" if relation_name == "read_rows" else "writeAccess"
        return {"ok": True, key: self._access_json(relation_name)}

    @staticmethod
    def assign_document(document, who, block_ids):
        if not Documents.can_see_document(document, who):
            return {"ok": False, "message": "Документ не найден"}
        posted = _link_id_list(block_ids)
        if posted is None:
            return {"ok": False, "message": "Блок недоступен"}
        Documents.ensure_block_link(document)
        existing = {link.block_id: link.block for link in document.block_links.select_related("block")}
        locked = [pk for pk, block in existing.items() if block and not block.can_write(who)]
        posted_set = set(posted)
        if any(pk not in posted_set for pk in locked):
            return {"ok": False, "message": "Нет прав"}
        desired = []
        seen = set()
        for pk in posted:
            if pk in seen:
                continue
            seen.add(pk)
            if pk in locked:
                desired.append(pk)
                continue
            block = existing.get(pk) or DocumentBlock.objects.filter(pk=pk).first()
            if not block:
                return {"ok": False, "message": "Блок недоступен"}
            if not block.can_write(who):
                return {"ok": False, "message": "Нет прав"}
            desired.append(pk)
        with transaction.atomic():
            document.block_links.exclude(block_id__in=desired).delete()
            have = set(document.block_links.values_list("block_id", flat=True))
            DocumentBlockLink.objects.bulk_create([DocumentBlockLink(document=document, block_id=pk) for pk in desired if pk not in have])
            Documents.sync_block_field(document)
        blocks = Documents.linked_blocks(document)
        return {"ok": True, "blockIds": [block.pk for block in blocks], "blockId": document.block_id}


class DocumentBlockWriteAccess(models.Model):
    block = models.ForeignKey(DocumentBlock, related_name="write_rows", on_delete=models.CASCADE)
    doctor = models.ForeignKey(DoctorProfile, related_name="document_block_write_access", on_delete=models.CASCADE)

    class Meta:
        verbose_name = "Запись в блок"
        verbose_name_plural = "Запись в блоки"
        unique_together = ("block", "doctor")

    def __str__(self):
        return f"{self.block_id} {self.doctor_id}"


class DocumentBlockReadAccess(models.Model):
    block = models.ForeignKey(DocumentBlock, related_name="read_rows", on_delete=models.CASCADE)
    doctor = models.ForeignKey(DoctorProfile, related_name="document_block_read_access", on_delete=models.CASCADE)

    class Meta:
        verbose_name = "Чтение блока"
        verbose_name_plural = "Чтение блоков"
        unique_together = ("block", "doctor")

    def __str__(self):
        return f"{self.block_id} {self.doctor_id}"


class DocumentPickerFavorite(models.Model):
    doctor = models.ForeignKey(DoctorProfile, related_name="document_picker_favorites", on_delete=models.CASCADE)
    type_document = models.ForeignKey(
        TypeDocuments,
        related_name="picker_favorites",
        blank=True,
        null=True,
        on_delete=models.CASCADE,
    )
    type_case = models.ForeignKey(
        TypeCases,
        related_name="picker_favorites",
        blank=True,
        null=True,
        on_delete=models.CASCADE,
    )

    class Meta:
        verbose_name = "Избранный вид"
        verbose_name_plural = "Избранные виды"
        constraints = [
            models.CheckConstraint(
                check=(models.Q(type_document__isnull=False, type_case__isnull=True) | models.Q(type_document__isnull=True, type_case__isnull=False)),
                name="document_picker_favorite_one_target",
            ),
            models.UniqueConstraint(
                fields=("doctor", "type_document"),
                condition=models.Q(type_document__isnull=False),
                name="document_picker_favorite_doctor_type",
            ),
            models.UniqueConstraint(
                fields=("doctor", "type_case"),
                condition=models.Q(type_case__isnull=False),
                name="document_picker_favorite_doctor_case",
            ),
        ]

    def __str__(self):
        return f"{self.doctor_id} {self.type_document_id or self.type_case_id}"

    @property
    def picker_pk(self):
        if self.type_document_id:
            return TypeDocuments.PICKER_PK_SHIFT + self.type_document_id
        if self.type_case_id:
            return TypeCases.PICKER_PK_SHIFT + self.type_case_id
        return None

    @staticmethod
    def picker_pks(doctor):
        if not doctor:
            return []
        rows = DocumentPickerFavorite.objects.filter(doctor=doctor).order_by("pk")
        return [row.picker_pk for row in rows if row.picker_pk]

    @staticmethod
    def _target(pk):
        try:
            value = int(pk)
        except (TypeError, ValueError):
            return None, None
        if value >= TypeCases.PICKER_PK_SHIFT:
            case = TypeCases.objects.filter(pk=value - TypeCases.PICKER_PK_SHIFT).first()
            return ("case", case) if case else (None, None)
        if value >= TypeDocuments.PICKER_PK_SHIFT:
            doc = TypeDocuments.objects.filter(pk=value - TypeDocuments.PICKER_PK_SHIFT).first()
            return ("document", doc) if doc else (None, None)
        return None, None

    @staticmethod
    def toggle(doctor, pk):
        if not doctor:
            return {"ok": False, "message": "Пользователь не найден"}
        kind, target = DocumentPickerFavorite._target(pk)
        if not target:
            return {"ok": False, "message": "Вид не найден"}
        if kind == "case":
            existing = DocumentPickerFavorite.objects.filter(doctor=doctor, type_case=target)
        else:
            existing = DocumentPickerFavorite.objects.filter(doctor=doctor, type_document=target)
        if existing.exists():
            existing.delete()
            favorite = False
        elif kind == "case":
            DocumentPickerFavorite.objects.create(doctor=doctor, type_case=target)
            favorite = True
        else:
            DocumentPickerFavorite.objects.create(doctor=doctor, type_document=target)
            favorite = True
        return {"ok": True, "favorite": favorite, "pks": DocumentPickerFavorite.picker_pks(doctor)}


class UserFavoriteDocument(models.Model):
    doctor = models.ForeignKey(DoctorProfile, related_name="favorite_documents", on_delete=models.CASCADE)
    document = models.ForeignKey("Documents", related_name="user_favorites", on_delete=models.CASCADE)

    class Meta:
        verbose_name = "Избранный документ"
        verbose_name_plural = "Избранные документы"
        unique_together = ("doctor", "document")

    def __str__(self):
        return f"{self.doctor_id} {self.document_id}"


class UserFavoriteCase(models.Model):
    doctor = models.ForeignKey(DoctorProfile, related_name="favorite_cases", on_delete=models.CASCADE)
    document = models.ForeignKey("Documents", related_name="user_case_favorites", on_delete=models.CASCADE)
    document_case = models.ForeignKey(DocumentCase, related_name="favorite_rows", blank=True, null=True, on_delete=models.CASCADE)

    class Meta:
        verbose_name = "Избранное дело"
        verbose_name_plural = "Избранные дела"
        unique_together = ("doctor", "document", "document_case")

    def __str__(self):
        return f"{self.doctor_id} {self.document_id}"


def _record_favorite_model(document):
    return UserFavoriteCase if document.type_case_id else UserFavoriteDocument


def record_favorite_status(doctor, document):
    if not doctor or not document:
        return False
    return _record_favorite_model(document).objects.filter(doctor=doctor, document=document).exists()


def toggle_favorite_case(doctor, case_id):
    if not doctor:
        return {"ok": False, "message": "Пользователь не найден"}
    try:
        case_pk = int(case_id)
    except (TypeError, ValueError):
        return {"ok": False, "message": "Дело не найдено"}
    case = DocumentCase.objects.filter(pk=case_pk).first()
    if not case or not case.can_access(doctor):
        return {"ok": False, "message": "Дело не найдено"}
    existing = UserFavoriteCase.objects.filter(doctor=doctor).filter(models.Q(document_case=case) | models.Q(document_case__isnull=True, document__document_case=case))
    if existing.exists():
        existing.delete()
        return {"ok": True, "favorite": False, "isCase": True}
    document = Documents.objects.filter(models.Q(case_links__document_case=case) | models.Q(document_case=case)).order_by("pk").distinct().first()
    if not document:
        return {"ok": False, "message": "В деле нет документа"}
    UserFavoriteCase.objects.create(doctor=doctor, document=document, document_case=case)
    return {"ok": True, "favorite": True, "isCase": True}


def toggle_case_favorite(doctor, document_id):
    if not doctor:
        return {"ok": False, "message": "Пользователь не найден"}
    document = Documents.objects.filter(pk=document_id).first()
    cases = Documents.linked_cases(document) if document else []
    case = cases[0] if len(cases) == 1 else None
    if not document or not case or not Documents.can_see_document(document, doctor):
        return {"ok": False, "message": "Дело не найдено"}
    existing = UserFavoriteCase.objects.filter(doctor=doctor).filter(models.Q(document_case=case) | models.Q(document_case__isnull=True, document__document_case=case))
    if existing.exists():
        existing.delete()
        favorite = False
    else:
        UserFavoriteCase.objects.create(doctor=doctor, document=document, document_case=case)
        favorite = True
    return {"ok": True, "favorite": favorite, "isCase": True}


def toggle_record_favorite(doctor, document_id):
    if not doctor:
        return {"ok": False, "message": "Пользователь не найден"}
    document = Documents.objects.select_related("document_case").filter(pk=document_id).first()
    if not document or not Documents.can_see_document(document, doctor):
        return {"ok": False, "message": "Документ не найден"}
    model = _record_favorite_model(document)
    existing = model.objects.filter(doctor=doctor, document=document)
    if existing.exists():
        existing.delete()
        favorite = False
    else:
        model.objects.create(doctor=doctor, document=document)
        favorite = True
    return {"ok": True, "favorite": favorite, "isCase": bool(document.type_case_id)}


def list_record_favorites(doctor, cases=False):
    if not doctor:
        return []
    model = UserFavoriteCase if cases else UserFavoriteDocument
    rows = (
        model.objects.filter(doctor=doctor)
        .filter(Documents.case_access_q(doctor, "document__"))
        .select_related("document", "document__type_document", "document__document_case")
        .order_by("-pk")
        .distinct()
    )
    documents = [row.document for row in rows if row.document_id]
    if not documents:
        return []
    from directions.models import Issledovaniya

    iss_by_doc = {}
    for iss in Issledovaniya.objects.filter(document_id__in=[doc.pk for doc in documents]).order_by("pk"):
        iss_by_doc.setdefault(iss.document_id, iss)
    topic_by_doc = Documents.topics_for_documents(documents, iss_by_doc)
    result = []
    for document in documents:
        topic = (topic_by_doc.get(document.pk) or "").strip()
        type_doc = document.type_document.title if document.type_document else ""
        title = " ".join(part for part in (topic, type_doc) if part)
        result.append({"id": document.pk, "title": Documents.title_with_id(document.pk, title)})
    return result


class Plans(models.Model):
    title = models.CharField(max_length=128, blank=True, null=True)
    user_created = models.BooleanField(default=False, db_index=True, help_text="Создан пользователем")
    who_create = models.ForeignKey(
        DoctorProfile,
        related_name="created_plans",
        null=True,
        blank=True,
        default=None,
        db_index=True,
        help_text="Автор плана",
        on_delete=models.SET_NULL,
    )

    class Meta:
        verbose_name = "План"
        verbose_name_plural = "Планы"

    def __str__(self):
        return f"{self.title}"

    @property
    def json(self):
        return {"id": self.id, "title": self.title or ""}

    @staticmethod
    def get_list():
        return [row.json for row in Plans.objects.filter(user_created=False).order_by("title", "pk")]

    @staticmethod
    def parse_ids(value):
        if value in (None, ""):
            return []
        raw = value
        if isinstance(value, str):
            try:
                raw = json.loads(value)
            except (TypeError, ValueError):
                return None
        if not isinstance(raw, list):
            return None
        ids = []
        for item in raw:
            if isinstance(item, dict):
                item = item.get("id")
            try:
                pk = int(item)
            except (TypeError, ValueError):
                return None
            if pk > 0 and pk not in ids:
                ids.append(pk)
        return ids

    @staticmethod
    def choice_lists(who):
        global_rows = Plans.objects.filter(user_created=False).order_by("title", "pk")
        mine_rows = Plans.objects.none()
        if who and getattr(who, "pk", None):
            mine_rows = Plans.objects.filter(user_created=True, who_create=who).order_by("title", "pk")
        return {"global": [row.json for row in global_rows], "mine": [row.json for row in mine_rows]}

    @staticmethod
    def allowed_ids(who):
        ids = set(Plans.objects.filter(user_created=False).values_list("pk", flat=True))
        if who and getattr(who, "pk", None):
            ids.update(Plans.objects.filter(user_created=True, who_create=who).values_list("pk", flat=True))
        return ids

    @staticmethod
    def ids_on_document(document):
        from directions.models import ParaclinicResult

        found = set()
        rows = ParaclinicResult.objects.filter(issledovaniye__document=document, field_type=PLAN_FIELD_TYPE).only("value")
        for row in rows:
            parsed = Plans.parse_ids(row.value)
            if parsed:
                found.update(parsed)
        return found

    @staticmethod
    def preview_one(plan):
        def indicator_row(row):
            data = row.as_json()
            data["title"] = row.indicator.title if row.indicator_id and getattr(row, "indicator", None) else ""
            return data

        def test_row(row):
            data = row.as_json()
            data["title"] = row.test.title if row.test_id and getattr(row, "test", None) else ""
            return data

        groups = []
        for group in plan.indicator_groups.all():
            groups.append(
                {
                    "id": group.pk,
                    "title": group.title or "",
                    "order": group.order,
                    "indicators": [indicator_row(row) for row in group.indicators.all()],
                    "tests": [test_row(row) for row in group.tests.all()],
                }
            )
        loose_indicators = [indicator_row(row) for row in plan.indicators.all() if row.group_id is None]
        loose_tests = [test_row(row) for row in plan.tests.all() if row.group_id is None]
        return {
            "id": plan.pk,
            "title": plan.title or "",
            "groups": groups,
            "indicators": loose_indicators,
            "tests": loose_tests,
        }

    @staticmethod
    def preview_for(ids, who, document=None):
        from django.db.models import Prefetch

        allowed = Plans.allowed_ids(who)
        if document is not None:
            allowed |= Plans.ids_on_document(document)
        wanted = [pk for pk in ids if pk in allowed]
        if not wanted:
            return []
        plans = Plans.objects.filter(pk__in=wanted).prefetch_related(
            Prefetch(
                "indicator_groups",
                queryset=PlanIndicatorGroup.objects.order_by("order", "pk").prefetch_related(
                    Prefetch("indicators", queryset=PlanIndicator.objects.select_related("indicator").order_by("order", "pk")),
                    Prefetch("tests", queryset=PlanTest.objects.select_related("test").order_by("order", "pk")),
                ),
            ),
            Prefetch("indicators", queryset=PlanIndicator.objects.select_related("indicator").order_by("order", "pk")),
            Prefetch("tests", queryset=PlanTest.objects.select_related("test").order_by("order", "pk")),
        )
        by_id = {plan.pk: plan for plan in plans}
        return [Plans.preview_one(by_id[pk]) for pk in wanted if pk in by_id]

    @staticmethod
    def validate_document_fields(groups, iss, who):
        from directions.models import ParaclinicResult

        allowed = Plans.allowed_ids(who)
        for group in groups or []:
            for field in group.get("fields") or []:
                if field.get("field_type") != PLAN_FIELD_TYPE:
                    continue
                ids = Plans.parse_ids(field.get("value"))
                if ids is None:
                    return "Некорректный список планов"
                previous = []
                existing = ParaclinicResult.objects.filter(issledovaniye=iss, field_id=field.get("pk")).only("value").first()
                if existing:
                    previous = Plans.parse_ids(existing.value) or []
                if any(pk not in allowed and pk not in previous for pk in ids):
                    return "План недоступен"
                field["value"] = json.dumps(ids)
        return None

    @staticmethod
    def indicator_options():
        options = []
        for group in DouIndicator.catalog()["globalGroups"]:
            options.extend(group["indicators"])
        return options

    @staticmethod
    def indicator_catalog(who=None, extra_ids=None):
        return DouIndicator.catalog(who, extra_ids)

    @staticmethod
    def create_user_indicator(who, title):
        return DouIndicator.create_for_user(who, title)

    @staticmethod
    def test_options():
        rows = DouTest.objects.all().order_by("title", "pk")
        return [{"id": row.pk, "label": row.title or ""} for row in rows]

    @staticmethod
    def get_details(pk, who=None):
        options = Plans.indicator_options()
        test_options = Plans.test_options()
        if pk in (None, -1, "-1"):
            return {
                "ok": True,
                "id": -1,
                "title": "",
                "groups": [],
                "indicators": [],
                "tests": [],
                "indicatorOptions": options,
                "indicatorCatalog": Plans.indicator_catalog(who),
                "testOptions": test_options,
            }
        plan = Plans.objects.filter(pk=pk).prefetch_related("indicator_groups__indicators", "indicator_groups__tests", "indicators", "tests").first()
        if not plan:
            return {"ok": False, "message": "План не найден"}
        groups = []
        for group in plan.indicator_groups.all().order_by("order", "pk"):
            groups.append(
                {
                    "id": group.pk,
                    "title": group.title or "",
                    "order": group.order,
                    "indicators": [row.as_json() for row in group.indicators.all().order_by("order", "pk")],
                    "tests": [row.as_json() for row in group.tests.all().order_by("order", "pk")],
                }
            )
        loose = [row.as_json() for row in plan.indicators.filter(group__isnull=True).order_by("order", "pk")]
        loose_tests = [row.as_json() for row in plan.tests.filter(group__isnull=True).order_by("order", "pk")]
        return {
            "ok": True,
            "id": plan.pk,
            "title": plan.title or "",
            "groups": groups,
            "indicators": loose,
            "tests": loose_tests,
            "indicatorOptions": options,
            "indicatorCatalog": Plans.indicator_catalog(who, plan.indicators.values_list("indicator_id", flat=True)),
            "testOptions": test_options,
            "userCreated": bool(plan.user_created),
        }

    @staticmethod
    def save_plan(pk, title, groups, indicators, tests=None, allow_user=False, who=None):
        from django.db import transaction

        title = (title or "").strip()
        if not title:
            return {"ok": False, "message": "Укажите название"}
        groups = groups or []
        indicators = indicators or []
        tests = tests or []
        kept_ids = set()
        if pk not in (None, -1, "-1"):
            kept_ids = set(PlanIndicator.objects.filter(plan_id=pk).values_list("indicator_id", flat=True))
        cleaned_groups = []
        for index, group in enumerate(groups):
            group_title = (group.get("title") or "").strip()
            if not group_title:
                return {"ok": False, "message": "Укажите название группы"}
            cleaned_rows = []
            for row_index, row in enumerate(group.get("indicators") or []):
                cleaned, error = PlanIndicator.clean_row(row, row_index, who=who, allow_user=allow_user, kept_ids=kept_ids)
                if error:
                    return {"ok": False, "message": error}
                cleaned_rows.append(cleaned)
            cleaned_tests = []
            for row_index, row in enumerate(group.get("tests") or []):
                cleaned, error = PlanTest.clean_row(row, row_index)
                if error:
                    return {"ok": False, "message": error}
                cleaned_tests.append(cleaned)
            cleaned_groups.append({"title": group_title, "order": index, "indicators": cleaned_rows, "tests": cleaned_tests})
        cleaned_loose = []
        for row_index, row in enumerate(indicators):
            cleaned, error = PlanIndicator.clean_row(row, row_index, who=who, allow_user=allow_user, kept_ids=kept_ids)
            if error:
                return {"ok": False, "message": error}
            cleaned_loose.append(cleaned)
        cleaned_loose_tests = []
        for row_index, row in enumerate(tests):
            cleaned, error = PlanTest.clean_row(row, row_index)
            if error:
                return {"ok": False, "message": error}
            cleaned_loose_tests.append(cleaned)
        with transaction.atomic():
            if pk in (None, -1, "-1"):
                if allow_user:
                    return {"ok": False, "message": "План не найден"}
                plan = Plans(title=title, user_created=False)
            else:
                plan = Plans.objects.filter(pk=pk).first()
                if not plan or bool(plan.user_created) != bool(allow_user):
                    return {"ok": False, "message": "План не найден"}
                plan.title = title
            plan.save()
            PlanTest.objects.filter(plan=plan).delete()
            PlanIndicator.objects.filter(plan=plan).delete()
            PlanIndicatorGroup.objects.filter(plan=plan).delete()
            for group_row in cleaned_groups:
                group = PlanIndicatorGroup.objects.create(plan=plan, title=group_row["title"], order=group_row["order"])
                for row in group_row["indicators"]:
                    PlanIndicator.objects.create(plan=plan, group=group, **row)
                for row in group_row["tests"]:
                    PlanTest.objects.create(plan=plan, group=group, **row)
            for row in cleaned_loose:
                PlanIndicator.objects.create(plan=plan, group=None, **row)
            for row in cleaned_loose_tests:
                PlanTest.objects.create(plan=plan, group=None, **row)
        return {"ok": True, "id": plan.pk, "title": plan.title}


class DouIndicatorGroup(models.Model):
    title = models.CharField(max_length=255, blank=True, default="")
    order = models.IntegerField(default=0)
    who_create = models.ForeignKey(
        "users.DoctorProfile",
        related_name="created_dou_indicator_groups",
        null=True,
        blank=True,
        default=None,
        db_index=True,
        on_delete=models.SET_NULL,
        help_text="Автор группы",
    )

    class Meta:
        verbose_name = "Группа показателей"
        verbose_name_plural = "Группы показателей"
        ordering = ("order", "pk")

    def __str__(self):
        return self.title or ""


class DouIndicator(models.Model):
    title = models.CharField(max_length=255, blank=True, default="")
    group = models.ForeignKey(DouIndicatorGroup, related_name="indicators", null=True, blank=True, default=None, on_delete=models.CASCADE)
    order = models.IntegerField(default=0)
    who_create = models.ForeignKey(
        "users.DoctorProfile",
        related_name="created_dou_indicators",
        null=True,
        blank=True,
        default=None,
        db_index=True,
        on_delete=models.SET_NULL,
        help_text="Автор своей задачи",
    )

    class Meta:
        verbose_name = "Показатель"
        verbose_name_plural = "Показатели"
        ordering = ("order", "pk")

    def __str__(self):
        return self.title or ""

    @staticmethod
    def _label(row):
        return {"id": row.pk, "label": row.title or ""}

    @staticmethod
    def _ids(raw_ids):
        found = []
        for raw in raw_ids or []:
            try:
                pk = int(raw)
            except (TypeError, ValueError):
                continue
            if pk > 0 and pk not in found:
                found.append(pk)
        return found

    @staticmethod
    def catalog(who=None, extra_ids=None):
        groups = list(DouIndicatorGroup.objects.filter(who_create__isnull=True).order_by("order", "title", "pk"))
        rows = list(DouIndicator.objects.filter(who_create__isnull=True).order_by("order", "title", "pk"))
        by_group = {}
        loose = []
        for row in rows:
            if row.group_id:
                by_group.setdefault(row.group_id, []).append(DouIndicator._label(row))
            else:
                loose.append(DouIndicator._label(row))
        global_groups = [{"id": 0, "title": "Без группы", "indicators": loose}]
        for group in groups:
            global_groups.append({"id": group.pk, "title": group.title or "", "indicators": by_group.get(group.pk, [])})
        known = {row.pk for row in rows}
        mine_rows = []
        if getattr(who, "pk", None):
            mine_rows = list(DouIndicator.objects.filter(who_create_id=who.pk).order_by("title", "pk"))
            known.update(row.pk for row in mine_rows)
        missing = [pk for pk in DouIndicator._ids(extra_ids) if pk not in known]
        extra_rows = list(DouIndicator.objects.filter(pk__in=missing).order_by("title", "pk")) if missing else []
        return {
            "globalGroups": global_groups,
            "mine": [DouIndicator._label(row) for row in mine_rows],
            "extra": [DouIndicator._label(row) for row in extra_rows],
        }

    @staticmethod
    def create_for_user(who, title):
        title = (title or "").strip()
        if not title:
            return {"ok": False, "message": "Укажите название задачи"}
        if len(title) > 255:
            return {"ok": False, "message": "Название слишком длинное"}
        if not getattr(who, "pk", None):
            return {"ok": False, "message": "Задача недоступна"}
        row = DouIndicator.objects.create(title=title, group=None, who_create=who, order=0)
        return {"ok": True, "id": row.pk, "label": row.title}

    @staticmethod
    def get_list():
        groups = DouIndicatorGroup.objects.filter(who_create__isnull=True).order_by("order", "title", "pk")
        return [{"id": 0, "title": "Без группы"}, *[{"id": group.pk, "title": group.title or ""} for group in groups]]

    @staticmethod
    def get_details(pk):
        if pk in (0, "0"):
            rows = DouIndicator.objects.filter(group__isnull=True, who_create__isnull=True).order_by("order", "title", "pk")
            return {"ok": True, "id": 0, "title": "Без группы", "loose": True, "indicators": [DouIndicator._editor_row(row) for row in rows]}
        if pk in (None, -1, "-1"):
            return {"ok": True, "id": -1, "title": "", "loose": False, "indicators": []}
        group = DouIndicatorGroup.objects.filter(pk=pk, who_create__isnull=True).first()
        if not group:
            return {"ok": False, "message": "Группа не найдена"}
        rows = group.indicators.filter(who_create__isnull=True).order_by("order", "title", "pk")
        return {"ok": True, "id": group.pk, "title": group.title or "", "loose": False, "indicators": [DouIndicator._editor_row(row) for row in rows]}

    @staticmethod
    def _editor_row(row):
        return {"id": row.pk, "title": row.title or ""}

    @staticmethod
    def _clean_titles(rows):
        cleaned = []
        for index, row in enumerate(rows or []):
            title = (row.get("title") or "").strip()
            if not title:
                return None, "Укажите название задачи"
            if len(title) > 255:
                return None, "Название слишком длинное"
            try:
                row_id = int(row.get("id") or 0)
            except (TypeError, ValueError):
                row_id = 0
            cleaned.append({"id": row_id if row_id > 0 else None, "title": title, "order": index})
        return cleaned, None

    @staticmethod
    def save_catalog(pk, title, rows):
        from django.db import transaction

        cleaned, error = DouIndicator._clean_titles(rows)
        if error:
            return {"ok": False, "message": error}
        loose = pk in (0, "0")
        title = (title or "").strip()
        if not loose and not title:
            return {"ok": False, "message": "Укажите название группы"}
        if not loose and len(title) > 255:
            return {"ok": False, "message": "Название слишком длинное"}

        class CatalogError(Exception):
            def __init__(self, message):
                super().__init__(message)
                self.message = message

        try:
            with transaction.atomic():
                if loose:
                    group = None
                    existing = DouIndicator.objects.filter(group__isnull=True, who_create__isnull=True)
                elif pk in (None, -1, "-1"):
                    order = DouIndicatorGroup.objects.filter(who_create__isnull=True).count()
                    group = DouIndicatorGroup.objects.create(title=title, order=order, who_create=None)
                    existing = DouIndicator.objects.none()
                else:
                    group = DouIndicatorGroup.objects.filter(pk=pk, who_create__isnull=True).first()
                    if not group:
                        raise CatalogError("Группа не найдена")
                    group.title = title
                    group.save(update_fields=["title"])
                    existing = group.indicators.filter(who_create__isnull=True)
                keep = []
                for row in cleaned:
                    obj = existing.filter(pk=row["id"]).first() if row["id"] else None
                    moved = False
                    if row["id"] and not obj and group is not None:
                        obj = DouIndicator.objects.filter(pk=row["id"], group__isnull=True, who_create__isnull=True).first()
                        moved = obj is not None
                    if row["id"] and not obj:
                        raise CatalogError("Показатель не найден")
                    if obj and moved:
                        obj.group = group
                        obj.title = row["title"]
                        obj.order = row["order"]
                        obj.save(update_fields=["group", "title", "order"])
                    elif obj:
                        obj.title = row["title"]
                        obj.order = row["order"]
                        obj.save(update_fields=["title", "order"])
                    else:
                        obj = DouIndicator.objects.create(title=row["title"], group=group, order=row["order"], who_create=None)
                    keep.append(obj.pk)
                removed = existing.exclude(pk__in=keep)
                if PlanIndicator.objects.filter(indicator_id__in=removed.values("pk")).exists():
                    raise CatalogError("Показатель используется в плане")
                removed.delete()
        except CatalogError as exc:
            return {"ok": False, "message": exc.message}
        return {"ok": True, "id": 0 if loose else group.pk, "title": "Без группы" if loose else group.title}

    @staticmethod
    def delete_group(pk):
        group = DouIndicatorGroup.objects.filter(pk=pk, who_create__isnull=True).first()
        if not group:
            return {"ok": False, "message": "Группа не найдена"}
        if PlanIndicator.objects.filter(indicator__group=group).exists():
            return {"ok": False, "message": "Показатель используется в плане"}
        group.delete()
        return {"ok": True, "id": group.pk}


class PlanIndicatorGroup(models.Model):
    plan = models.ForeignKey(Plans, related_name="indicator_groups", on_delete=models.CASCADE)
    title = models.CharField(max_length=128, blank=True, null=True)
    order = models.IntegerField(default=0)

    class Meta:
        verbose_name = "Группа показателей плана"
        verbose_name_plural = "Группы показателей плана"
        ordering = ("order", "pk")

    def __str__(self):
        return f"{self.plan} – {self.title}"


class PlanIndicator(models.Model):
    DUE_ABSOLUTE = "absolute"
    DUE_EMPLOYMENT = "employment"
    UNIT_DAYS = "days"
    UNIT_MONTHS = "months"
    UNIT_YEARS = "years"

    plan = models.ForeignKey(Plans, related_name="indicators", on_delete=models.CASCADE)
    group = models.ForeignKey(PlanIndicatorGroup, related_name="indicators", null=True, blank=True, default=None, on_delete=models.CASCADE)
    indicator = models.ForeignKey("DouIndicator", related_name="plan_rows", on_delete=models.PROTECT)
    order = models.IntegerField(default=0)
    due_kind = models.CharField(max_length=16, default=DUE_ABSOLUTE)
    due_date = models.DateField(null=True, blank=True)
    offset_value = models.IntegerField(null=True, blank=True)
    offset_unit = models.CharField(max_length=16, default=UNIT_DAYS, blank=True)
    executors = models.JSONField(default=list, blank=True, help_text="Исполнители")

    class Meta:
        verbose_name = "Показатель плана"
        verbose_name_plural = "Показатели плана"
        ordering = ("order", "pk")

    def __str__(self):
        return f"{self.plan} – {self.indicator_id}"

    def as_json(self):
        return {
            "id": self.pk,
            "indicatorId": self.indicator_id,
            "order": self.order,
            "dueKind": self.due_kind,
            "dueDate": self.due_date.strftime("%Y-%m-%d") if self.due_date else "",
            "offsetValue": self.offset_value,
            "offsetUnit": self.offset_unit or self.UNIT_DAYS,
            "executors": self.executors if isinstance(self.executors, list) else [],
        }

    @staticmethod
    def _parse_date(value):
        from datetime import datetime

        text = str(value or "").strip()
        if not text:
            return None
        for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
            try:
                return datetime.strptime(text[:10], fmt).date()
            except ValueError:
                continue
        return None

    @staticmethod
    def clean_row(row, order, who=None, allow_user=False, kept_ids=None):
        indicator = DouIndicator.objects.filter(pk=row.get("indicatorId")).first()
        if not indicator:
            return None, "Выберите показатель"
        if indicator.who_create_id:
            own = bool(allow_user and who is not None and indicator.who_create_id == getattr(who, "pk", None))
            kept = bool(allow_user and indicator.pk in (kept_ids or set()))
            if not own and not kept:
                return None, "Задача недоступна"
        due_kind = row.get("dueKind") or PlanIndicator.DUE_ABSOLUTE
        if due_kind not in (PlanIndicator.DUE_ABSOLUTE, PlanIndicator.DUE_EMPLOYMENT):
            return None, "Укажите вид срока"
        due_date = None
        offset_value = None
        offset_unit = PlanIndicator.UNIT_DAYS
        if due_kind == PlanIndicator.DUE_ABSOLUTE:
            due_date = PlanIndicator._parse_date(row.get("dueDate"))
            if not due_date:
                return None, "Укажите дату выполнения"
        else:
            raw = row.get("offsetValue")
            if raw in (None, ""):
                return None, "Укажите сдвиг от даты приема на работу"
            try:
                offset_value = int(raw)
            except (TypeError, ValueError):
                return None, "Сдвиг должен быть числом"
            offset_unit = row.get("offsetUnit") or PlanIndicator.UNIT_DAYS
            if offset_unit not in (PlanIndicator.UNIT_DAYS, PlanIndicator.UNIT_MONTHS, PlanIndicator.UNIT_YEARS):
                return None, "Укажите единицу сдвига"
        return {
            "indicator": indicator,
            "order": order,
            "due_kind": due_kind,
            "due_date": due_date,
            "offset_value": offset_value,
            "offset_unit": offset_unit,
            "executors": PlanIndicator.clean_executors(row.get("executors")),
        }, None

    @staticmethod
    def clean_executors(value):
        ids = DocumentReview.parse_addressee_ids(value)
        if not ids:
            return []
        doctors = {item.pk: item for item in DoctorProfile.objects.filter(pk__in=ids).select_related("podrazdeleniye")}
        result = []
        for doctor_id in ids:
            doctor = doctors.get(doctor_id)
            if not doctor:
                continue
            department = ""
            if doctor.podrazdeleniye_id and doctor.podrazdeleniye:
                department = doctor.podrazdeleniye.title or ""
            result.append({"id": doctor.pk, "fio": doctor.get_full_fio() or "", "department": department})
        return result


class PlanTest(models.Model):
    plan = models.ForeignKey(Plans, related_name="tests", on_delete=models.CASCADE)
    group = models.ForeignKey(PlanIndicatorGroup, related_name="tests", null=True, blank=True, default=None, on_delete=models.CASCADE)
    test = models.ForeignKey("DouTest", related_name="plan_rows", on_delete=models.PROTECT)
    order = models.IntegerField(default=0)
    due_kind = models.CharField(max_length=16, default=PlanIndicator.DUE_ABSOLUTE)
    due_date = models.DateField(null=True, blank=True)
    offset_value = models.IntegerField(null=True, blank=True)
    offset_unit = models.CharField(max_length=16, default=PlanIndicator.UNIT_DAYS, blank=True)

    class Meta:
        verbose_name = "Тест плана"
        verbose_name_plural = "Тесты плана"
        ordering = ("order", "pk")

    def __str__(self):
        return f"{self.plan} – {self.test_id}"

    def as_json(self):
        return {
            "id": self.pk,
            "testId": self.test_id,
            "order": self.order,
            "dueKind": self.due_kind,
            "dueDate": self.due_date.strftime("%Y-%m-%d") if self.due_date else "",
            "offsetValue": self.offset_value,
            "offsetUnit": self.offset_unit or PlanIndicator.UNIT_DAYS,
        }

    @staticmethod
    def clean_row(row, order):
        test = DouTest.objects.filter(pk=row.get("testId")).first()
        if not test:
            return None, "Выберите тест"
        due_kind = row.get("dueKind") or PlanIndicator.DUE_ABSOLUTE
        if due_kind not in (PlanIndicator.DUE_ABSOLUTE, PlanIndicator.DUE_EMPLOYMENT):
            return None, "Укажите вид срока"
        due_date = None
        offset_value = None
        offset_unit = PlanIndicator.UNIT_DAYS
        if due_kind == PlanIndicator.DUE_ABSOLUTE:
            due_date = PlanIndicator._parse_date(row.get("dueDate"))
            if not due_date:
                return None, "Укажите дату выполнения"
        else:
            raw = row.get("offsetValue")
            if raw in (None, ""):
                return None, "Укажите сдвиг от даты приема на работу"
            try:
                offset_value = int(raw)
            except (TypeError, ValueError):
                return None, "Сдвиг должен быть числом"
            offset_unit = row.get("offsetUnit") or PlanIndicator.UNIT_DAYS
            if offset_unit not in (PlanIndicator.UNIT_DAYS, PlanIndicator.UNIT_MONTHS, PlanIndicator.UNIT_YEARS):
                return None, "Укажите единицу сдвига"
        return {
            "test": test,
            "order": order,
            "due_kind": due_kind,
            "due_date": due_date,
            "offset_value": offset_value,
            "offset_unit": offset_unit,
        }, None


class DouTest(models.Model):
    ORDER_SEQUENTIAL = "sequential"
    ORDER_RANDOM = "random"
    PICK_ALL = "all"
    PICK_RANDOM = "random"

    title = models.CharField(max_length=128, blank=True, null=True)
    order_mode = models.CharField(max_length=16, default=ORDER_SEQUENTIAL)
    question_limit = models.PositiveIntegerField(null=True, blank=True)
    question_pick = models.CharField(max_length=16, default=PICK_ALL)

    class Meta:
        verbose_name = "Тест"
        verbose_name_plural = "Тесты"

    def __str__(self):
        return f"{self.title}"

    @property
    def json(self):
        return {"id": self.id, "title": self.title or ""}

    @staticmethod
    def get_list():
        return [row.json for row in DouTest.objects.all().order_by("title", "pk")]

    @staticmethod
    def get_details(pk):
        if pk in (None, -1, "-1"):
            return {"ok": True, "id": -1, "title": "", "questions": [], **DouTest.empty_settings()}
        test = DouTest.objects.filter(pk=pk).prefetch_related("questions__answers").first()
        if not test:
            return {"ok": False, "message": "Тест не найден"}
        questions = []
        for question in test.questions.all().order_by("order", "pk"):
            questions.append(
                {
                    "id": question.pk,
                    "text": question.text or "",
                    "order": question.order,
                    "answers": [row.as_json() for row in question.answers.all().order_by("order", "pk")],
                }
            )
        return {"ok": True, "id": test.pk, "title": test.title or "", "questions": questions, **test.settings_json()}

    @staticmethod
    def empty_settings():
        return {"orderMode": DouTest.ORDER_SEQUENTIAL, "questionLimit": None, "questionPick": DouTest.PICK_ALL}

    def settings_json(self):
        order_mode = self.order_mode if self.order_mode in (self.ORDER_SEQUENTIAL, self.ORDER_RANDOM) else self.ORDER_SEQUENTIAL
        question_pick = self.question_pick if self.question_pick in (self.PICK_ALL, self.PICK_RANDOM) else self.PICK_ALL
        return {"orderMode": order_mode, "questionLimit": self.question_limit, "questionPick": question_pick}

    @staticmethod
    def clean_settings(order_mode, question_limit, question_pick, question_count):
        if order_mode not in (DouTest.ORDER_SEQUENTIAL, DouTest.ORDER_RANDOM):
            return None, "Укажите порядок вопросов"
        if question_pick not in (DouTest.PICK_ALL, DouTest.PICK_RANDOM):
            return None, "Укажите набор вопросов"
        if question_limit in (None, ""):
            limit = None
        else:
            try:
                limit = int(question_limit)
            except (TypeError, ValueError):
                return None, "Укажите количество вопросов"
            if limit < 1:
                return None, "Укажите количество вопросов"
            if limit > question_count:
                return None, "Количество вопросов больше, чем в тесте"
        return {"order_mode": order_mode, "question_limit": limit, "question_pick": question_pick}, None

    @staticmethod
    def save_test(pk, title, questions, order_mode=None, question_limit=None, question_pick=None):
        from django.db import transaction

        title = (title or "").strip()
        if not title:
            return {"ok": False, "message": "Укажите название"}
        cleaned_questions = []
        for index, question in enumerate(questions or []):
            text = (question.get("text") or "").strip()
            if not text:
                return {"ok": False, "message": "Укажите текст вопроса"}
            cleaned_answers = []
            has_correct = False
            for answer_index, answer in enumerate(question.get("answers") or []):
                answer_text = (answer.get("text") or "").strip()
                if not answer_text:
                    return {"ok": False, "message": "Укажите текст варианта ответа"}
                is_correct = bool(answer.get("isCorrect"))
                if is_correct:
                    has_correct = True
                cleaned_answers.append({"text": answer_text, "is_correct": is_correct, "order": answer_index})
            if not cleaned_answers:
                return {"ok": False, "message": "Добавьте вариант ответа"}
            if not has_correct:
                return {"ok": False, "message": "Отметьте правильный ответ"}
            cleaned_questions.append({"text": text, "order": index, "answers": cleaned_answers})
        settings, error = DouTest.clean_settings(order_mode or DouTest.ORDER_SEQUENTIAL, question_limit, question_pick or DouTest.PICK_ALL, len(cleaned_questions))
        if error:
            return {"ok": False, "message": error}
        with transaction.atomic():
            if pk in (None, -1, "-1"):
                test = DouTest(title=title)
            else:
                test = DouTest.objects.filter(pk=pk).first()
                if not test:
                    return {"ok": False, "message": "Тест не найден"}
                test.title = title
            test.order_mode = settings["order_mode"]
            test.question_limit = settings["question_limit"]
            test.question_pick = settings["question_pick"]
            test.save()
            DouTestQuestion.objects.filter(test=test).delete()
            for row in cleaned_questions:
                question = DouTestQuestion.objects.create(test=test, text=row["text"], order=row["order"])
                for answer in row["answers"]:
                    DouTestAnswer.objects.create(question=question, **answer)
        return {"ok": True, "id": test.pk, "title": test.title}

    @staticmethod
    def _cell_text(value):
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    @staticmethod
    def _is_correct_mark(value):
        return DouTest._cell_text(value) == "1"

    @staticmethod
    def _question_error(question, message):
        text = question["text"]
        short = text if len(text) <= 80 else f"{text[:80]}…"
        return {"ok": False, "message": f"{message}: {short}"}

    @staticmethod
    def parse_xlsx(uploaded):
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(filename=uploaded, data_only=True, read_only=True)
        except Exception:
            return {"ok": False, "message": "Не удалось прочитать файл"}
        try:
            if not workbook.sheetnames:
                return {"ok": False, "message": "В файле нет листа"}
            sheet = workbook[workbook.sheetnames[0]]
            questions = []
            current = None
            for row in sheet.iter_rows(max_col=3, values_only=True):
                cells = list(row) + [None, None, None]
                question_text = DouTest._cell_text(cells[0])
                answer_text = DouTest._cell_text(cells[1])
                if not questions and not current and question_text.lower() == "вопрос":
                    continue
                if not question_text and not answer_text:
                    continue
                if question_text:
                    if current is not None:
                        error = DouTest._closed_question(current)
                        if error:
                            return error
                    current = {"text": question_text, "order": len(questions), "answers": []}
                    questions.append(current)
                if not answer_text:
                    continue
                if current is None:
                    return {"ok": False, "message": "Вариант ответа без вопроса"}
                if len(answer_text) > 512:
                    return DouTest._question_error(current, "Вариант ответа длиннее 512 символов")
                current["answers"].append(
                    {
                        "text": answer_text,
                        "isCorrect": DouTest._is_correct_mark(cells[2]),
                        "order": len(current["answers"]),
                    }
                )
            if current is not None:
                error = DouTest._closed_question(current)
                if error:
                    return error
            if not questions:
                return {"ok": False, "message": "В файле нет вопросов"}
            return {"ok": True, "questions": questions}
        finally:
            workbook.close()

    @staticmethod
    def _closed_question(question):
        if not question["answers"]:
            return DouTest._question_error(question, "Добавьте вариант ответа")
        if not any(row["isCorrect"] for row in question["answers"]):
            return DouTest._question_error(question, "Отметьте правильный ответ")
        return None


class DouTestQuestion(models.Model):
    test = models.ForeignKey(DouTest, related_name="questions", on_delete=models.CASCADE)
    text = models.TextField(blank=True, default="")
    order = models.IntegerField(default=0)

    class Meta:
        verbose_name = "Вопрос теста"
        verbose_name_plural = "Вопросы теста"
        ordering = ("order", "pk")

    def __str__(self):
        return f"{self.test} – {self.text}"


class DouTestAnswer(models.Model):
    question = models.ForeignKey(DouTestQuestion, related_name="answers", on_delete=models.CASCADE)
    text = models.CharField(max_length=512, blank=True, default="")
    is_correct = models.BooleanField(default=False)
    order = models.IntegerField(default=0)

    class Meta:
        verbose_name = "Вариант ответа"
        verbose_name_plural = "Варианты ответа"
        ordering = ("order", "pk")

    def __str__(self):
        return f"{self.question_id} – {self.text}"

    def as_json(self):
        return {"id": self.pk, "text": self.text or "", "isCorrect": self.is_correct, "order": self.order}


class DouTestAttempt(models.Model):
    test = models.ForeignKey(DouTest, related_name="attempts", db_index=True, on_delete=models.CASCADE)
    doctor = models.ForeignKey(DoctorProfile, related_name="dou_test_attempts", db_index=True, on_delete=models.CASCADE)
    answer_ids = models.JSONField(default=list, blank=True, help_text="Отмеченные варианты")
    confirmed_at = models.DateTimeField(null=True, blank=True, default=None)
    document = models.OneToOneField("Documents", related_name="test_attempt", db_index=True, null=True, blank=True, default=None, on_delete=models.SET_NULL)

    class Meta:
        verbose_name = "Прохождение теста"
        verbose_name_plural = "Прохождения тестов"
        unique_together = ("test", "doctor")

    def __str__(self):
        return f"{self.test_id} – {self.doctor_id}"

    @staticmethod
    def _clean_ids(raw_ids, valid_ids):
        cleaned = []
        seen = set()
        for raw in raw_ids or []:
            try:
                pk = int(raw)
            except (TypeError, ValueError):
                return None
            if pk not in valid_ids or pk in seen:
                if pk not in valid_ids:
                    return None
                continue
            seen.add(pk)
            cleaned.append(pk)
        return cleaned

    @staticmethod
    def _score(questions, selected_ids):
        selected = set(selected_ids)
        correct = 0
        for question in questions:
            right = {row["id"] for row in question["answers"] if row["correct"]}
            picked = {row["id"] for row in question["answers"] if row["id"] in selected}
            if picked == right:
                correct += 1
        total = len(questions)
        percent = 0 if total == 0 else int((correct * 100 + total // 2) // total)
        return correct, total, percent

    @staticmethod
    def _questions(test):
        questions = []
        valid_ids = set()
        for question in test.questions.all().order_by("order", "pk"):
            answers = []
            for row in question.answers.all().order_by("order", "pk"):
                valid_ids.add(row.pk)
                answers.append({"id": row.pk, "text": row.text or "", "correct": row.is_correct})
            questions.append({"id": question.pk, "text": question.text or "", "answers": answers})
        return questions, valid_ids

    @staticmethod
    def _can_reset(doctor):
        user = getattr(doctor, "user", None) if doctor else None
        return bool(user and getattr(user, "is_superuser", False))

    @staticmethod
    def as_quiz(test, attempt, questions, can_reset=False, document=None, editable=None):
        selected = []
        if attempt:
            for raw in attempt.answer_ids or []:
                try:
                    pk = int(raw)
                except (TypeError, ValueError):
                    continue
                if any(pk == row["id"] for question in questions for row in question["answers"]) and pk not in selected:
                    selected.append(pk)
        linked = document or (attempt.document if attempt and attempt.document_id else None)
        confirmed = bool(attempt and attempt.confirmed_at) or bool(linked and linked.time_confirm)
        if confirmed:
            public_questions = [
                {
                    "id": question["id"],
                    "text": question["text"],
                    "answers": [{"id": row["id"], "text": row["text"], "isCorrect": bool(row["correct"])} for row in question["answers"]],
                }
                for question in questions
            ]
        else:
            public_questions = [
                {
                    "id": question["id"],
                    "text": question["text"],
                    "answers": [{"id": row["id"], "text": row["text"]} for row in question["answers"]],
                }
                for question in questions
            ]
        payload = {
            "ok": True,
            "id": test.pk,
            "title": test.title or "",
            "questions": public_questions,
            "selected": selected,
            "confirmed": confirmed,
            "canReset": bool(can_reset),
            "documentId": linked.pk if linked else None,
            "whoConfirmed": Documents.confirmation_text(linked) if confirmed and linked and linked.time_confirm else "",
        }
        if confirmed:
            correct, total, percent = DouTestAttempt._score(questions, selected)
            payload["correct"] = correct
            payload["total"] = total
            payload["percent"] = percent
        if editable is None:
            editable = bool(can_reset or not confirmed)
        payload["editable"] = bool(editable)
        return payload

    @staticmethod
    def _quiz(test, attempt, doctor, document=None, editable=None):
        questions, _valid = DouTestAttempt._questions(test)
        linked = document or (attempt.document if attempt and attempt.document_id else None)
        return DouTestAttempt.as_quiz(test, attempt, questions, DouTestAttempt._can_reset(doctor), linked, editable)

    @staticmethod
    def get_for(test_id, doctor):
        if not doctor:
            return {"ok": False, "message": "Сотрудник не найден"}
        test = DouTest.objects.filter(pk=test_id).prefetch_related("questions__answers").first()
        if not test:
            return {"ok": False, "message": "Тест не найден"}
        attempt = DouTestAttempt.objects.filter(test=test, doctor=doctor).select_related("document", "document__who_confirm", "document__who_confirm__podrazdeleniye").first()
        return DouTestAttempt._quiz(test, attempt, doctor)

    @staticmethod
    def get_for_document(document_id, doctor):
        if not doctor:
            return {"ok": False, "message": "Сотрудник не найден"}
        document = Documents.objects.filter(pk=document_id).select_related("type_document", "type_document__test", "who_confirm", "who_confirm__podrazdeleniye").first()
        if not document or not Documents.can_see_document(document, doctor):
            return {"ok": False, "message": "Документ не найден"}
        linked_test = document.type_document.test if document.type_document_id and document.type_document and document.type_document.test_id else None
        test = DouTest.objects.filter(pk=linked_test.pk).prefetch_related("questions__answers").first() if linked_test else None
        if not test:
            return {"ok": False, "message": "Тест не найден"}
        attempt = DouTestAttempt.objects.filter(document=document).first()
        confirmed = bool(document.time_confirm or (attempt and attempt.confirmed_at))
        owner_id = attempt.doctor_id if attempt else document.who_create_id
        editable = bool(DouTestAttempt._can_reset(doctor) or (owner_id == doctor.pk and not confirmed))
        return DouTestAttempt._quiz(test, attempt, doctor, document, editable)

    @staticmethod
    def _test_type(test):
        return TypeDocuments.objects.filter(test=test).order_by("title", "pk").first()

    @staticmethod
    def _create_document(type_doc, doctor):
        from directions.models import Issledovaniya

        schema = TypeDocumentsSchema.objects.filter(type_document=type_doc).order_by("-created_at", "-pk").first()
        if schema is None:
            schema = TypeDocumentsSchema.get_or_create_for_type(type_doc)
        obj = Documents.objects.create(type_document=type_doc, who_create=doctor, body_values={}, schema=schema)
        Issledovaniya.objects.create(document=obj, research=None, creator=doctor)
        title = f"{type_doc.title or 'Документ'} №{obj.pk}"
        Documents.log_action(obj.pk, Documents.LOG_CREATE, doctor, {"title": title, "typeId": type_doc.pk, "typeTitle": type_doc.title or ""})
        return obj

    @staticmethod
    def _bind_document(test, doctor, attempt):
        if attempt and attempt.document_id:
            return attempt.document, None
        type_doc = DouTestAttempt._test_type(test)
        if not type_doc:
            return None, "Нет вида документа для этого теста"
        existing = Documents.objects.filter(type_document=type_doc, who_create=doctor, test_attempt__isnull=True).order_by("-pk").first()
        if existing:
            return existing, None
        return DouTestAttempt._create_document(type_doc, doctor), None

    @staticmethod
    def _log_title(document):
        type_title = "Документ"
        if document.type_document_id:
            type_doc = TypeDocuments.objects.filter(pk=document.type_document_id).only("title").first()
            if type_doc and type_doc.title:
                type_title = type_doc.title
        return f"{type_title} №{document.pk}", type_title

    @staticmethod
    def _stamp_confirm(document, doctor):
        from django.utils import timezone

        from directions.models import Issledovaniya

        now = timezone.now()
        document.time_confirm = now
        document.who_confirm = doctor
        document.save(update_fields=["time_confirm", "who_confirm"])
        iss = Issledovaniya.objects.filter(document=document).order_by("pk").first()
        if iss:
            iss.doc_confirmation = doctor
            iss.time_confirmation = now
            iss.doc_save = doctor
            iss.time_save = now
            iss.save(update_fields=["doc_confirmation", "time_confirmation", "doc_save", "time_save"])
        title, type_title = DouTestAttempt._log_title(document)
        Documents.log_action(
            document.pk,
            Documents.LOG_CONFIRM,
            doctor,
            {"title": title, "typeId": document.type_document_id, "typeTitle": type_title, "confirmed": True},
        )
        return now

    @staticmethod
    def _clear_confirm(document_id, doctor):
        from directions.models import Issledovaniya

        document = Documents.objects.filter(pk=document_id).first()
        if not document or not document.time_confirm:
            return
        iss = Issledovaniya.objects.filter(document=document).order_by("pk").first()
        if iss:
            iss.time_confirmation = None
            iss.doc_confirmation = None
            iss.executor_confirmation = None
            iss.save(update_fields=["time_confirmation", "doc_confirmation", "executor_confirmation"])
        document.time_confirm = None
        document.who_confirm = None
        document.save(update_fields=["time_confirm", "who_confirm"])
        title, type_title = DouTestAttempt._log_title(document)
        Documents.log_action(document.pk, Documents.LOG_RESET, doctor, {"title": title, "typeId": document.type_document_id, "typeTitle": type_title})

    @staticmethod
    def save_for(test_id, doctor, answer_ids, confirm, document_id=None):
        if not doctor:
            return {"ok": False, "message": "Сотрудник не найден"}
        with transaction.atomic():
            document = None
            if document_id not in (None, "", -1, "-1"):
                document = Documents.objects.select_for_update().filter(pk=document_id).first()
                if not document or not Documents.can_see_document(document, doctor):
                    return {"ok": False, "message": "Документ не найден"}
                type_doc = TypeDocuments.objects.filter(pk=document.type_document_id).select_related("test").first()
                test = type_doc.test if type_doc and type_doc.test_id else None
                if not test:
                    return {"ok": False, "message": "Тест не найден"}
            else:
                test = DouTest.objects.select_for_update().filter(pk=test_id).first()
                if not test:
                    return {"ok": False, "message": "Тест не найден"}
            attempt = DouTestAttempt.objects.select_for_update().filter(document=document).first() if document is not None else None
            if attempt is None:
                attempt = DouTestAttempt.objects.select_for_update().filter(test=test, doctor=doctor).first()
            if attempt and attempt.confirmed_at and not DouTestAttempt._can_reset(doctor):
                return {"ok": False, "message": "Тест подтверждён"}
            if attempt and document is not None and attempt.document_id and attempt.document_id != document.pk:
                return {"ok": False, "message": "Прохождение уже сохранено в другом документе"}
            valid_ids = set(DouTestAnswer.objects.filter(question__test=test).values_list("pk", flat=True))
            cleaned = DouTestAttempt._clean_ids(answer_ids, valid_ids)
            if cleaned is None:
                return {"ok": False, "message": "Некорректный ответ"}
            if attempt is None:
                attempt = DouTestAttempt(test=test, doctor=doctor)
            attempt.answer_ids = cleaned
            if confirm:
                if document is not None:
                    bound, error = document, None
                else:
                    bound, error = DouTestAttempt._bind_document(test, doctor, attempt)
                if error:
                    return {"ok": False, "message": error}
                attempt.document = bound
                if not bound.time_confirm:
                    attempt.confirmed_at = DouTestAttempt._stamp_confirm(bound, doctor)
                elif not attempt.confirmed_at:
                    attempt.confirmed_at = bound.time_confirm
            elif document is not None:
                attempt.document = document
            attempt.save()
            if not confirm and attempt.document_id:
                title, type_title = DouTestAttempt._log_title(attempt.document)
                Documents.log_action(
                    attempt.document_id,
                    Documents.LOG_SAVE,
                    doctor,
                    {"title": title, "typeId": attempt.document.type_document_id, "typeTitle": type_title, "confirmed": bool(attempt.confirmed_at)},
                )
            saved_id = attempt.document_id
        test = DouTest.objects.filter(pk=test.pk).prefetch_related("questions__answers").first()
        attempt = DouTestAttempt.objects.filter(pk=attempt.pk).select_related("document", "document__who_confirm", "document__who_confirm__podrazdeleniye").first()
        linked = attempt.document if attempt and attempt.document_id else None
        editable = None
        if saved_id and linked:
            confirmed = bool(linked.time_confirm or (attempt and attempt.confirmed_at))
            owner_id = attempt.doctor_id if attempt else linked.who_create_id
            editable = bool(DouTestAttempt._can_reset(doctor) or (owner_id == doctor.pk and not confirmed))
        return DouTestAttempt._quiz(test, attempt, doctor, linked, editable)

    @staticmethod
    def reset_for(test_id, doctor):
        if not doctor:
            return {"ok": False, "message": "Сотрудник не найден"}
        if not DouTestAttempt._can_reset(doctor):
            return {"ok": False, "message": "Недостаточно прав"}
        with transaction.atomic():
            test = DouTest.objects.select_for_update().filter(pk=test_id).first()
            if not test:
                return {"ok": False, "message": "Тест не найден"}
            rows = list(DouTestAttempt.objects.select_for_update().filter(test=test))
            for row in rows:
                if row.document_id:
                    DouTestAttempt._clear_confirm(row.document_id, doctor)
            DouTestAttempt.objects.filter(pk__in=[row.pk for row in rows]).delete()
        test = DouTest.objects.filter(pk=test.pk).prefetch_related("questions__answers").first()
        return DouTestAttempt._quiz(test, None, doctor)


class DocumentFieldGroups(models.Model):
    title = models.CharField(max_length=550, help_text="Название группы")
    show_title = models.BooleanField(default=True, blank=True)
    type_document = models.ForeignKey(TypeDocuments, db_index=True, on_delete=models.CASCADE, help_text="Вид документа")
    order = models.IntegerField()
    hide = models.BooleanField(default=False, blank=True)
    visibility = models.TextField(default="", blank=True)
    fields_inline = models.BooleanField(default=False, blank=True)

    class Meta:
        verbose_name = "Группа полей документа"
        verbose_name_plural = "Группы полей документа"

    def __str__(self):
        return f"{self.type_document} – {self.title}"

    def as_json(self):
        return {
            "pk": self.pk,
            "order": self.order,
            "title": self.title,
            "show_title": self.show_title,
            "hide": self.hide,
            "visibility": self.visibility,
            "fieldsInline": self.fields_inline,
            "fields": [field.as_json() for field in self.fields.all().order_by("order")],
        }

    @staticmethod
    def get_structure(type_document_id):
        type_doc = TypeDocuments.objects.filter(pk=type_document_id).first()
        if not type_doc:
            return {"ok": False, "message": "Вид документа не найден"}
        groups = [group.as_json() for group in DocumentFieldGroups.objects.filter(type_document=type_doc).prefetch_related("fields").order_by("order")]
        orphan_fields = DocumentFields.objects.filter(type_document=type_doc, group__isnull=True).order_by("order")
        if orphan_fields.exists():
            next_order = max([group["order"] for group in groups], default=0) + 1
            groups.append(
                {
                    "pk": -1,
                    "order": next_order,
                    "title": "Основное",
                    "show_title": True,
                    "hide": False,
                    "visibility": "",
                    "fieldsInline": False,
                    "fields": [field.as_json() for field in orphan_fields],
                }
            )
        return {"ok": True, "id": type_doc.pk, "title": type_doc.title or "", "code": type_doc.code or "", "groups": groups}

    @staticmethod
    @transaction.atomic
    def save_structure(type_document_id, groups):
        type_doc = TypeDocuments.objects.filter(pk=type_document_id).first()
        if not type_doc:
            return {"ok": False, "message": "Вид документа не найден"}
        for group_data in groups or []:
            group_pk = group_data.get("pk", -1)
            if group_pk in (None, -1, "-1"):
                group = DocumentFieldGroups(
                    title=group_data.get("title") or "",
                    show_title=bool(group_data.get("show_title", True)),
                    type_document=type_doc,
                    order=int(group_data.get("order") or 0),
                    hide=bool(group_data.get("hide", False)),
                    visibility=group_data.get("visibility") or "",
                    fields_inline=bool(group_data.get("fieldsInline", False)),
                )
            else:
                group = DocumentFieldGroups.objects.filter(pk=group_pk, type_document=type_doc).first()
                if not group:
                    continue
                group.title = group_data.get("title") or ""
                group.show_title = bool(group_data.get("show_title", True))
                group.order = int(group_data.get("order") or 0)
                group.hide = bool(group_data.get("hide", False))
                group.visibility = group_data.get("visibility") or ""
                group.fields_inline = bool(group_data.get("fieldsInline", False))
            group.save()
            for field_data in group_data.get("fields") or []:
                DocumentFields.save_field(type_doc, group, field_data)
        return {"ok": True, "id": type_doc.pk}


class Nomenclature(models.Model):
    title = models.CharField(max_length=128, blank=True, null=True)
    type_document = models.ForeignKey(TypeDocuments, default=None, blank=True, null=True, db_index=True, help_text="Тип документа", on_delete=models.SET_NULL)
    department = models.ForeignKey(
        Podrazdeleniya, related_name="department_nomenclature", help_text="Подразделение", db_index=True, null=True, blank=True, default=None, on_delete=models.SET_NULL
    )

    class Meta:
        verbose_name = "Номенклатура"
        verbose_name_plural = "Носенклатуры"

    def __str__(self):
        return f"{self.title}"


class Generators(models.Model):
    title = models.CharField(max_length=400, help_text="Название генератора номера")
    hospital = models.ForeignKey(Hospitals, on_delete=models.CASCADE, db_index=True, verbose_name="Больница")
    year = models.IntegerField(verbose_name="Год", db_index=True)
    is_active = models.BooleanField(verbose_name="Активность диапазона", db_index=True)
    start = models.PositiveBigIntegerField(verbose_name="Начало диапазона")
    end = models.PositiveBigIntegerField(verbose_name="Конец диапазона", null=True, blank=True, default=None)
    last = models.PositiveBigIntegerField(verbose_name="Последнее значение диапазона", null=True, blank=True)
    free_numbers = ArrayField(models.PositiveBigIntegerField(verbose_name="Свободные номера"), default=list, blank=True)
    prepend_length = models.PositiveSmallIntegerField(verbose_name="Длина номера", help_text='Если номер короче, впереди будет добавлено недостающее кол-во "0"')
    department = models.ForeignKey(
        Podrazdeleniya, related_name="department_for_generator", help_text="Подразделение", db_index=True, null=True, blank=True, default=None, on_delete=models.SET_NULL
    )

    def __str__(self):
        return f"{self.hospital} {self.year} {self.is_active} {self.start} — {self.end} ({self.last})"

    class Meta:
        verbose_name = "Генератор номеров"
        verbose_name_plural = "Генераторы номеров"


class MetaFieldsForAllTypeDocuments(models.Model):
    TYPES = (
        (0, "Text"),
        (1, "Date"),
        (2, "Calc"),
        (3, "List"),
        (4, "Dict"),
        (5, "Radio"),
        (6, "Number"),
        (7, "Number range"),
        (8, "Time HH:MM"),
        (9, "Table"),
        (10, "Исполнитель"),
    )
    title = models.CharField(max_length=400, help_text="Название поля ввода")
    short_title = models.CharField(max_length=400, default="", blank=True, help_text="Синоним-короткое название поля ввода")
    order = models.IntegerField()
    default_value = models.TextField(blank=True, default="")
    input_templates = models.TextField()
    hide = models.BooleanField()
    lines = models.IntegerField(default=3)
    field_type = models.SmallIntegerField(default=0, choices=TYPES, blank=True)
    required = models.BooleanField(default=False, blank=True)

    class Meta:
        verbose_name = "Мета реквизит для всех документов"
        verbose_name_plural = "Мета реквизиты для всех документов"

    def __str__(self):
        return f"{self.title}"


class DocumentFields(models.Model):
    TYPES = (
        (0, "Text"),
        (1, "Date"),
        (2, "Calc"),
        (3, "List"),
        (4, "Dict"),
        (5, "Radio"),
        (6, "Number"),
        (7, "Number range"),
        (8, "Time HH:MM"),
        (9, "Table"),
        (10, "Исполнитель"),
    )
    type_document = models.ForeignKey(TypeDocuments, default=None, db_index=True, blank=True, null=True, help_text="Тип документа", on_delete=models.SET_NULL)
    group = models.ForeignKey("DocumentFieldGroups", related_name="fields", default=None, blank=True, null=True, db_index=True, help_text="Группа полей", on_delete=models.CASCADE)
    title = models.CharField(max_length=400, help_text="Название поля ввода")
    short_title = models.CharField(max_length=400, default="", blank=True, help_text="Синоним-короткое название поля ввода")
    order = models.IntegerField()
    default_value = models.TextField(blank=True, default="")
    input_templates = models.TextField()
    hide = models.BooleanField()
    lines = models.IntegerField(default=3)
    field_type = models.SmallIntegerField(default=0, choices=TYPES, blank=True)
    required = models.BooleanField(default=False, blank=True)
    is_meta_attributes = models.BooleanField(default=False, blank=True)
    is_content_attributes = models.BooleanField(default=False, blank=True)

    class Meta:
        verbose_name = "Реквизит для документа"
        verbose_name_plural = "Реквизит для документов"

    def __str__(self):
        return f"{self.type_document}"

    @staticmethod
    def _parse_templates(raw):
        if not raw:
            return []
        try:
            data = json.loads(raw)
            return data if isinstance(data, list) else []
        except (TypeError, ValueError):
            return []

    def as_json(self):
        return {
            "pk": self.pk,
            "order": self.order,
            "title": self.title,
            "short_title": self.short_title,
            "default": self.default_value,
            "values_to_input": self._parse_templates(self.input_templates),
            "new_value": "",
            "hide": self.hide,
            "lines": self.lines,
            "field_type": self.field_type,
            "required": self.required,
            "is_meta_attributes": self.is_meta_attributes,
            "is_content_attributes": self.is_content_attributes,
        }

    @staticmethod
    def save_field(type_doc, group, field_data):
        field_pk = field_data.get("pk", -1)
        values_to_input = field_data.get("values_to_input") or []
        if not isinstance(values_to_input, list):
            values_to_input = []
        templates = json.dumps(values_to_input)
        if field_pk in (None, -1, "-1"):
            field = DocumentFields(
                type_document=type_doc,
                group=group,
                title=field_data.get("title") or "",
                short_title=field_data.get("short_title") or "",
                order=int(field_data.get("order") or 0),
                default_value=field_data.get("default") or "",
                input_templates=templates,
                hide=bool(field_data.get("hide", False)),
                lines=int(field_data.get("lines") or 3),
                field_type=int(field_data.get("field_type") or 0),
                required=bool(field_data.get("required", False)),
                is_meta_attributes=bool(field_data.get("is_meta_attributes", False)),
                is_content_attributes=bool(field_data.get("is_content_attributes", False)),
            )
        else:
            field = DocumentFields.objects.filter(pk=field_pk).first()
            if not field:
                return
            field.type_document = type_doc
            field.group = group
            field.title = field_data.get("title") or ""
            field.short_title = field_data.get("short_title") or ""
            field.order = int(field_data.get("order") or 0)
            field.default_value = field_data.get("default") or ""
            field.input_templates = templates
            field.hide = bool(field_data.get("hide", False))
            field.lines = int(field_data.get("lines") or 3)
            field.field_type = int(field_data.get("field_type") or 0)
            field.required = bool(field_data.get("required", False))
            field.is_meta_attributes = bool(field_data.get("is_meta_attributes", False))
            field.is_content_attributes = bool(field_data.get("is_content_attributes", False))
        field.save()


class TypeDocumentsSchema(models.Model):
    type_document = models.ForeignKey(TypeDocuments, db_index=True, related_name="schemas", on_delete=models.CASCADE, help_text="Вид документа")
    version = models.UUIDField(default=uuid.uuid4, unique=True, db_index=True, help_text="Версия схемы")
    schema = models.JSONField(default=dict, blank=True, help_text="Снимок шаблона со всеми правилами")
    created_at = models.DateTimeField(auto_now_add=True, db_index=True, help_text="Дата создания версии")

    class Meta:
        verbose_name = "Схема вида документа"
        verbose_name_plural = "Схемы видов документов"
        ordering = ["-created_at"]

    def __str__(self):
        title = self.type_document.title if self.type_document else "Документ"
        return f"{title} {self.version}"

    @staticmethod
    def _schema_key(schema):
        return json.dumps(schema, sort_keys=True, ensure_ascii=False, default=str)

    @staticmethod
    def _parse_templates(raw):
        if not raw:
            return []
        try:
            data = json.loads(raw)
            return data if isinstance(data, list) else []
        except (TypeError, ValueError):
            return []

    @staticmethod
    def serialize_field(field):
        from directory.models import ParaclinicInputFieldFileSettings

        values_to_input = TypeDocumentsSchema._parse_templates(field.input_templates)
        field_type = field.field_type
        if field.required and field_type in [10, 12] and "- Не выбрано" not in values_to_input:
            values_to_input = ["- Не выбрано", *values_to_input]
        file_settings = None
        if field_type == 42:
            file_settings = ParaclinicInputFieldFileSettings.get_file_field_settings(field)
        return {
            "pk": field.pk,
            "order": field.order,
            "lines": field.lines,
            "title": field.title,
            "short_title": field.short_title or "",
            "hide": field.hide,
            "values_to_input": values_to_input,
            "field_type": field_type,
            "can_edit": field.can_edit_computed,
            "default_value": field.default_value or "",
            "visibility": field.visibility or "",
            "required": bool(field.required or field.required_set_by_admin),
            "required_set_by_admin": bool(field.required_set_by_admin),
            "helper": field.helper or "",
            "controlParam": field.control_param or "",
            "not_edit": field.not_edit,
            "operator_enter_param": field.operator_enter_param,
            "deniedGroup": field.denied_group.name if field.denied_group else "",
            "isDiagTable": field.is_diag_table,
            "file_settings": file_settings,
            "for_talon": field.for_talon,
            "sign_organization": field.sign_organization,
            "for_extract_card": field.for_extract_card,
            "for_med_certificate": field.for_med_certificate,
            "attached": field.attached or "",
            "layout_link_research_id": field.layout_link_research_id,
            "cdaOption": field.cda_option_id if field.cda_option_id else -1,
        }

    @staticmethod
    def serialize_group(group):
        fields = [TypeDocumentsSchema.serialize_field(field) for field in group.paraclinicinputfield_set.all() if field.field_type != LAYOUT_TEMPLATE_FIELD_TYPE]
        fields.sort(key=lambda row: row["order"])
        return {
            "pk": group.pk,
            "order": group.order,
            "title": group.title,
            "show_title": group.show_title,
            "hide": group.hide,
            "visibility": group.visibility or "",
            "fieldsInline": group.fields_inline,
            "cdaOption": group.cda_option_id if group.cda_option_id else -1,
            "fields": fields,
        }

    @staticmethod
    def dump_from_type(type_doc):
        from django.db.models import Prefetch

        from directory.models import ParaclinicInputField, ParaclinicInputGroups

        monitoring = None
        if type_doc and type_doc.is_plan:
            return {
                "pk": None,
                "title": type_doc.title or "",
                "wide_headers": False,
                "layoutTemplateIds": [],
                "isPlan": True,
                "groups": [],
            }
        if type_doc and type_doc.test_id and type_doc.test:
            return {
                "pk": None,
                "title": type_doc.test.title or "",
                "wide_headers": False,
                "layoutTemplateIds": [],
                "testId": type_doc.test_id,
                "groups": [],
            }
        if type_doc and type_doc.monitoring_id:
            linked = type_doc.monitoring
            if linked and linked.is_monitoring:
                monitoring = linked
        templates = [] if monitoring else (type_doc.get_layout_templates() if type_doc else [])
        sources = [monitoring] if monitoring else templates
        if not sources:
            return None
        groups = []
        group_order = 0
        for template in sources:
            group_qs = (
                ParaclinicInputGroups.objects.filter(research=template)
                .order_by("order")
                .prefetch_related(
                    Prefetch(
                        "paraclinicinputfield_set",
                        queryset=ParaclinicInputField.objects.select_related("denied_group", "file_settings").order_by("order"),
                    )
                )
            )
            for group in group_qs:
                if group.hide:
                    continue
                serialized = TypeDocumentsSchema.serialize_group(group)
                if not serialized["fields"]:
                    continue
                serialized["order"] = group_order
                group_order += 1
                groups.append(serialized)
        first = sources[0]
        result = {
            "pk": first.pk,
            "title": first.title,
            "wide_headers": any(bool(getattr(row, "wide_headers", False)) for row in sources),
            "layoutTemplateIds": [] if monitoring else [row.pk for row in templates],
            "groups": groups,
        }
        if monitoring:
            result["monitoringId"] = monitoring.pk
        return result

    @staticmethod
    def get_or_create_for_type(type_doc):
        dumped = TypeDocumentsSchema.dump_from_type(type_doc)
        latest = TypeDocumentsSchema.objects.filter(type_document=type_doc).order_by("-created_at", "-pk").first()
        if not dumped:
            if latest and isinstance(latest.schema, dict) and (latest.schema.get("monitoringId") or latest.schema.get("testId") or latest.schema.get("isPlan")):
                dumped = {
                    "pk": None,
                    "title": type_doc.title if type_doc else "",
                    "wide_headers": False,
                    "layoutTemplateIds": [],
                    "monitoringId": None,
                    "testId": None,
                    "groups": [],
                }
            else:
                return None
        if latest and TypeDocumentsSchema._schema_key(latest.schema) == TypeDocumentsSchema._schema_key(dumped):
            return latest
        return TypeDocumentsSchema.objects.create(type_document=type_doc, schema=dumped)

    @staticmethod
    def sync_for_layout_template(template):
        if not template:
            return
        type_ids = TypeDocumentLayoutTemplate.objects.filter(layout_template=template).values_list("type_document_id", flat=True)
        qs = TypeDocuments.objects.filter(models.Q(pk__in=type_ids) | models.Q(layout_template=template)).distinct()
        for type_doc in qs:
            TypeDocumentsSchema.get_or_create_for_type(type_doc)

    @staticmethod
    def sync_for_monitoring(research):
        if not research or not research.is_monitoring:
            return
        qs = TypeDocuments.objects.filter(monitoring=research).select_related("monitoring")
        for type_doc in qs:
            TypeDocumentsSchema.get_or_create_for_type(type_doc)


def _link_id_list(raw):
    if raw in (None, "", 0, "0"):
        return []
    if not isinstance(raw, (list, tuple)):
        raw = [raw]
    seen = set()
    result = []
    for item in raw:
        if item in (None, "", 0, "0"):
            continue
        try:
            value = int(item)
        except (TypeError, ValueError):
            return None
        if value <= 0 or value in seen:
            continue
        seen.add(value)
        result.append(value)
    return result


class DocumentCaseLink(models.Model):
    document = models.ForeignKey("Documents", related_name="case_links", on_delete=models.CASCADE)
    document_case = models.ForeignKey(DocumentCase, related_name="document_links", on_delete=models.CASCADE)

    class Meta:
        verbose_name = "Документ в деле"
        verbose_name_plural = "Документы в делах"
        unique_together = ("document", "document_case")

    def __str__(self):
        return f"{self.document_id} {self.document_case_id}"


class DocumentBlockLink(models.Model):
    document = models.ForeignKey("Documents", related_name="block_links", on_delete=models.CASCADE)
    block = models.ForeignKey(DocumentBlock, related_name="document_links", on_delete=models.CASCADE)

    class Meta:
        verbose_name = "Документ в блоке"
        verbose_name_plural = "Документы в блоках"
        unique_together = ("document", "block")

    def __str__(self):
        return f"{self.document_id} {self.block_id}"


class Documents(models.Model):
    type_document = models.ForeignKey(TypeDocuments, db_index=True, default=None, blank=True, null=True, help_text="Тип документа", on_delete=models.SET_NULL)
    schema = models.ForeignKey(
        TypeDocumentsSchema,
        related_name="documents",
        db_index=True,
        default=None,
        blank=True,
        null=True,
        help_text="Снимок схемы шаблона",
        on_delete=models.PROTECT,
    )
    who_create = models.ForeignKey(DoctorProfile, db_index=True, default=None, blank=True, null=True, help_text="Создатель документа", on_delete=models.SET_NULL)
    create_at = models.DateTimeField(auto_now_add=True, help_text="Дата создания документа", db_index=True)
    time_registration = models.DateTimeField(null=True, blank=True, db_index=True, help_text="Время регистрации")
    number_registration = models.CharField(max_length=20, blank=True, help_text="Регистрационный номер документа", db_index=True)
    is_registered = models.BooleanField(default=False, blank=True, help_text="Прошел регистрацию", db_index=True)
    need_approve = models.BooleanField(default=False, blank=True, help_text="Находится на согласовании", db_index=True)
    need_control = models.BooleanField(default=False, blank=True, help_text="Находится на контроле", db_index=True)
    total_approved = models.BooleanField(default=False, blank=True, help_text="Полностью согласован", db_index=True)
    total_completed = models.BooleanField(default=False, blank=True, help_text="Полностью завершен/снят с контроля", db_index=True)
    parent_document = models.ForeignKey("self", db_index=True, related_name="document_p", help_text="Документ основание", blank=True, null=True, default=None, on_delete=models.SET_NULL)
    body_values = models.JSONField(default=dict, blank=True, help_text="Значения полей шаблона")
    time_confirm = models.DateTimeField(null=True, blank=True, db_index=True, help_text="Время подтверждения")
    who_confirm = models.ForeignKey(
        DoctorProfile,
        related_name="document_who_confirm",
        db_index=True,
        default=None,
        blank=True,
        null=True,
        help_text="Кто подтвердил",
        on_delete=models.SET_NULL,
    )
    is_hidden = models.BooleanField(default=False, blank=True, db_index=True, help_text="Скрыт из списков")
    type_case = models.ForeignKey(
        TypeCases,
        related_name="created_documents",
        blank=True,
        null=True,
        db_index=True,
        help_text="Вид дела, из которого создан документ",
        on_delete=models.SET_NULL,
    )
    document_case = models.ForeignKey(
        DocumentCase,
        related_name="documents",
        blank=True,
        null=True,
        default=None,
        db_index=True,
        help_text="Дело",
        on_delete=models.SET_NULL,
    )
    block = models.ForeignKey(
        "DocumentBlock",
        related_name="documents",
        blank=True,
        null=True,
        default=None,
        db_index=True,
        help_text="Блок",
        on_delete=models.SET_NULL,
    )
    plan = models.ForeignKey(
        Plans,
        related_name="documents",
        blank=True,
        null=True,
        default=None,
        db_index=True,
        help_text="План, составленный в документе",
        on_delete=models.SET_NULL,
    )

    HIDDEN_DOCS_GROUP = "Скрытие документа"
    RESET_GROUP = "Сброс документов"
    CANCEL_APPROVE_GROUP = "ДОУ: отмена согласования"
    HISTORY_GROUP = "История документа"
    LOG_CREATE = 260000
    LOG_SAVE = 260001
    LOG_CONFIRM = 260002
    LOG_RESET = 260003
    LOG_HIDE = 260004
    LOG_SHOW = 260005
    LOG_APPROVERS = 260006
    LOG_APPROVE = 260007
    LOG_REJECT = 260008
    LOG_APPROVE_CANCEL = 260009
    LOG_TYPES = (LOG_CREATE, LOG_SAVE, LOG_CONFIRM, LOG_RESET, LOG_HIDE, LOG_SHOW, LOG_APPROVERS, LOG_APPROVE, LOG_REJECT, LOG_APPROVE_CANCEL)

    class Meta:
        verbose_name = "Документ-экземпляр"
        verbose_name_plural = "Документы экземляров"

    def __str__(self):
        return f"{self.type_document}"

    def list_title(self, iss=None, topic=None):
        type_title = self.type_document.title if self.type_document else "Документ"
        fallback = self.number_registration or f"{type_title} №{self.pk}"
        is_saved = bool(self.time_confirm or (iss and (iss.time_confirmation or iss.time_save)))
        topic_text = "" if topic is None else str(topic).strip()
        if is_saved and topic_text:
            return topic_text
        return fallback

    @staticmethod
    def title_with_id(pk, title):
        marker = str(pk)
        text = (title or "").strip()
        suffix = f"№{marker}"
        if text.endswith(suffix):
            text = text[: -len(suffix)].rstrip()
        if text.startswith(f"{marker} ") or text == marker:
            return text
        return f"{marker} {text}" if text else marker

    @staticmethod
    def topic_cda_id():
        from laboratory.settings import CDA_TOPIC_ID_FOR_DOCUMENT_MANAGER

        raw = CDA_TOPIC_ID_FOR_DOCUMENT_MANAGER
        if raw in (None, "", -1, "-1"):
            return None
        try:
            value = int(raw)
        except (TypeError, ValueError):
            return None
        return value if value > 0 else None

    @staticmethod
    def topic_from_research(research):
        topic_id = Documents.topic_cda_id()
        if not topic_id or not research:
            return ""
        for group in research.get("groups") or []:
            for field in group.get("fields") or []:
                cda = field.get("cdaOption")
                if cda in (None, "", -1, "-1"):
                    continue
                try:
                    if int(cda) != topic_id:
                        continue
                except (TypeError, ValueError):
                    continue
                value = (field.get("value") or "").strip()
                if value:
                    return value
        return ""

    @staticmethod
    def _topic_field_ids_from_schema(doc, topic_id):
        schema = getattr(doc, "schema", None)
        snapshot = getattr(schema, "schema", None) if schema else None
        if not isinstance(snapshot, dict):
            return []
        ids = []
        for group in snapshot.get("groups") or []:
            for field in group.get("fields") or []:
                cda = field.get("cdaOption")
                if cda in (None, "", -1, "-1"):
                    continue
                try:
                    if int(cda) != topic_id:
                        continue
                except (TypeError, ValueError):
                    continue
                pk = field.get("pk")
                if pk:
                    ids.append(pk)
        return ids

    @staticmethod
    def topics_for_documents(docs, iss_by_doc=None):
        from directory.models import ParaclinicInputField
        from directions.models import ParaclinicResult

        topic_id = Documents.topic_cda_id()
        if not topic_id or not docs:
            return {}
        iss_by_doc = iss_by_doc or {}
        topic_by_doc = {}
        field_ids = list(ParaclinicInputField.objects.filter(cda_option_id=topic_id).values_list("pk", flat=True))
        snapshot_field_ids_by_doc = {}
        for doc in docs:
            extra = Documents._topic_field_ids_from_schema(doc, topic_id)
            if extra:
                snapshot_field_ids_by_doc[doc.pk] = extra
                field_ids = list({*field_ids, *extra})
        iss_ids = [iss.pk for iss in iss_by_doc.values()]
        if field_ids and iss_ids:
            rows = ParaclinicResult.objects.filter(issledovaniye_id__in=iss_ids, field_id__in=field_ids).exclude(value="").select_related("issledovaniye")
            for row in rows:
                value = (row.value or "").strip()
                if value:
                    topic_by_doc.setdefault(row.issledovaniye.document_id, value)
        for doc in docs:
            if doc.pk in topic_by_doc:
                continue
            body = doc.body_values if isinstance(doc.body_values, dict) else {}
            ids = list(dict.fromkeys([*snapshot_field_ids_by_doc.get(doc.pk, []), *field_ids]))
            for field_id in ids:
                value = body.get(str(field_id), body.get(field_id))
                if value and str(value).strip():
                    topic_by_doc[doc.pk] = str(value).strip()
                    break
        return topic_by_doc

    def topic_value(self, iss=None, research=None):
        if research is not None:
            return Documents.topic_from_research(research)
        iss_by_doc = {self.pk: iss} if iss else {}
        return Documents.topics_for_documents([self], iss_by_doc).get(self.pk, "")

    @property
    def json(self):
        type_title = self.type_document.title if self.type_document else "Документ"
        group = self.type_document.group_document if self.type_document else None
        return {
            "id": self.id,
            "title": self.list_title(),
            "typeId": self.type_document_id,
            "typeTitle": type_title,
            "groupId": group.pk if group else None,
            "groupTitle": group.title if group else "",
            "confirmed": bool(self.time_confirm),
            "isHidden": self.is_hidden,
        }

    @classmethod
    def _who_group_names(cls, who, *group_names):
        names = [name for name in group_names if name]
        if not who or not names:
            return set()
        user = getattr(who, "user", None)
        if user and getattr(user, "is_superuser", False):
            return set(names)
        if not user:
            return set()
        return set(user.groups.filter(name__in=names).values_list("name", flat=True))

    @classmethod
    def can_view_all_hidden(cls, who):
        if not who:
            return False
        user = getattr(who, "user", None)
        if user and getattr(user, "is_superuser", False):
            return True
        return who.has_group(cls.HIDDEN_DOCS_GROUP)

    @classmethod
    def document_is_confirmed(cls, obj, iss=None):
        if obj.time_confirm:
            return True
        if iss is not None:
            return bool(iss.time_confirmation)
        from directions.models import Issledovaniya

        return Issledovaniya.objects.filter(document=obj, time_confirmation__isnull=False).exists()

    @staticmethod
    def confirmation_text(obj, iss=None):
        from laboratory.utils import strdatetimeru

        doctor = None
        moment = None
        if iss is not None and getattr(iss, "time_confirmation", None) and getattr(iss, "doc_confirmation_id", None):
            doctor = iss.doc_confirmation
            moment = iss.time_confirmation
        elif obj.time_confirm and obj.who_confirm_id:
            doctor = obj.who_confirm
            moment = obj.time_confirm
        if not doctor or not moment:
            return ""
        return f"{doctor}, {strdatetimeru(moment)}"

    @staticmethod
    def linked_cases(obj):
        if not obj:
            return []
        prefetched = getattr(obj, "_prefetched_objects_cache", {}).get("case_links")
        links = prefetched if prefetched is not None else obj.case_links.select_related("document_case").order_by("pk")
        cases = []
        seen = set()
        for link in links:
            case = link.document_case if link.document_case_id else None
            if case and case.pk not in seen:
                seen.add(case.pk)
                cases.append(case)
        if obj.document_case_id and obj.document_case_id not in seen:
            case = obj.document_case if getattr(obj, "document_case", None) is not None and obj.document_case.pk == obj.document_case_id else None
            if case is None:
                case = DocumentCase.objects.filter(pk=obj.document_case_id).first()
            if case:
                cases.append(case)
        return cases

    @staticmethod
    def linked_blocks(obj):
        if not obj:
            return []
        prefetched = getattr(obj, "_prefetched_objects_cache", {}).get("block_links")
        links = prefetched if prefetched is not None else obj.block_links.select_related("block").order_by("pk")
        blocks = []
        seen = set()
        for link in links:
            block = link.block if link.block_id else None
            if block and block.pk not in seen:
                seen.add(block.pk)
                blocks.append(block)
        if obj.block_id and obj.block_id not in seen:
            block = obj.block if getattr(obj, "block", None) is not None and obj.block.pk == obj.block_id else None
            if block is None:
                block = DocumentBlock.objects.filter(pk=obj.block_id).first()
            if block:
                blocks.append(block)
        return blocks

    @staticmethod
    def ensure_case_link(obj):
        if obj.document_case_id and not obj.case_links.filter(document_case_id=obj.document_case_id).exists():
            DocumentCaseLink.objects.get_or_create(document=obj, document_case_id=obj.document_case_id)

    @staticmethod
    def ensure_block_link(obj):
        if obj.block_id and not obj.block_links.filter(block_id=obj.block_id).exists():
            DocumentBlockLink.objects.get_or_create(document=obj, block_id=obj.block_id)

    @staticmethod
    def sync_case_field(obj):
        ids = list(obj.case_links.order_by("pk").values_list("document_case_id", flat=True))
        if obj.document_case_id in ids:
            return
        obj.document_case_id = ids[0] if ids else None
        Documents.objects.filter(pk=obj.pk).update(document_case_id=obj.document_case_id)

    @staticmethod
    def sync_block_field(obj):
        ids = list(obj.block_links.order_by("pk").values_list("block_id", flat=True))
        if obj.block_id in ids:
            return
        obj.block_id = ids[0] if ids else None
        Documents.objects.filter(pk=obj.pk).update(block_id=obj.block_id)

    @classmethod
    def can_see_document(cls, obj, who):
        if who and any(block.is_public() for block in Documents.linked_blocks(obj)):
            return True
        cases = Documents.linked_cases(obj)
        if cases and not any(case.can_access(who) for case in cases):
            if not DocumentEmployeeApprove.is_assigned(obj, who):
                return False
        if not obj.is_hidden:
            return True
        return cls.can_view_all_hidden(who)

    @staticmethod
    def _public_block(obj):
        block = obj.block if getattr(obj, "block", None) is not None else None
        if block is None or block.pk != obj.block_id:
            block = DocumentBlock.objects.filter(pk=obj.block_id).first()
        return bool(block and block.is_public())

    @staticmethod
    def case_access_q(who, related=""):
        case_field = f"{related}document_case"
        link_field = f"{related}case_links__document_case"
        if DocumentCase._is_superuser(who):
            return models.Q()
        owner = "pk" if not related else f"{related[:-2]}_id"
        has_link = DocumentCaseLink.objects.filter(document_id=models.OuterRef(owner))
        no_case = models.Q(**{f"{case_field}__isnull": True}) & ~models.Exists(has_link)
        if not who:
            return no_case
        accessible = (
            models.Q(**{f"{link_field}__who_create_id": who.pk})
            | models.Q(**{f"{link_field}__access_rows__doctor_id": who.pk})
            | models.Q(**{f"{case_field}__who_create_id": who.pk})
            | models.Q(**{f"{case_field}__access_rows__doctor_id": who.pk})
        )
        return no_case | accessible

    @classmethod
    def can_set_hidden(cls, obj, who, confirmed=None):
        if not who:
            return False
        user = getattr(who, "user", None)
        if user and getattr(user, "is_superuser", False):
            return True
        if obj.who_create_id == getattr(who, "pk", None):
            if confirmed is None:
                confirmed = cls.document_is_confirmed(obj)
            if not confirmed:
                return True
        return who.has_group(cls.HIDDEN_DOCS_GROUP)

    @classmethod
    def can_reset_confirm(cls, who):
        if not who:
            return False
        user = getattr(who, "user", None)
        if user and getattr(user, "is_superuser", False):
            return True
        return who.has_group(cls.RESET_GROUP)

    @staticmethod
    def find_by_id(pk, who=None):
        try:
            pk = int(pk)
        except (TypeError, ValueError):
            return {"ok": False, "message": "Введите номер документа"}
        if pk <= 0:
            return {"ok": False, "message": "Введите номер документа"}
        obj = Documents.objects.select_related("type_document", "type_document__group_document", "document_case").filter(pk=pk).first()
        if not obj or not Documents.can_see_document(obj, who):
            return {"ok": False, "message": "Документ не найден"}
        payload = obj.json
        payload["ok"] = True
        return payload

    @staticmethod
    def _search_dates(date_from, date_to):
        from datetime import datetime

        def parse(value):
            if value in (None, ""):
                return None
            try:
                return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()
            except ValueError:
                return False

        start = parse(date_from)
        end = parse(date_to)
        if start is False or end is False:
            return None, None, "Некорректная дата"
        if start and end and start > end:
            return None, None, "Дата начала позже даты окончания"
        return start, end, None

    @staticmethod
    def search(query, by_number=False, by_text=False, by_case=False, who=None, limit=50, date_from=None, date_to=None):
        query = (query or "").strip()
        if not query:
            return {"ok": False, "message": "Введите запрос"}
        if not by_number and not by_text and not by_case:
            return {"ok": False, "message": "Выберите способ поиска"}
        start, end, date_error = Documents._search_dates(date_from, date_to)
        if date_error:
            return {"ok": False, "message": date_error}
        found_ids = []
        seen = set()

        number_only = bool(by_number) and not bool(by_text) and not bool(by_case)
        case_only = bool(by_case) and not bool(by_number) and not bool(by_text)

        def in_period(qs, field="create_at"):
            if number_only:
                return qs
            if start:
                qs = qs.filter(**{f"{field}__date__gte": start})
            if end:
                qs = qs.filter(**{f"{field}__date__lte": end})
            return qs

        def add_ids(pks):
            for pk in pks:
                if not pk or pk in seen:
                    continue
                seen.add(pk)
                found_ids.append(pk)
                if len(found_ids) >= limit:
                    return True
            return False

        if case_only:
            import re

            from document_management.sql_func import search_cases_by_topic

            found = search_cases_by_topic(re.escape(query), limit)
            case_ids = [row.id for row in found]
            cases = [case for case in DocumentCase.objects.filter(pk__in=case_ids) if case.can_access(who)]
            order = {pk: index for index, pk in enumerate(case_ids)}
            cases.sort(key=lambda case: order.get(case.pk, 0))
            if not cases:
                return {"ok": False, "message": "Дело не найдено"}
            rows = []
            for case in cases:
                rows.append({"id": case.pk, "title": Documents.title_with_id(case.pk, case.topic or ""), "isCase": True})
            message = rows[0]["title"] if len(rows) == 1 else f"Найдено {len(rows)}"
            return {"ok": True, "result": rows, "cases": True, "message": message}
        elif by_number:
            digits = "".join(ch for ch in query if ch.isdigit())
            if not digits:
                if not by_text:
                    return {"ok": False, "message": "Введите номер документа"}
            else:
                try:
                    pk = int(digits)
                except ValueError:
                    pk = 0
                if pk > 0:
                    add_ids(in_period(Documents.objects.filter(pk=pk)).values_list("pk", flat=True))
                if len(found_ids) < limit:
                    numbered = in_period(Documents.objects.filter(number_registration=digits))
                    add_ids(numbered.values_list("pk", flat=True)[:limit])
        if by_text and len(found_ids) < limit:
            from directory.models import ParaclinicInputField
            from directions.models import Issledovaniya, ParaclinicResult

            titled = in_period(Documents.objects.filter(type_document__title__icontains=query)).order_by("-pk")
            add_ids(titled.values_list("pk", flat=True)[:limit])
            if len(found_ids) < limit:
                recent_ids = in_period(DocumentRecent.objects.filter(topic__icontains=query), "document__create_at").order_by("-opened_at").values_list("document_id", flat=True)[:limit]
                add_ids(recent_ids)
            topic_id = Documents.topic_cda_id()
            if topic_id and len(found_ids) < limit:
                field_ids = list(ParaclinicInputField.objects.filter(cda_option_id=topic_id).values_list("pk", flat=True))
                if field_ids:
                    iss_ids = list(
                        ParaclinicResult.objects.filter(field_id__in=field_ids, value__icontains=query).values_list(
                            "issledovaniye_id",
                            flat=True,
                        )[:limit]
                    )
                    if iss_ids:
                        doc_ids = in_period(
                            Issledovaniya.objects.filter(pk__in=iss_ids, document_id__isnull=False),
                            "document__create_at",
                        ).values_list("document_id", flat=True)
                        add_ids(doc_ids)
        if not found_ids:
            return {"ok": False, "message": "Документ не найден"}
        visible = Documents.objects.select_related("type_document", "document_case").filter(pk__in=found_ids[:limit])
        docs = [doc for doc in visible if Documents.can_see_document(doc, who)]
        order = {pk: index for index, pk in enumerate(found_ids)}
        docs.sort(key=lambda doc: order.get(doc.pk, 0))
        if not docs:
            return {"ok": False, "message": "Документ не найден"}
        from directions.models import Issledovaniya

        iss_by_doc = {}
        for iss in Issledovaniya.objects.filter(document_id__in=[doc.pk for doc in docs]).order_by("pk"):
            iss_by_doc.setdefault(iss.document_id, iss)
        topic_by_doc = Documents.topics_for_documents(docs, iss_by_doc)
        rows = []
        for doc in docs:
            title = doc.list_title(iss_by_doc.get(doc.pk), topic_by_doc.get(doc.pk))
            rows.append({"id": doc.pk, "title": Documents.title_with_id(doc.pk, title)})
        message = rows[0]["title"] if len(rows) == 1 else f"Найдено {len(rows)}"
        return {"ok": True, "result": rows, "message": message}

    @classmethod
    def log_action(cls, pk, log_type, who, extra=None):
        DocumentManagerLog.write(log_type, who, extra or {}, document_id=pk)

    @staticmethod
    def participant_cases(who, date_from=None, date_to=None):
        if not who:
            return {"ok": True, "result": []}
        start, end, date_error = Documents._search_dates(date_from, date_to)
        if date_error:
            return {"ok": False, "message": date_error, "result": []}
        qs = DocumentCase.objects.filter(access_rows__doctor_id=who.pk).distinct()
        if start:
            qs = qs.filter(created_at__date__gte=start)
        if end:
            qs = qs.filter(created_at__date__lte=end)
        cases = list(qs.order_by("-created_at", "-pk"))
        case_ids = [case.pk for case in cases]
        mine_ids = set(UserMyCase.objects.filter(doctor=who, document_case_id__in=case_ids).values_list("document_case_id", flat=True))
        favorite_ids = set(UserFavoriteCase.objects.filter(doctor=who, document_case_id__in=case_ids).values_list("document_case_id", flat=True))
        favorite_ids.update(
            UserFavoriteCase.objects.filter(doctor=who, document_case__isnull=True, document__document_case_id__in=case_ids).values_list("document__document_case_id", flat=True)
        )
        rows = []
        for case in cases:
            is_mine = case.who_create_id == who.pk
            rows.append(
                {
                    "id": case.pk,
                    "title": Documents.title_with_id(case.pk, case.topic or ""),
                    "isCase": True,
                    "isMine": is_mine,
                    "inMyCases": is_mine or case.pk in mine_ids,
                    "isFavorite": case.pk in favorite_ids,
                }
            )
        return {"ok": True, "result": rows}

    @staticmethod
    def get_list(type_id=None, group_id=None, role_filter=None, who=None, hidden=False, my_cases=False):
        qs = (
            Documents.objects.select_related("type_document", "type_document__group_document", "schema", "document_case")
            .prefetch_related(models.Prefetch("case_links", queryset=DocumentCaseLink.objects.select_related("document_case").order_by("pk")))
            .order_by("-pk")
        )
        if role_filter == "toBeAgreed":
            if not who:
                return []
            qs = qs.filter(
                is_hidden=False,
                documentemployeeapprove__employee_id=who.pk,
                documentemployeeapprove__is_cancel=False,
                documentemployeeapprove__time_approve__isnull=True,
            )
        else:
            qs = qs.filter(Documents.case_access_q(who))
            if my_cases:
                if not who:
                    return []
                qs = qs.filter(
                    models.Q(case_links__document_case__who_create_id=who.pk)
                    | models.Q(case_links__document_case__user_my_cases__doctor_id=who.pk)
                    | models.Q(document_case__who_create_id=who.pk)
                    | models.Q(document_case__user_my_cases__doctor_id=who.pk)
                )
            elif role_filter == "toReview":
                if not who:
                    return []
                qs = qs.filter(pk__in=DocumentReview.pending_document_ids(who))
            else:
                if type_id not in (None, "", -1, "-1"):
                    qs = qs.filter(type_document_id=type_id)
                elif group_id in (-1, "-1"):
                    qs = qs.filter(models.Q(type_document__isnull=True) | models.Q(type_document__group_document__isnull=True))
                elif group_id not in (None, "", 0, "0"):
                    qs = qs.filter(type_document__group_document_id=group_id)
                if role_filter == "created" and who:
                    qs = qs.filter(who_create=who)
            if hidden:
                if not Documents.can_view_all_hidden(who):
                    return []
                qs = qs.filter(is_hidden=True)
            else:
                qs = qs.filter(is_hidden=False)
        docs = list(qs.distinct())
        if not docs:
            return []

        from directions.models import Issledovaniya

        iss_by_doc = {}
        for iss in Issledovaniya.objects.filter(document_id__in=[doc.pk for doc in docs]).order_by("pk"):
            iss_by_doc.setdefault(iss.document_id, iss)

        topic_by_doc = {}
        favorite_case_ids = set()
        mine_case_ids = set()
        if not my_cases:
            topic_by_doc = Documents.topics_for_documents(docs, iss_by_doc)
        elif who:
            case_ids = []
            for doc in docs:
                case_ids.extend(case.pk for case in Documents.linked_cases(doc))
            if case_ids:
                favorite_case_ids = set(UserFavoriteCase.objects.filter(doctor=who, document_case_id__in=case_ids).values_list("document_case_id", flat=True))
                favorite_case_ids.update(
                    UserFavoriteCase.objects.filter(doctor=who, document_case__isnull=True, document__document_case_id__in=case_ids).values_list(
                        "document__document_case_id",
                        flat=True,
                    )
                )
                mine_case_ids = set(UserMyCase.objects.filter(doctor=who, document_case_id__in=case_ids).values_list("document_case_id", flat=True))

        result = []
        for doc in docs:
            payload = doc.json
            if my_cases:
                for case in Documents.linked_cases(doc):
                    if case.who_create_id != getattr(who, "pk", None) and case.pk not in mine_case_ids:
                        continue
                    row = dict(payload)
                    case_topic = (case.topic or "").strip()
                    row["title"] = f"{case_topic} {doc.pk}".strip() if case_topic else str(doc.pk)
                    row["caseId"] = case.pk
                    row["isFavorite"] = case.pk in favorite_case_ids
                    result.append(row)
                continue
            else:
                topic = (topic_by_doc.get(doc.pk) or "").strip()
                type_doc = doc.type_document.title if doc.type_document else ""
                title = " ".join(part for part in (topic, type_doc) if part)
                payload["title"] = Documents.title_with_id(doc.pk, title)
            result.append(payload)
        return result

    @staticmethod
    def create_document(type_id, who_create, case_id=None, topic=None):
        from directions.models import Issledovaniya

        case = None
        if case_id not in (None, "", -1, "-1"):
            case = TypeCases.objects.filter(pk=case_id).select_related("default_type_document").first()
            if not case:
                return {"ok": False, "message": "Вид дела не найден"}
            topic_text = (topic or "").strip()
            if not topic_text:
                return {"ok": False, "message": "Укажите тему дела"}
            if case.default_type_document_id:
                type_id = case.default_type_document_id
            elif type_id in (None, "", -1, "-1"):
                with transaction.atomic():
                    document_case = DocumentCase.objects.create(type_case=case, topic=topic_text, who_create=who_create)
                return {
                    "ok": True,
                    "id": None,
                    "caseId": document_case.pk,
                    "emptyCase": True,
                    "title": topic_text,
                }

        if type_id in (None, "", -1, "-1"):
            return {"ok": False, "message": "Выберите вид документа"}
        type_doc = TypeDocuments.objects.select_related("layout_template", "monitoring").filter(pk=type_id).first()
        if not type_doc:
            return {"ok": False, "message": "Вид документа не найден"}
        if not type_doc.can_create(who_create):
            return {"ok": False, "message": "Нет прав на создание этого вида документа"}
        schema = TypeDocumentsSchema.objects.filter(type_document=type_doc).order_by("-created_at", "-pk").first()
        if schema is None:
            schema = TypeDocumentsSchema.get_or_create_for_type(type_doc)
        with transaction.atomic():
            document_case = None
            if case is not None:
                document_case = DocumentCase.objects.create(type_case=case, topic=(topic or "").strip(), who_create=who_create)
            plan = None
            if type_doc.is_plan:
                plan = Plans.objects.create(title=type_doc.title or "", user_created=True, who_create=who_create)
            obj = Documents.objects.create(
                type_document=type_doc,
                who_create=who_create,
                body_values={},
                schema=schema,
                type_case=case,
                document_case=document_case,
                plan=plan,
            )
            if document_case:
                DocumentCaseLink.objects.create(document=obj, document_case=document_case)
            Issledovaniya.objects.create(
                document=obj,
                research=type_doc.form_research(),
                creator=who_create,
            )
        type_title = type_doc.title or "Документ"
        title = f"{type_title} №{obj.pk}"
        Documents.log_action(
            obj.pk,
            Documents.LOG_CREATE,
            who_create,
            {"title": title, "typeId": type_doc.pk, "typeTitle": type_doc.title or ""},
        )
        return {"ok": True, "id": obj.pk, "title": title, "caseId": document_case.pk if document_case else None}

    def save_user_plan(self, who, title, groups, indicators, tests, with_confirm=False):
        from django.utils import timezone

        if not Documents.can_see_document(self, who):
            return {"ok": False, "message": "Документ не найден"}
        if not self.plan_id or not self.plan or not self.plan.user_created:
            return {"ok": False, "message": "План не найден"}
        if Documents.document_is_confirmed(self):
            return {"ok": False, "message": "Документ подтверждён"}
        result = Plans.save_plan(self.plan_id, title, groups, indicators, tests, allow_user=True, who=who)
        if not result.get("ok"):
            return result
        if with_confirm:
            now = timezone.now()
            self.time_confirm = now
            self.who_confirm = who
            self.save(update_fields=["time_confirm", "who_confirm"])
            iss = self.get_issledovaniye()
            if iss:
                iss.doc_confirmation = who
                iss.time_confirmation = now
                iss.save(update_fields=["doc_confirmation", "time_confirmation"])
        Documents.log_action(
            self.pk,
            Documents.LOG_CONFIRM if with_confirm else Documents.LOG_SAVE,
            who,
            {
                "title": self.json.get("title") if isinstance(self.json, dict) else "",
                "typeId": self.type_document_id,
                "typeTitle": self.type_document.title if self.type_document else "",
                "planId": self.plan_id,
                "confirmed": bool(with_confirm),
            },
        )
        return {"ok": True, "id": self.pk, "planId": self.plan_id, "confirmed": bool(self.time_confirm)}

    def get_issledovaniye(self):
        from directions.models import Issledovaniya

        prefetched = getattr(self, "_prefetched_objects_cache", {}).get("issledovaniya_set")
        if prefetched is not None:
            iss = prefetched[0] if prefetched else None
        else:
            iss = Issledovaniya.objects.filter(document=self).only("pk", "time_confirmation", "document_id", "research_id").order_by("pk").first()
        if iss:
            return iss
        type_doc = self.type_document
        return Issledovaniya.objects.create(
            document=self,
            research=type_doc.form_research() if type_doc else None,
            creator=self.who_create,
        )

    def _schema_snapshot(self):
        if self.schema_id:
            snapshot = getattr(self.schema, "schema", None)
            if isinstance(snapshot, dict) and snapshot.get("groups") is not None:
                return snapshot
        if not self.type_document_id:
            return None
        latest = TypeDocumentsSchema.objects.filter(type_document_id=self.type_document_id).order_by("-created_at", "-pk").only("pk", "schema").first()
        if latest and isinstance(latest.schema, dict) and latest.schema.get("groups") is not None:
            if not self.schema_id:
                Documents.objects.filter(pk=self.pk, schema_id__isnull=True).update(schema_id=latest.pk)
                self.schema_id = latest.pk
            return latest.schema
        return TypeDocumentsSchema.dump_from_type(self.type_document)

    def build_research(self, iss=None):
        snapshot = self._schema_snapshot()
        if not snapshot:
            return None
        return self._research_from_schema(snapshot, iss)

    def _research_from_schema(self, snapshot, iss=None):
        from directions.models import ParaclinicResult, ParaclinicResultFile

        saved = self.body_values if isinstance(self.body_values, dict) else {}
        result_by_field = {}
        files_by_result = {}
        if iss:
            rows = list(ParaclinicResult.objects.filter(issledovaniye_id=iss.pk).only("pk", "field_id", "value", "field_type"))
            result_by_field = {row.field_id: row for row in rows}
            type42_ids = set()
            for group in snapshot.get("groups") or []:
                for field in group.get("fields") or []:
                    if field.get("field_type") == 42:
                        type42_ids.add(field.get("pk"))
            result_pks = [result_by_field[field_id].pk for field_id in type42_ids if field_id in result_by_field]
            if result_pks:
                for row in ParaclinicResultFile.objects.filter(result_id__in=result_pks).order_by("pk"):
                    files_by_result.setdefault(row.result_id, []).append(row)
        confirmed = bool(iss.time_confirmation) if iss else False
        groups = []
        for group in snapshot.get("groups") or []:
            if group.get("hide"):
                continue
            g = {
                "pk": group.get("pk"),
                "order": group.get("order"),
                "title": group.get("title") if group.get("show_title") else "",
                "show_title": group.get("show_title"),
                "hide": group.get("hide"),
                "display_hidden": False,
                "fields": [],
                "visibility": group.get("visibility") or "",
                "fieldsInline": group.get("fieldsInline"),
                "cdaOption": group.get("cdaOption", -1),
            }
            for field in group.get("fields") or []:
                if field.get("hide") or field.get("field_type") == LAYOUT_TEMPLATE_FIELD_TYPE:
                    continue
                field_pk = field.get("pk")
                values_to_input = field.get("values_to_input") or []
                if not isinstance(values_to_input, list):
                    values_to_input = []
                result_field = result_by_field.get(field_pk)
                field_type = field.get("field_type") or 0
                if confirmed and result_field is not None and result_field.field_type is not None:
                    field_type = result_field.field_type
                if field.get("required") and field_type in [10, 12] and "- Не выбрано" not in values_to_input:
                    values_to_input = ["- Не выбрано", *values_to_input]
                default_value = field.get("default_value") or ""
                if field_type in [3, 11, 13, 14, 30, 42, 44, ADDRESSEE_FIELD_TYPE, MENTEE_FIELD_TYPE, MENTOR_FIELD_TYPE]:
                    default_value = ""
                key = str(field_pk)
                if result_field:
                    value = result_field.value
                elif key in saved:
                    value = saved[key]
                else:
                    value = saved.get(field_pk, default_value)
                if field_type == ADDRESSEE_FIELD_TYPE and not value:
                    value = "[]"
                if field_type == PLAN_FIELD_TYPE and not value:
                    value = "[]"
                file_settings = field.get("file_settings")
                files = []
                if field_type == 42:
                    value = ""
                    if result_field:
                        files = [ParaclinicResultFile.serialize(row) for row in files_by_result.get(result_field.pk, [])]
                g["fields"].append(
                    {
                        "pk": field_pk,
                        "order": field.get("order"),
                        "lines": field.get("lines"),
                        "title": field.get("short_title") or field.get("title") or "",
                        "hide": field.get("hide"),
                        "values_to_input": values_to_input,
                        "value": value if value is not None else "",
                        "field_type": field_type,
                        "can_edit": field.get("can_edit"),
                        "default_value": field.get("default_value") or "",
                        "visibility": field.get("visibility") or "",
                        "required": field.get("required"),
                        "helper": field.get("helper") or "",
                        "controlParam": field.get("controlParam") or "",
                        "not_edit": field.get("not_edit"),
                        "operator_enter_param": field.get("operator_enter_param"),
                        "deniedGroup": field.get("deniedGroup") or "",
                        "isDiagTable": field.get("isDiagTable"),
                        "file_settings": file_settings,
                        "files": files,
                        "cdaOption": field.get("cdaOption", -1),
                    }
                )
            if not g["fields"]:
                continue
            groups.append(g)
        return {
            "pk": snapshot.get("pk"),
            "title": snapshot.get("title"),
            "version": 0,
            "wide_headers": bool(snapshot.get("wide_headers")),
            "groups": groups,
            "is_gistology": False,
            "show_more_services": False,
        }

    @staticmethod
    def get_details(pk, who=None):
        from django.db.models import Prefetch
        from directions.models import Issledovaniya

        obj = (
            Documents.objects.select_related(
                "schema",
                "type_document",
                "type_document__layout_template",
                "type_document__group_document",
                "document_case",
                "document_case__who_create",
                "document_case__who_close",
                "block",
                "plan",
                "who_confirm",
                "who_confirm__podrazdeleniye",
            )
            .prefetch_related(
                Prefetch(
                    "case_links",
                    queryset=DocumentCaseLink.objects.select_related("document_case", "document_case__who_create", "document_case__who_close")
                    .prefetch_related(Prefetch("document_case__access_rows", queryset=DocumentCaseAccess.objects.select_related("doctor", "doctor__podrazdeleniye")))
                    .order_by("pk"),
                ),
                Prefetch(
                    "block_links",
                    queryset=DocumentBlockLink.objects.select_related("block").order_by("pk"),
                ),
                Prefetch(
                    "document_case__access_rows",
                    queryset=DocumentCaseAccess.objects.select_related("doctor", "doctor__podrazdeleniye"),
                ),
                Prefetch(
                    "issledovaniya_set",
                    queryset=Issledovaniya.objects.only("pk", "time_confirmation", "document_id", "research_id").order_by("pk"),
                ),
                Prefetch(
                    "documentemployeeapprove_set",
                    queryset=DocumentEmployeeApprove.objects.filter(is_cancel=False, employee_id__isnull=False).select_related("employee", "employee__podrazdeleniye", "who_create"),
                ),
            )
            .filter(pk=pk)
            .first()
        )
        if not obj:
            return {"ok": False, "message": "Документ не найден"}
        if not Documents.can_see_document(obj, who):
            return {"ok": False, "message": "Документ не найден"}
        group_names = Documents._who_group_names(who, Documents.HIDDEN_DOCS_GROUP, Documents.RESET_GROUP)
        iss = obj.get_issledovaniye()
        research = obj.build_research(iss)
        payload = obj.json
        payload["ok"] = True
        payload["issPk"] = iss.pk if iss else None
        payload["research"] = research
        payload["testId"] = obj.type_document.test_id if obj.type_document_id and obj.type_document and obj.type_document.test_id else None
        payload["confirmed"] = bool(iss and iss.time_confirmation) or bool(obj.time_confirm)
        payload["whoConfirmed"] = Documents.confirmation_text(obj, iss) if payload["confirmed"] else ""
        can_hide = Documents.HIDDEN_DOCS_GROUP in group_names
        if not can_hide and who and obj.who_create_id == getattr(who, "pk", None) and not payload["confirmed"]:
            can_hide = True
        payload["canHide"] = can_hide
        payload["canReset"] = Documents.RESET_GROUP in group_names
        payload["isCase"] = bool(obj.type_case_id)
        cases = Documents.linked_cases(obj)
        case_options = DocumentCase.available_options(who)
        known_cases = {item["id"] for item in case_options}
        for case in cases:
            if case.pk not in known_cases:
                case_options.insert(0, {"id": case.pk, "label": f"{case.pk} {case.topic}".strip()})
                known_cases.add(case.pk)
        payload["caseIds"] = [case.pk for case in cases]
        payload["case"] = cases[0].as_json(who) if len(cases) == 1 else None
        payload["availableCases"] = case_options
        blocks = Documents.linked_blocks(obj)
        block_options = DocumentBlock.write_options(who)
        known_blocks = {item["id"] for item in block_options}
        for block in blocks:
            if block.pk not in known_blocks:
                block_options.insert(0, {"id": block.pk, "label": block.title or f"Блок {block.pk}"})
                known_blocks.add(block.pk)
        payload["blockIds"] = [block.pk for block in blocks]
        payload["blockId"] = obj.block_id
        payload["availableBlocks"] = block_options
        locked_block_ids = {block.pk for block in blocks if not block.can_write(who)}
        payload["canChangeBlock"] = any(block.can_write(who) for block in blocks) or any(item["id"] not in locked_block_ids for item in block_options)
        payload["approvers"] = DocumentEmployeeApprove.approvers_json(obj)
        payload["approvalStatuses"] = DocumentEmployeeApprove.statuses_json(obj)
        payload["canEditApprovers"] = DocumentEmployeeApprove.can_edit(obj, who)
        payload["canCancelApproval"] = DocumentEmployeeApprove.can_cancel(who)
        payload["isFavorite"] = record_favorite_status(who, obj)
        payload["hasPrintTemplate"] = bool(obj.type_document_id and obj.type_document.print_docx)
        plan = obj.plan if obj.plan_id and obj.plan and obj.plan.user_created else None
        payload["isPlan"] = bool(plan) or bool(obj.type_document_id and obj.type_document.is_plan)
        if plan:
            plan_details = Plans.get_details(plan.pk, who)
            if plan_details.get("ok"):
                payload["plan"] = plan_details
        payload["reviewedNow"] = DocumentReview.mark_opened(obj, who, iss=iss) if payload["confirmed"] else False
        type_doc = obj.type_document.title if obj.type_document else ""
        DocumentRecent.remember(who, obj, topic=obj.topic_value(iss, research=research), type_doc=type_doc)
        return payload

    @staticmethod
    def save_paraclinic_result(iss_pk, research, with_confirm, visibility_state, who, request_files=None, case_topic=None):
        from django.utils import timezone

        from directory.models import ParaclinicInputField
        from directions.models import Issledovaniya, ParaclinicResult, ParaclinicResultFile

        iss = Issledovaniya.objects.filter(pk=iss_pk, document__isnull=False).select_related("document", "document__document_case", "research").first()
        if not iss or not iss.document:
            return {"ok": False, "message": "Исследование не найдено"}
        if not Documents.can_see_document(iss.document, who):
            return {"ok": False, "message": "Документ не найден"}
        if iss.time_confirmation or iss.document.time_confirm:
            return {"ok": False, "message": "Документ подтверждён"}
        document = iss.document
        case_row = document.document_case if document.document_case_id else None
        next_topic = None
        if case_row is not None and case_topic is not None:
            next_topic = str(case_topic).strip()
        if case_row is not None:
            topic = next_topic if next_topic is not None else (case_row.topic or "").strip()
            if not topic:
                return {"ok": False, "message": "Укажите тему дела"}
        v_g = (visibility_state or {}).get("groups") or {}
        v_f = (visibility_state or {}).get("fields") or {}
        groups = (research or {}).get("groups") or []
        plan_error = Plans.validate_document_fields(groups, iss, who)
        if plan_error:
            return {"ok": False, "message": plan_error}
        files_by_field = {}
        body_values = dict(document.body_values) if isinstance(document.body_values, dict) else {}
        addressee_ids = []
        addressee_seen = set()
        has_addressee_field = False
        with transaction.atomic():
            if case_row is not None and next_topic is not None and case_row.topic != next_topic:
                case_row.topic = next_topic
                case_row.save(update_fields=["topic"])
            for group in groups:
                group_pk = group.get("pk")
                if not v_g.get(str(group_pk), True):
                    ParaclinicResult.objects.filter(issledovaniye=iss, field__group__pk=group_pk).delete()
                    for field in group.get("fields") or []:
                        field_pk = field.get("pk")
                        if field_pk:
                            body_values.pop(str(field_pk), None)
                    continue
                for field in group.get("fields") or []:
                    field_pk = field.get("pk")
                    if not field_pk:
                        continue
                    if not v_f.get(str(field_pk), True):
                        ParaclinicResult.objects.filter(issledovaniye=iss, field__pk=field_pk).delete()
                        body_values.pop(str(field_pk), None)
                        continue
                    value = field.get("value")
                    stored_value = "" if not value else value
                    payload_field_type = field.get("field_type")
                    f = ParaclinicInputField.objects.filter(pk=field_pk).first()
                    if not f or f.field_type == 21:
                        if payload_field_type != 21:
                            body_values[str(field_pk)] = stored_value
                        continue
                    body_values[str(field_pk)] = stored_value
                    f_result = ParaclinicResult.objects.filter(issledovaniye=iss, field=f).first()
                    if not f_result:
                        f_result = ParaclinicResult(issledovaniye=iss, field=f, value="")
                    f_result.value = stored_value
                    f_result.field_type = payload_field_type if payload_field_type is not None else f.field_type
                    field_type = f_result.field_type
                    if field_type == ADDRESSEE_FIELD_TYPE:
                        has_addressee_field = True
                        if isinstance(stored_value, (dict, list)):
                            stored_value = json.dumps(stored_value)
                            f_result.value = stored_value
                            body_values[str(field_pk)] = stored_value
                        for doctor_id in DocumentReview.parse_addressee_ids(value):
                            if doctor_id not in addressee_seen:
                                addressee_seen.add(doctor_id)
                                addressee_ids.append(doctor_id)
                    if field_type in [27, 28, 29, 32, 33, 34, 35, 44, ADDRESSEE_FIELD_TYPE, MENTEE_FIELD_TYPE, MENTOR_FIELD_TYPE, PLAN_FIELD_TYPE]:
                        if isinstance(value, (dict, list)):
                            val = value
                        else:
                            try:
                                val = json.loads(value)
                            except Exception:
                                val = []
                        f_result.value_json = val
                    f_result.save()
                    if field_type == 42:
                        files_by_field[field_pk] = ParaclinicResultFile.sync_field_files(
                            f_result=f_result,
                            payload_files=field.get("files") or [],
                            uploaded_files=request_files or {},
                            field_pk=field_pk,
                        )
            iss.doc_save = who
            iss.time_save = timezone.now()
            document.body_values = body_values
            update_fields = ["body_values"]
            if with_confirm:
                now = timezone.now()
                iss.doc_confirmation = who
                iss.time_confirmation = now
                document.time_confirm = now
                document.who_confirm = who
                update_fields.extend(["time_confirm", "who_confirm"])
            document.save(update_fields=update_fields)
            iss.save()
            if has_addressee_field:
                DocumentReview.sync_for_document(document, addressee_ids)
            DocumentRecent.objects.filter(document=document).update(topic=document.topic_value(iss, research=research))
        Documents.log_action(
            document.pk,
            Documents.LOG_CONFIRM if with_confirm else Documents.LOG_SAVE,
            who,
            {
                "title": document.json.get("title") if isinstance(document.json, dict) else "",
                "typeId": document.type_document_id,
                "typeTitle": document.type_document.title if document.type_document else "",
                "issPk": iss.pk,
                "confirmed": bool(iss.time_confirmation),
            },
        )
        return {
            "ok": True,
            "id": iss.document_id,
            "issPk": iss.pk,
            "confirmed": bool(iss.time_confirmation),
            "whoConfirmed": Documents.confirmation_text(document, iss) if iss.time_confirmation else "",
            "files_by_field": files_by_field,
        }

    @staticmethod
    def confirm_reset(pk, who=None):
        from directions.models import Issledovaniya

        obj = Documents.objects.filter(pk=pk).first()
        if not obj:
            return {"ok": False, "message": "Документ не найден"}
        if not Documents.can_reset_confirm(who):
            return {"ok": False, "message": "Нет прав"}
        if not Documents.can_see_document(obj, who):
            return {"ok": False, "message": "Документ не найден"}
        iss = Issledovaniya.objects.filter(document=obj).order_by("pk").first()
        if (not iss or not iss.time_confirmation) and not obj.time_confirm:
            return {"ok": False, "message": "Документ не подтверждён"}
        if iss:
            iss.time_confirmation = None
            iss.doc_confirmation = None
            iss.executor_confirmation = None
            iss.save(update_fields=["time_confirmation", "doc_confirmation", "executor_confirmation"])
        obj.time_confirm = None
        obj.who_confirm = None
        obj.save(update_fields=["time_confirm", "who_confirm"])
        Documents.log_action(
            obj.pk,
            Documents.LOG_RESET,
            who,
            {
                "title": obj.json.get("title") if isinstance(obj.json, dict) else "",
                "typeId": obj.type_document_id,
                "typeTitle": obj.type_document.title if obj.type_document else "",
            },
        )
        return {"ok": True, "id": obj.pk, "confirmed": False}

    @staticmethod
    def set_hidden(pk, hidden, who):
        obj = Documents.objects.filter(pk=pk).first()
        if not obj:
            return {"ok": False, "message": "Документ не найден"}
        if not Documents.can_see_document(obj, who):
            return {"ok": False, "message": "Документ не найден"}
        if not Documents.can_set_hidden(obj, who):
            return {"ok": False, "message": "Нет прав"}
        obj.is_hidden = bool(hidden)
        obj.save(update_fields=["is_hidden"])
        Documents.log_action(
            obj.pk,
            Documents.LOG_HIDE if obj.is_hidden else Documents.LOG_SHOW,
            who,
            {
                "title": obj.json.get("title") if isinstance(obj.json, dict) else "",
                "typeId": obj.type_document_id,
                "typeTitle": obj.type_document.title if obj.type_document else "",
                "isHidden": obj.is_hidden,
            },
        )
        return {"ok": True, "id": obj.pk, "isHidden": obj.is_hidden}

    @classmethod
    def get_history_for_document(cls, pk):
        from laboratory.utils import strdatetime
        from slog.models import Log

        try:
            pk = int(pk)
        except (TypeError, ValueError):
            return []
        if pk < 1 or not cls.objects.filter(pk=pk).exists():
            return []

        body_labels = {
            "title": "Заголовок",
            "typeId": "Код вида",
            "typeTitle": "Вид",
            "issPk": "Исследование",
            "confirmed": "Подтверждён",
            "isHidden": "Скрыт",
            "added": "Добавлены",
            "removed": "Сняты",
            "employeeFio": "Согласующий",
            "status": "Статус",
            "comment": "Комментарий",
            "caseId": "Дело",
            "blockId": "Блок",
            "favorite": "Избранное",
            "topic": "Тема",
            "fileName": "Файл",
            "error": "Ошибка",
            "message": "Сообщение",
            "id": "Код",
            "isPublished": "Опубликован",
            "inMyCases": "В моих делах",
            "deleted": "Удалён",
            "memberId": "Участник",
            "closedAt": "Закрыто",
        }
        yesno = {True: "да", False: "нет"}

        def one(lg, label):
            event = [["title", f"{strdatetime(lg.time)} {label}"]]
            if lg.user:
                podr = lg.user.podrazdeleniye.title if lg.user.podrazdeleniye else ""
                username = lg.user.user.username if lg.user.user else ""
                event.append(["Пользователь", f"{lg.user.fio}, {username}, {podr}"])
            if lg.body:
                try:
                    parsed = json.loads(lg.body)
                except Exception:
                    parsed = None
                if isinstance(parsed, dict):
                    for key, value in parsed.items():
                        if isinstance(value, (dict, list)):
                            event.append(["json_data", json.dumps(value, ensure_ascii=False)])
                            continue
                        if isinstance(value, bool):
                            value = yesno[value]
                        if value in ("", None):
                            continue
                        event.append([body_labels.get(key, key), str(value)])
                elif parsed is not None:
                    event.append(["json_data", lg.body])
                else:
                    event.append(["Данные", lg.body])
            return event

        packed = []
        fresh = DocumentManagerLog.objects.filter(document_id=pk).select_related("user__user", "user__podrazdeleniye")
        for lg in fresh:
            packed.append((lg.time, lg.pk, one(lg, lg.get_action_display())))
        logs = Log.objects.filter(key=str(pk), type__in=cls.LOG_TYPES).select_related("user__user", "user__podrazdeleniye")
        for lg in logs:
            packed.append((lg.time, lg.pk, one(lg, lg.get_type_display())))
        packed.sort(key=lambda item: (item[0], item[1]))
        return [{"type": f"Документ №{pk}", "events": [item[2] for item in packed]}]


class DocumentEmployeeApprove(models.Model):
    document = models.ForeignKey(Documents, db_index=True, on_delete=models.CASCADE)
    employee = models.ForeignKey(DoctorProfile, related_name="employee_approve", db_index=True, default=None, blank=True, null=True, help_text="Согласующий", on_delete=models.SET_NULL)
    time_approve = models.DateTimeField(null=True, blank=True, db_index=True, help_text="Время согласования исполнителем")
    create_at_approve = models.DateTimeField(null=True, blank=True, db_index=True, help_text="Время запроса согласования")
    who_create = models.ForeignKey(
        DoctorProfile, related_name="who_create_to_approve", db_index=True, default=None, blank=True, null=True, help_text="кто запросил согласование", on_delete=models.SET_NULL
    )
    is_cancel = models.BooleanField(default=True, blank=True, db_index=True)
    note_cancel = models.CharField(max_length=130, blank=True, help_text="Причина отмены")

    class Meta:
        verbose_name = "Согласование документа"
        verbose_name_plural = "Согласования документов"

    def __str__(self):
        return f"{self.document} {self.employee} {self.time_approve} {self.is_cancel}"

    @classmethod
    def is_assigned(cls, document, who):
        if not document or not who:
            return False
        return cls.objects.filter(document=document, employee_id=who.pk).filter(models.Q(is_cancel=False) | models.Q(time_approve__isnull=False)).exists()

    @staticmethod
    def _status_code(row):
        if row.time_approve and not row.is_cancel:
            return "approved"
        if row.is_cancel:
            return "rejected"
        return "waiting"

    @staticmethod
    def _prefer(prev, row):
        if prev is None:
            return row
        if prev.is_cancel != row.is_cancel:
            return row if not row.is_cancel else prev
        if bool(row.time_approve) != bool(prev.time_approve):
            return row if row.time_approve else prev
        return row if row.pk > prev.pk else prev

    @classmethod
    def can_cancel(cls, who):
        if not who:
            return False
        return Documents.CANCEL_APPROVE_GROUP in Documents._who_group_names(who, Documents.CANCEL_APPROVE_GROUP)

    @classmethod
    def chosen_row(cls, document, employee_id):
        chosen = None
        rows = cls.objects.filter(document=document, employee_id=employee_id).select_related("employee").order_by("pk")
        for row in rows:
            chosen = cls._prefer(chosen, row)
        return chosen

    @classmethod
    def statuses_json(cls, document):
        rows = list(cls.objects.filter(document=document, employee_id__isnull=False).select_related("employee").order_by("pk"))
        from laboratory.utils import strdatetimeru

        chosen = {}
        order = []
        for row in rows:
            prev = chosen.get(row.employee_id)
            if prev is None:
                chosen[row.employee_id] = row
                order.append(row.employee_id)
                continue
            chosen[row.employee_id] = cls._prefer(prev, row)
        result = []
        for employee_id in order:
            row = chosen[employee_id]
            if not row.employee:
                continue
            status = cls._status_code(row)
            decided_at = ""
            if status != "waiting":
                decided_at = strdatetimeru(row.time_approve) or ""
            result.append(
                {
                    "id": row.employee_id,
                    "fio": row.employee.get_fio(),
                    "status": status,
                    "decidedAt": decided_at,
                    "comment": (row.note_cancel or "") if status == "rejected" else "",
                }
            )
        return result

    @classmethod
    def approve(cls, document, who, comment=None):
        if not document or not who or not Documents.can_see_document(document, who):
            return {"ok": False, "message": "Документ не найден"}
        text = comment.strip() if isinstance(comment, str) else ""
        if len(text) > 130:
            return {"ok": False, "message": "Комментарий длиннее 130 символов"}
        from django.utils import timezone

        now = timezone.now()
        waiting = cls.objects.filter(document=document, employee=who, is_cancel=False, time_approve__isnull=True)
        if not waiting.exists():
            return {"ok": False, "message": "Согласование не требуется"}
        if text:
            waiting.update(time_approve=now, is_cancel=True, note_cancel=text)
        else:
            waiting.update(time_approve=now, is_cancel=False, note_cancel="")
        cache = getattr(document, "_prefetched_objects_cache", None)
        if cache is not None:
            cache.pop("documentemployeeapprove_set", None)
        extra = {"employeeFio": who.get_fio()}
        log_type = Documents.LOG_APPROVE
        if text:
            extra["comment"] = text
            log_type = Documents.LOG_REJECT
        Documents.log_action(document.pk, log_type, who, extra)
        return {"ok": True, "approvalStatuses": cls.statuses_json(document)}

    @classmethod
    def cancel(cls, document, who, employee_id):
        if not document or not who or not Documents.can_see_document(document, who):
            return {"ok": False, "message": "Документ не найден"}
        if not cls.can_cancel(who):
            return {"ok": False, "message": "Нет прав"}
        try:
            employee_id = int(employee_id)
        except (TypeError, ValueError):
            return {"ok": False, "message": "Согласующий не найден"}
        row = cls.chosen_row(document, employee_id)
        if not row or not row.employee_id:
            return {"ok": False, "message": "Согласующий не найден"}
        status = cls._status_code(row)
        if status == "waiting":
            return {"ok": False, "message": "Решения ещё нет"}
        status_label = "Согласовано" if status == "approved" else "Не согласовано"
        extra = {"employeeFio": row.employee.get_fio() if row.employee else "", "status": status_label}
        if row.note_cancel:
            extra["comment"] = row.note_cancel
        cls.objects.filter(document=document, employee_id=employee_id, is_cancel=False).update(time_approve=None, note_cancel="")
        if status == "rejected":
            cls.objects.filter(pk=row.pk).update(is_cancel=False, time_approve=None, note_cancel="")
        cache = getattr(document, "_prefetched_objects_cache", None)
        if cache is not None:
            cache.pop("documentemployeeapprove_set", None)
        Documents.log_action(document.pk, Documents.LOG_APPROVE_CANCEL, who, extra)
        return {"ok": True, "approvalStatuses": cls.statuses_json(document)}

    @staticmethod
    def can_edit(document, who):
        if not document or not who:
            return False
        if document.time_confirm:
            return False
        cache = getattr(document, "_prefetched_objects_cache", None) or {}
        prefetched_iss = cache.get("issledovaniya_set")
        if prefetched_iss is not None:
            if any(iss.time_confirmation for iss in prefetched_iss):
                return False
        elif Documents.document_is_confirmed(document):
            return False
        if DocumentCase._is_superuser(who):
            return True
        return document.who_create_id == who.pk

    @classmethod
    def _active_rows(cls, document):
        cache = getattr(document, "_prefetched_objects_cache", None) or {}
        prefetched = cache.get("documentemployeeapprove_set")
        if prefetched is not None:
            return [row for row in prefetched if not row.is_cancel and row.employee_id]
        return list(cls.objects.filter(document=document, is_cancel=False, employee_id__isnull=False).select_related("employee", "employee__podrazdeleniye", "who_create"))

    @classmethod
    def approvers_json(cls, document):
        from laboratory.utils import strdatetimeru

        seen = set()
        result = []
        for row in cls._active_rows(document):
            if row.employee_id in seen or not row.employee:
                continue
            seen.add(row.employee_id)
            item = AddresseeGroup._employee_json(row.employee)
            item["whoCreate"] = row.who_create.get_fio() if row.who_create_id and row.who_create else ""
            item["createAtApprove"] = strdatetimeru(row.create_at_approve) or ""
            result.append(item)
        return result

    @classmethod
    def set_approvers(cls, document, who, members):
        if not Documents.can_see_document(document, who):
            return {"ok": False, "message": "Документ не найден"}
        if not cls.can_edit(document, who):
            if Documents.document_is_confirmed(document):
                return {"ok": False, "message": "Документ подтверждён"}
            return {"ok": False, "message": "Нет прав"}
        if not isinstance(members, list):
            return {"ok": False, "message": "Некорректный список"}
        ids = []
        seen = set()
        for item in members:
            raw = item.get("id") if isinstance(item, dict) else item
            try:
                doctor_id = int(raw)
            except (TypeError, ValueError):
                continue
            if doctor_id > 0 and doctor_id not in seen:
                seen.add(doctor_id)
                ids.append(doctor_id)
        found = set(DoctorProfile.objects.filter(pk__in=ids).values_list("pk", flat=True))
        valid_ids = [doctor_id for doctor_id in ids if doctor_id in found]
        valid_set = set(valid_ids)
        before_rows = cls.objects.filter(document=document, is_cancel=False, employee_id__isnull=False).order_by("pk").values_list("employee_id", flat=True)
        before_ids = list(dict.fromkeys(before_rows))
        before_set = set(before_ids)
        added_ids = [doctor_id for doctor_id in valid_ids if doctor_id not in before_set]
        removed_ids = [doctor_id for doctor_id in before_ids if doctor_id not in valid_set]
        from django.utils import timezone

        now = timezone.now()
        with transaction.atomic():
            active = cls.objects.filter(document=document, is_cancel=False)
            if valid_set:
                active.exclude(employee_id__in=valid_set).delete()
            else:
                active.delete()
            existing = set(cls.objects.filter(document=document, is_cancel=False, employee_id__in=valid_set).values_list("employee_id", flat=True)) if valid_set else set()
            cls.objects.bulk_create(
                [
                    cls(
                        document=document,
                        employee_id=doctor_id,
                        who_create=who,
                        create_at_approve=now,
                        is_cancel=False,
                    )
                    for doctor_id in valid_ids
                    if doctor_id not in existing
                ]
            )
        cache = getattr(document, "_prefetched_objects_cache", None)
        if cache is not None:
            cache.pop("documentemployeeapprove_set", None)
        if added_ids or removed_ids:
            fio_by_id = {doctor.pk: doctor.get_fio() for doctor in DoctorProfile.objects.filter(pk__in=added_ids + removed_ids)}
            extra = {}
            if added_ids:
                extra["added"] = ", ".join(fio_by_id.get(doctor_id) or str(doctor_id) for doctor_id in added_ids)
            if removed_ids:
                extra["removed"] = ", ".join(fio_by_id.get(doctor_id) or str(doctor_id) for doctor_id in removed_ids)
            Documents.log_action(document.pk, Documents.LOG_APPROVERS, who, extra)
        return {"ok": True, "approvers": cls.approvers_json(document), "approvalStatuses": cls.statuses_json(document)}


class DocumentResolution(models.Model):
    document = models.ForeignKey(Documents, db_index=True, on_delete=models.CASCADE)
    employee = models.ForeignKey(DoctorProfile, related_name="employee_resolution", db_index=True, default=None, blank=True, null=True, help_text="Исполнитель", on_delete=models.SET_NULL)
    time_resolution = models.DateTimeField(null=True, blank=True, db_index=True, help_text="Время создания резолюции")
    who_create = models.ForeignKey(
        DoctorProfile, related_name="who_create_resolution", db_index=True, default=None, blank=True, null=True, help_text="Кто создал резолюцию", on_delete=models.SET_NULL
    )
    is_cancel = models.BooleanField(default=True, blank=True, db_index=True)
    note_cancel = models.CharField(max_length=128, blank=True, help_text="Причина отмены")
    note_resolution = models.CharField(max_length=255, blank=True, help_text="Содержание резолюции")
    parent_resolution = models.ForeignKey("self", db_index=True, related_name="resolution_p", help_text="Резолюция основание", blank=True, null=True, default=None, on_delete=models.SET_NULL)

    class Meta:
        verbose_name = "Резолюция документа"
        verbose_name_plural = "Резолюции документов"

    def __str__(self):
        return f"{self.document} {self.employee} {self.time_resolution} {self.is_cancel}"


class DocumentControl(models.Model):
    document = models.ForeignKey(Documents, db_index=True, on_delete=models.CASCADE)
    time_control = models.DateTimeField(null=True, blank=True, db_index=True, help_text="Дата контроля")
    resolution = models.ForeignKey(DocumentResolution, db_index=True, on_delete=models.CASCADE, help_text="Резолюция")
    is_completed = models.BooleanField(default=False, blank=True, db_index=True, help_text="Выполнена резолюция")

    class Meta:
        verbose_name = "Контроль документа"
        verbose_name_plural = "Контроли документов"

    def __str__(self):
        return f"{self.document} {self.document} {self.time_control}"


class DocumentReview(models.Model):
    document = models.ForeignKey(Documents, db_index=True, on_delete=models.CASCADE)
    time_review = models.DateTimeField(null=True, blank=True, db_index=True, help_text="Дата ознакомления")
    doctor_review = models.ForeignKey(DoctorProfile, default=None, blank=True, null=True, help_text="Кто  должен ознакомлен", on_delete=models.SET_NULL)

    class Meta:
        verbose_name = "Ознакомление документа"
        verbose_name_plural = "Ознакомления документов"

    def __str__(self):
        return f"{self.document} {self.doctor_review} {self.time_review}"

    @classmethod
    def pending_qs(cls, who):
        qs = cls.objects.filter(time_review__isnull=True, document__is_hidden=False).filter(Documents.case_access_q(who, "document__"))
        if who:
            qs = qs.filter(doctor_review=who)
        else:
            qs = qs.none()
        return qs.distinct()

    @classmethod
    def pending_document_ids(cls, who):
        return cls.pending_qs(who).values("document_id")

    @classmethod
    def pending_count(cls, who):
        if not who:
            return 0
        return cls.pending_qs(who).values("document_id").distinct().count()

    @classmethod
    def mark_opened(cls, document, who, iss=None):
        if not document or not who:
            return False
        if not document.time_confirm:
            if iss is None:
                iss = document.get_issledovaniye()
            if not iss or not iss.time_confirmation:
                return False
        from django.utils import timezone

        updated = cls.objects.filter(document=document, doctor_review=who, time_review__isnull=True).update(time_review=timezone.now())
        return updated > 0

    @classmethod
    def parse_addressee_ids(cls, value):
        if isinstance(value, str):
            try:
                value = json.loads(value)
            except Exception:
                value = []
        if not isinstance(value, list):
            return []
        ids = []
        seen = set()
        for item in value:
            raw = item.get("id") if isinstance(item, dict) else item
            try:
                doctor_id = int(raw)
            except (TypeError, ValueError):
                continue
            if doctor_id < 1 or doctor_id in seen:
                continue
            seen.add(doctor_id)
            ids.append(doctor_id)
        return ids

    @classmethod
    def collect_ids_from_research(cls, research):
        ids = []
        seen = set()
        for group in (research or {}).get("groups") or []:
            for field in group.get("fields") or []:
                field_type = field.get("field_type")
                if field_type != ADDRESSEE_FIELD_TYPE:
                    continue
                for doctor_id in cls.parse_addressee_ids(field.get("value")):
                    if doctor_id in seen:
                        continue
                    seen.add(doctor_id)
                    ids.append(doctor_id)
        return ids

    @classmethod
    def sync_for_document(cls, document, doctor_ids):
        wanted = set(doctor_ids or [])
        if wanted:
            wanted = set(DoctorProfile.objects.filter(pk__in=wanted).values_list("pk", flat=True))
        existing = list(cls.objects.filter(document=document))
        keep = set()
        to_delete = []
        for row in existing:
            if row.doctor_review_id and row.doctor_review_id in wanted:
                keep.add(row.doctor_review_id)
            else:
                to_delete.append(row.pk)
        if to_delete:
            cls.objects.filter(pk__in=to_delete).delete()
        cls.objects.bulk_create([cls(document=document, doctor_review_id=doctor_id) for doctor_id in wanted if doctor_id not in keep])

    @classmethod
    def sync_from_research(cls, document, research):
        cls.sync_for_document(document, cls.collect_ids_from_research(research))


class DocumentRecent(models.Model):
    PAGE_SIZE = 50
    MAX_ITEMS = 200

    doctor = models.ForeignKey(DoctorProfile, related_name="document_recents", on_delete=models.CASCADE, help_text="Пользователь")
    document = models.ForeignKey(Documents, related_name="recent_opens", on_delete=models.CASCADE, help_text="Документ")
    topic = models.CharField(max_length=255, blank=True, default="", help_text="Тема документа")
    type_doc = models.CharField(max_length=128, blank=True, default="", help_text="Вид документа")
    opened_at = models.DateTimeField(db_index=True, help_text="Дата и время последнего открытия")

    class Meta:
        verbose_name = "Последний документ"
        verbose_name_plural = "Последние документы"
        unique_together = ("doctor", "document")
        indexes = (models.Index(fields=("doctor", "-opened_at")),)
        ordering = ("-opened_at", "-pk")

    def __str__(self):
        return f"{self.doctor_id} {self.document_id} {self.opened_at}"

    @classmethod
    def remember(cls, who, document, topic="", type_doc=""):
        if not who or not document:
            return
        from django.utils import timezone

        defaults = {
            "topic": topic or "",
            "type_doc": type_doc or "",
            "opened_at": timezone.now(),
        }
        if cls.objects.filter(doctor=who, document=document).update(**defaults):
            return
        try:
            cls.objects.create(doctor=who, document=document, **defaults)
        except IntegrityError:
            cls.objects.filter(doctor=who, document=document).update(**defaults)
            return
        extra_ids = list(cls.objects.filter(doctor=who).order_by("-opened_at", "-pk").values_list("pk", flat=True)[cls.MAX_ITEMS :])
        if extra_ids:
            cls.objects.filter(pk__in=extra_ids).delete()

    @classmethod
    def get_page(cls, who, page=1, page_size=None):
        if not who:
            return {"result": [], "page": 1, "pageSize": cls.PAGE_SIZE, "hasMore": False}
        try:
            page = max(int(page or 1), 1)
        except (TypeError, ValueError):
            page = 1
        try:
            page_size = max(int(page_size or cls.PAGE_SIZE), 1)
        except (TypeError, ValueError):
            page_size = cls.PAGE_SIZE
        qs = (
            cls.objects.filter(doctor=who, document__is_hidden=False, document__time_confirm__isnull=False)
            .filter(Documents.case_access_q(who, "document__"))
            .select_related("document", "document__type_document", "document__schema")
            .order_by("-opened_at", "-pk")
            .distinct()
        )
        offset = (page - 1) * page_size
        rows = list(qs[offset : offset + page_size + 1])
        has_more = len(rows) > page_size
        rows = rows[:page_size]
        from directions.models import Issledovaniya

        docs = [row.document for row in rows]
        iss_by_doc = {}
        if docs:
            for iss in Issledovaniya.objects.filter(document_id__in=[doc.pk for doc in docs]).order_by("pk"):
                iss_by_doc.setdefault(iss.document_id, iss)
        live_topics = Documents.topics_for_documents(docs, iss_by_doc)
        result = []
        for row in rows:
            topic = (live_topics.get(row.document_id) or row.topic or "").strip()
            type_doc = (row.type_doc or "").strip()
            if not type_doc and row.document.type_document:
                type_doc = row.document.type_document.title
            title = " ".join(part for part in (topic, type_doc) if part)
            result.append(
                {
                    "id": row.document_id,
                    "title": Documents.title_with_id(row.document_id, title),
                    "topic": topic,
                    "typeDoc": type_doc,
                }
            )
        return {
            "result": result,
            "page": page,
            "pageSize": page_size,
            "hasMore": has_more,
        }


class DocumentResult(models.Model):
    document = models.ForeignKey(Documents, db_index=True, default=None, blank=True, null=True, help_text="Тип документа", on_delete=models.SET_NULL)
    create_at = models.DateTimeField(auto_now_add=True, help_text="Дата создания/изменения", db_index=True)
    who_create = models.ForeignKey(DoctorProfile, default=None, blank=True, null=True, help_text="Создатель", on_delete=models.SET_NULL)
    value = models.TextField(null=True, blank=True, help_text="Значение")
    document_field = models.ForeignKey(DocumentFields, db_index=True, default=None, blank=True, null=True, help_text="поле документа", on_delete=models.SET_NULL)
    field_type = models.SmallIntegerField(default=None, blank=True, choices=DocumentFields.TYPES, null=True)
    meta_document_field = models.ForeignKey(MetaFieldsForAllTypeDocuments, db_index=True, default=None, blank=True, null=True, help_text="поле документа", on_delete=models.SET_NULL)
    meta_field_type = models.SmallIntegerField(default=None, blank=True, choices=MetaFieldsForAllTypeDocuments.TYPES, null=True)

    class Meta:
        verbose_name = "Реквизит для документа"
        verbose_name_plural = "Реквизит для документов"

    def __str__(self):
        return f"{self.document}"


def get_file_path(instance: "DocumentFiles", filename):
    return os.path.join("document_files", str(instance.document.pk), str(uuid.uuid4()), filename)


class DocumentManagerLog(models.Model):
    CREATE = 260000
    SAVE = 260001
    CONFIRM = 260002
    RESET = 260003
    HIDE = 260004
    SHOW = 260005
    APPROVERS = 260006
    APPROVE = 260007
    REJECT = 260008
    APPROVE_CANCEL = 260009
    CASE_ASSIGN = 260010
    BLOCK_ASSIGN = 260011
    CASE_TYPE = 260012
    CASE_ACCESS = 260013
    CASE_COMMENT = 260014
    CASE_TOPIC = 260015
    CASE_CLOSE = 260016
    MY_CASE = 260017
    BLOCK_CREATE = 260018
    BLOCK_UPDATE = 260019
    BLOCK_READ = 260020
    BLOCK_WRITE = 260021
    FAVORITE = 260022
    CASE_FAVORITE = 260023
    GROUP = 260024
    TYPE = 260025
    PRINT_TEMPLATE = 260026
    SECTION_TYPE = 260027
    PLACE = 260028
    PLAN = 260029
    STRUCTURE = 260030
    ADDRESSEE_GROUP = 260031
    ADDRESSEE_MEMBER = 260032
    PRINT_ERROR = 260033
    PICKER_FAVORITE = 260034
    TEST = 260035
    ACTIONS = (
        (CREATE, "ДОУ: создание"),
        (SAVE, "ДОУ: сохранение"),
        (CONFIRM, "ДОУ: подтверждение"),
        (RESET, "ДОУ: сброс подтверждения"),
        (HIDE, "ДОУ: скрытие"),
        (SHOW, "ДОУ: показ"),
        (APPROVERS, "ДОУ: согласующие"),
        (APPROVE, "ДОУ: согласование"),
        (REJECT, "ДОУ: несогласие"),
        (APPROVE_CANCEL, "ДОУ: отмена согласования"),
        (CASE_ASSIGN, "ДОУ: дело документа"),
        (BLOCK_ASSIGN, "ДОУ: блок документа"),
        (CASE_TYPE, "ДОУ: вид дела"),
        (CASE_ACCESS, "ДОУ: доступ к делу"),
        (CASE_COMMENT, "ДОУ: комментарий дела"),
        (CASE_TOPIC, "ДОУ: тема дела"),
        (CASE_CLOSE, "ДОУ: закрытие дела"),
        (MY_CASE, "ДОУ: мои дела"),
        (BLOCK_CREATE, "ДОУ: создание блока"),
        (BLOCK_UPDATE, "ДОУ: изменение блока"),
        (BLOCK_READ, "ДОУ: чтение блока"),
        (BLOCK_WRITE, "ДОУ: запись блока"),
        (FAVORITE, "ДОУ: избранное"),
        (CASE_FAVORITE, "ДОУ: избранное дела"),
        (GROUP, "ДОУ: группа документов"),
        (TYPE, "ДОУ: вид документа"),
        (PRINT_TEMPLATE, "ДОУ: шаблон печати"),
        (SECTION_TYPE, "ДОУ: тип раздела"),
        (PLACE, "ДОУ: место"),
        (PLAN, "ДОУ: план"),
        (STRUCTURE, "ДОУ: структура"),
        (ADDRESSEE_GROUP, "ДОУ: набор адресатов"),
        (ADDRESSEE_MEMBER, "ДОУ: участник набора"),
        (PRINT_ERROR, "ДОУ: ошибка печати"),
        (PICKER_FAVORITE, "ДОУ: избранное подбора"),
        (TEST, "ДОУ: тест"),
    )

    action = models.IntegerField(choices=ACTIONS, db_index=True)
    document = models.ForeignKey(Documents, related_name="manager_logs", blank=True, null=True, on_delete=models.SET_NULL)
    document_case = models.ForeignKey(DocumentCase, related_name="manager_logs", blank=True, null=True, on_delete=models.SET_NULL)
    block = models.ForeignKey(DocumentBlock, related_name="manager_logs", blank=True, null=True, on_delete=models.SET_NULL)
    object_key = models.CharField(max_length=64, blank=True, default="", db_index=True)
    body = models.TextField(blank=True, default="")
    user = models.ForeignKey(DoctorProfile, related_name="document_manager_logs", blank=True, null=True, on_delete=models.SET_NULL)
    time = models.DateTimeField(auto_now_add=True, db_index=True)

    class Meta:
        verbose_name = "Журнал ДОУ"
        verbose_name_plural = "Журнал ДОУ"
        ordering = ("-time", "-pk")

    def __str__(self):
        return f"{self.get_action_display()} {self.document_id or self.object_key}"

    @staticmethod
    def _pk(value):
        if value in (None, "", 0, "0"):
            return None
        try:
            number = int(value)
        except (TypeError, ValueError):
            return None
        return number if number > 0 else None

    @classmethod
    def write(cls, action, who=None, body=None, document_id=None, case_id=None, block_id=None, object_key=""):
        payload = body if isinstance(body, dict) else {}
        cls.objects.create(
            action=action,
            user=who if getattr(who, "pk", None) else None,
            body=json.dumps(payload, ensure_ascii=False),
            document_id=cls._pk(document_id),
            document_case_id=cls._pk(case_id),
            block_id=cls._pk(block_id),
            object_key=str(object_key or "")[:64],
        )


class DocumentFiles(models.Model):
    document = models.ForeignKey(Documents, db_index=True, on_delete=models.CASCADE)
    uploaded_file = models.FileField(upload_to=get_file_path, blank=True, null=True, default=None)
    created_at = models.DateTimeField(auto_now_add=True)
    type_file = models.CharField(max_length=16, blank=True, help_text="Тип документа")
    who_add_files = models.ForeignKey(DoctorProfile, default=None, blank=True, null=True, related_name="document_who_add_files", help_text="Создатель направления", on_delete=models.SET_NULL)
    comment = models.CharField(max_length=130, blank=True, help_text="Комментарий")

    class Meta:
        verbose_name = "Файл на документ"
        verbose_name_plural = "Файлы на документы"
