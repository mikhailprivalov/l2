from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter

from api.gardening.views import _all_plots_export


def _cell_text(row, col):
    raw = row.get(col["key"])
    if col.get("empty_ok") and raw == "":
        return ""
    if col.get("missing_zero") and raw in (None, ""):
        return "0.00"
    if raw in (None, ""):
        return "—"
    return str(raw)


def form_02(request_data) -> Workbook:
    payload = _all_plots_export(request_data)
    columns = payload.get("columns") or []
    rows = payload.get("rows") or []
    title = payload.get("title") or "Учёт"

    work_book = Workbook()
    sheet = work_book[work_book.sheetnames[0]]
    sheet.title = "Учёт"
    sheet.append([title])
    header_font = Font(bold=True)
    title_cell = sheet.cell(row=1, column=1)
    title_cell.font = Font(bold=True, size=12)
    if columns:
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))

    sheet.append([col.get("label") or "" for col in columns] or ["Нет данных"])
    for cell in sheet[2]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    thin = Border(
        left=Side(style="thin", color="B1B1B1"),
        right=Side(style="thin", color="B1B1B1"),
        top=Side(style="thin", color="B1B1B1"),
        bottom=Side(style="thin", color="B1B1B1"),
    )
    if not columns:
        sheet.append(["Нет данных"])
    elif not rows:
        empty = ["Нет данных"] + [""] * (len(columns) - 1)
        sheet.append(empty)
    else:
        for row in rows:
            sheet.append([_cell_text(row, col) for col in columns])

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=1, max_col=max(len(columns), 1)):
        for idx, cell in enumerate(row):
            cell.border = thin
            cell.alignment = Alignment(vertical="center")
            col = columns[idx] if idx < len(columns) else None
            if col and col.get("numeric"):
                cell.alignment = Alignment(horizontal="right", vertical="center")

    widths = []
    for col in columns:
        key = col.get("key")
        if key in ("owner", "title", "period"):
            widths.append(28)
        elif key == "meter_title":
            widths.append(16)
        elif key == "num_object":
            widths.append(12)
        else:
            widths.append(14)
    if not widths:
        widths = [40]
    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width

    return work_book
