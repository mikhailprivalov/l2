import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from openpyxl import load_workbook

from clients.models import Individual
from directory.models import (
    GardeningElectricityMeter,
    GardeningElectricityMeterReading,
    OwnersRealEstate,
    RealEstate,
)

PLACEHOLDER_BIRTHDAY = date(1900, 1, 1)
READING_ORDER_ERROR = "Текущее показание не может быть меньше предыдущего"
MISSING_READING = "отсутствуют"
PLOT_NUMBER_MAX_LEN = 64

RU_MONTHS = {
    "январь": 1,
    "января": 1,
    "февраль": 2,
    "февраля": 2,
    "март": 3,
    "марта": 3,
    "апрель": 4,
    "апреля": 4,
    "май": 5,
    "мая": 5,
    "июнь": 6,
    "июня": 6,
    "июль": 7,
    "июля": 7,
    "август": 8,
    "августа": 8,
    "сентябрь": 9,
    "сентября": 9,
    "октябрь": 10,
    "октября": 10,
    "ноябрь": 11,
    "ноября": 11,
    "декабрь": 12,
    "декабря": 12,
}

COMPOUND_INITIALS_RE = re.compile(r"^([A-Za-zА-Яа-яЁё])\.?([A-Za-zА-Яа-яЁё])\.?$")
DATE_RE = re.compile(r"(\d{2})\.(\d{2})\.(\d{4})")
TITLE_MONTH_RE = re.compile(r"([А-Яа-яЁё]+)\s+(\d{4})")


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    return str(value).replace("\xa0", " ").replace("\u00a0", " ").strip()


def normalize_plot_number(raw):
    text = cell_text(raw)
    if not text or text.lower() == "none":
        return ""
    text = re.sub(r"\s+", " ", text).strip().upper()
    text = re.sub(r"^(\d+)\s+([A-ZА-ЯЁ])$", r"\1\2", text)
    return text[:PLOT_NUMBER_MAX_LEN]


def plot_sort_key(num_object):
    if num_object in (None, ""):
        return (1, 10**18, "")
    text = str(num_object)
    match = re.match(r"^(\d+)(.*)$", text)
    if match:
        return (0, int(match.group(1)), match.group(2))
    return (0, 10**18, text)


def ordered_real_estates():
    estates = list(RealEstate.objects.filter(hide=False))
    estates.sort(key=lambda item: (plot_sort_key(item.num_object), item.pk))
    return estates


def normalize_serial(raw):
    text = cell_text(raw)
    if not text or text.lower() == "none":
        return ""
    return re.sub(r"\s+", " ", text).strip()[:255]


def parse_subscriber_fio(raw):
    text = cell_text(raw)
    if not text or text.lower() == "none":
        return "", "", ""
    text = re.sub(r"\s+", " ", text).strip()
    parts = [part for part in text.split(" ") if part]
    if not parts:
        return "", "", ""
    if len(parts) == 1:
        return parts[0], "", ""
    if len(parts) == 2:
        initials = COMPOUND_INITIALS_RE.match(parts[1].replace(" ", ""))
        if initials:
            return parts[0], f"{initials.group(1)}.", f"{initials.group(2)}."
        return parts[0], parts[1], ""
    return parts[0], parts[1], " ".join(parts[2:])


def subscriber_is_plot_number(subscriber, plot_number):
    return normalize_plot_number(subscriber) == plot_number and bool(plot_number)


def parse_reading_value(raw):
    if isinstance(raw, Decimal):
        value = raw
    elif isinstance(raw, int):
        value = Decimal(raw)
    elif isinstance(raw, float):
        value = Decimal(str(raw))
    else:
        text = cell_text(raw)
        if not text or text.lower() == "none" or MISSING_READING in text.lower():
            return None
        text = text.replace(" ", "").replace(",", ".")
        try:
            value = Decimal(text)
        except (InvalidOperation, TypeError, ValueError):
            return None
    if value < 0:
        return None
    return value


def previous_year_month(year, month):
    if month == 1:
        return year - 1, 12
    return year, month - 1


def parse_cell_date(raw):
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = cell_text(raw)
    match = DATE_RE.search(text)
    if not match:
        return None
    try:
        return date(int(match.group(3)), int(match.group(2)), int(match.group(1)))
    except ValueError:
        return None


def parse_title_period(raw):
    text = cell_text(raw).lower()
    match = TITLE_MONTH_RE.search(text)
    if not match:
        return None, None
    month = RU_MONTHS.get(match.group(1))
    if not month:
        return None, None
    try:
        return int(match.group(2)), month
    except ValueError:
        return None, None


def _header_text(value):
    return re.sub(r"\s+", " ", cell_text(value).lower())


def _header_index(cells, predicate):
    for idx, value in enumerate(cells):
        if predicate(_header_text(value)):
            return idx
    return None


def _detect_layout(rows):
    account_idx, address_idx, subscriber_idx, device_idx, serial_idx = 0, 1, 2, 3, 4
    previous_idx, current_idx = 5, 7
    year, month = None, None
    data_start = 0
    dates = []

    for row_idx, row in enumerate(rows[:12]):
        cells = list(row)
        found_account = _header_index(cells, lambda text: "лицевого" in text)
        if found_account is not None:
            account_idx = found_account
            found_address = _header_index(cells, lambda text: "адрес" in text)
            if found_address is not None:
                address_idx = found_address
            found_subscriber = _header_index(cells, lambda text: text == "абонент")
            if found_subscriber is not None:
                subscriber_idx = found_subscriber
            found_device = _header_index(cells, lambda text: text == "тип" or text.startswith("тип "))
            if found_device is not None:
                device_idx = found_device
            found_serial = _header_index(cells, lambda text: "серийный" in text)
            if found_serial is not None:
                serial_idx = found_serial
            data_start = max(data_start, row_idx + 1)
        title_year, title_month = parse_title_period(cells[0] if cells else "")
        if title_year and title_month:
            year, month = title_year, title_month
        for col_idx, value in enumerate(cells):
            parsed = parse_cell_date(value)
            if parsed:
                dates.append((parsed, col_idx))

    dates.sort(key=lambda item: item[0])
    unique_dates = []
    seen = set()
    for parsed, col_idx in dates:
        key = (parsed, col_idx)
        if key in seen:
            continue
        seen.add(key)
        unique_dates.append((parsed, col_idx))
    if len(unique_dates) >= 2:
        previous_idx = unique_dates[0][1]
        current_idx = unique_dates[-1][1]
        start = unique_dates[0][0]
        year, month = start.year, start.month
    elif unique_dates:
        current_idx = unique_dates[0][1]
        year, month = unique_dates[0][0].year, unique_dates[0][0].month

    header_words = {"АБОНЕНТА", "ПРИБОРА", "ЛИЦЕВОГО СЧЁТА", "ЛИЦЕВОГО СЧЕТА"}
    for row_idx in range(data_start, min(len(rows), data_start + 8)):
        cells = list(rows[row_idx])
        plot_number = normalize_plot_number(_cell(cells, account_idx))
        if plot_number and plot_number not in header_words:
            data_start = row_idx
            break

    return {
        "account_idx": account_idx,
        "address_idx": address_idx,
        "subscriber_idx": subscriber_idx,
        "device_idx": device_idx,
        "serial_idx": serial_idx,
        "previous_idx": previous_idx,
        "current_idx": current_idx,
        "year": year,
        "month": month,
        "data_start": data_start,
    }


def _cell(cells, idx):
    if idx is None or idx < 0 or idx >= len(cells):
        return None
    return cells[idx]


def _get_or_create_plot(plot_number):
    created = False
    obj = RealEstate.objects.filter(num_object=plot_number, hide=False).first()
    if obj:
        return obj, created
    hidden = RealEstate.objects.filter(num_object=plot_number, hide=True).first()
    if hidden:
        hidden.hide = False
        hidden.title = plot_number
        hidden.save(update_fields=["hide", "title"])
        return hidden, False
    obj = RealEstate.objects.create(title=plot_number, num_object=plot_number)
    return obj, True


def _meters(real_estate):
    return list(GardeningElectricityMeter.objects.filter(real_estate=real_estate, hide=False).order_by("sort_weight", "pk"))


def meter_title_with_serial(title, serial_number):
    title = (title or "").strip()
    serial = (serial_number or "").strip()
    if serial:
        if not title or title == serial:
            return serial[:255]
        if f"({serial})" in title:
            return title[:255]
        return f"{title} ({serial})"[:255]
    return title[:255]


def _get_or_create_meter(real_estate, serial_number, address, subscriber, device_type):
    created = False
    meters = _meters(real_estate)
    meter = None
    if serial_number:
        for item in meters:
            if normalize_serial(item.serial_number) == serial_number:
                meter = item
                break
    if meter is None and serial_number:
        empty_serial = next((item for item in meters if not normalize_serial(item.serial_number)), None)
        if empty_serial is not None:
            meter = empty_serial
    if meter is None and not serial_number and meters:
        meter = meters[0]
    if meter is None:
        title = meter_title_with_serial(f"Счётчик {len(meters) + 1}", serial_number)
        sort_weight = (meters[-1].sort_weight if meters else -1) + 1
        meter = GardeningElectricityMeter.objects.create(
            real_estate=real_estate,
            title=title,
            sort_weight=sort_weight,
            subscriber_address=address,
            subscriber=subscriber,
            device_type=device_type,
            serial_number=serial_number,
        )
        created = True
        return meter, created

    title = meter_title_with_serial(meter.title, serial_number)

    update_fields = []
    if meter.title != title:
        meter.title = title
        update_fields.append("title")
    if address and meter.subscriber_address != address:
        meter.subscriber_address = address
        update_fields.append("subscriber_address")
    if subscriber and meter.subscriber != subscriber:
        meter.subscriber = subscriber
        update_fields.append("subscriber")
    if device_type and meter.device_type != device_type:
        meter.device_type = device_type
        update_fields.append("device_type")
    if serial_number and normalize_serial(meter.serial_number) != serial_number:
        meter.serial_number = serial_number
        update_fields.append("serial_number")
    if update_fields:
        meter.save(update_fields=update_fields)
    return meter, created


def _upsert_owner(real_estate, family, name, patronymic):
    family = (family or "")[:120]
    name = (name or "")[:120]
    patronymic = (patronymic or "")[:120]
    if not family:
        return False
    owner = OwnersRealEstate.objects.select_related("individual").filter(real_estate=real_estate, hide=False, date_end__isnull=True).order_by("date_start", "pk").first()
    if owner:
        individual = owner.individual
        if individual:
            if (individual.family or "").strip():
                return False
            individual.family = family
            individual.name = name
            individual.patronymic = patronymic
            individual.save(update_fields=["family", "name", "patronymic"])
            return False
        individual = Individual.objects.create(family=family, name=name, patronymic=patronymic, birthday=PLACEHOLDER_BIRTHDAY)
        owner.individual = individual
        owner.save(update_fields=["individual"])
        return True
    individual = Individual.objects.create(family=family, name=name, patronymic=patronymic, birthday=PLACEHOLDER_BIRTHDAY)
    OwnersRealEstate.objects.create(real_estate=real_estate, individual=individual)
    return True


def _upsert_reading(real_estate, meter, year, month, reading):
    existing = GardeningElectricityMeterReading.objects.filter(real_estate=real_estate, meter=meter, year=year, month=month, hide=False).first()
    if existing:
        if existing.reading == reading:
            return "unchanged"
        existing.reading = reading
        existing.save(update_fields=["reading"])
        return "updated"
    hidden = GardeningElectricityMeterReading.objects.filter(real_estate=real_estate, meter=meter, year=year, month=month, hide=True).first()
    if hidden:
        hidden.hide = False
        hidden.reading = reading
        hidden.previous_reading_manual = None
        hidden.save(update_fields=["hide", "reading", "previous_reading_manual"])
        return "created"
    GardeningElectricityMeterReading.objects.create(
        real_estate=real_estate,
        meter=meter,
        year=year,
        month=month,
        reading=reading,
    )
    return "created"


def _load_workbook(source):
    if hasattr(source, "read"):
        payload = source.read()
        return load_workbook(filename=BytesIO(payload), data_only=True)
    return load_workbook(filename=source, data_only=True)


def import_electricity_xlsx(source):
    wb = _load_workbook(source)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    layout = _detect_layout(rows)
    year, month = layout["year"], layout["month"]
    result = {
        "year": year,
        "month": month,
        "plots_created": 0,
        "meters_created": 0,
        "owners_created": 0,
        "readings_created": 0,
        "readings_updated": 0,
        "skipped_no_reading": 0,
        "errors": [],
    }
    if not year or not month:
        result["errors"].append({"row": 1, "plot": "", "message": "Не удалось определить год и месяц отчёта"})
        return result

    header_words = {"АБОНЕНТА", "ПРИБОРА", "ЛИЦЕВОГО СЧЁТА", "ЛИЦЕВОГО СЧЕТА"}
    with transaction.atomic():
        for row_idx, row in enumerate(rows[layout["data_start"] :], start=layout["data_start"] + 1):
            cells = list(row)
            plot_number = normalize_plot_number(_cell(cells, layout["account_idx"]))
            if not plot_number or plot_number in header_words:
                continue
            address = cell_text(_cell(cells, layout["address_idx"]))[:512]
            subscriber = cell_text(_cell(cells, layout["subscriber_idx"]))[:255]
            device_type = cell_text(_cell(cells, layout["device_idx"]))[:255]
            serial_number = normalize_serial(_cell(cells, layout["serial_idx"]))
            previous_reading = parse_reading_value(_cell(cells, layout["previous_idx"]))
            current_reading = parse_reading_value(_cell(cells, layout["current_idx"]))
            try:
                real_estate, plot_created = _get_or_create_plot(plot_number)
                if plot_created:
                    result["plots_created"] += 1
                meter, meter_created = _get_or_create_meter(real_estate, serial_number, address, subscriber, device_type)
                if meter_created:
                    result["meters_created"] += 1
                if subscriber and not subscriber_is_plot_number(subscriber, plot_number):
                    family, name, patronymic = parse_subscriber_fio(subscriber)
                    if family and _upsert_owner(real_estate, family, name, patronymic):
                        result["owners_created"] += 1
                if current_reading is None:
                    result["skipped_no_reading"] += 1
                    continue
                if previous_reading is not None and current_reading < previous_reading:
                    result["errors"].append({"row": row_idx, "plot": plot_number, "message": READING_ORDER_ERROR})
                    continue
                if previous_reading is not None:
                    prev_year, prev_month = previous_year_month(year, month)
                    status = _upsert_reading(real_estate, meter, prev_year, prev_month, previous_reading)
                    if status == "created":
                        result["readings_created"] += 1
                    elif status == "updated":
                        result["readings_updated"] += 1
                status = _upsert_reading(real_estate, meter, year, month, current_reading)
                if status == "created":
                    result["readings_created"] += 1
                elif status == "updated":
                    result["readings_updated"] += 1
            except Exception as exc:
                result["errors"].append({"row": row_idx, "plot": plot_number, "message": str(exc)})
    return result
